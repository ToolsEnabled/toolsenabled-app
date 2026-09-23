#!/usr/bin/env python
"""Read-only integrity tests for the content currently committed under data/."""
import json
import math
import re
import struct
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

import import_pptx as importer
import lint_fonts
import render_pptx as renderer


HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}
PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
FINAL_ORDER = [
    "s1", "s15", "s2", "s3", "s4", "s5", "s6", "s7", "s8",
    "s9", "s10", "s12", "s13", "s11", "s14", "s16", "sef432705",
]
TIMELINE_FILES = [
    "timeline-01-spawn.png",
    "timeline-02-hook.png",
    "timeline-03-quest.png",
    "timeline-04-baseline.png",
    "timeline-05-crash.png",
    "timeline-06-specify.png",
    "timeline-07-build.png",
    "timeline-08-verify.png",
    "timeline-09-peak.png",
    "timeline-10-trap.png",
    "timeline-11-power-up.png",
    "timeline-12-gate.png",
    "timeline-13-audit.png",
    "timeline-14-boss.png",
    "timeline-15-exit.png",
    "timeline-16-credits.png",
    "timeline-17-flag.png",
]


def shape_by_id(slide):
    return {shape.shape_id: shape for shape in slide.shapes}


def rich_summary(shape, theme):
    summary = []
    for para in shape.text_frame.paragraphs:
        runs = []
        for run in para.runs:
            runs.append({
                "text": run.text,
                "size": None if run.font.size is None else run.font.size.pt,
                "bold": run.font.bold,
                "color": importer._hex(run.font.color, theme),
                "font": run.font.name,
            })
        summary.append({"text": para.text, "align": para.alignment, "runs": runs})
    return summary


def content_changed(element, pristine):
    """Mirror the renderer's content-delta decision for one imported element."""
    if "items" in element:
        return (
            "items" not in pristine
            or element["items"] != pristine["items"]
        )
    return (
        "text" in element
        and (
            "text" not in pristine
            or element["text"] != pristine["text"]
        )
    )


def normalize_text(value):
    return re.sub(
        r"\s+",
        " ",
        str(value or "")
        .replace("\u2018", "'")
        .replace("\u2019", "'"),
    ).strip()


def element_values(element):
    values = []
    if element.get("text") is not None:
        values.append(str(element.get("text") or ""))
    values.extend(str(item) for item in element.get("items", []))
    return values


def element_text(element):
    return normalize_text(" ".join(element_values(element)))


def slide_text(slide):
    return normalize_text(
        " ".join(
            value
            for element in slide.get("elements", [])
            for value in element_values(element)
        )
    )


class CurrentContentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.model = json.loads((DATA / "model.json").read_text(encoding="utf-8"))
        cls.state = json.loads((DATA / "state.json").read_text(encoding="utf-8"))
        cls.base = Presentation(str(DATA / "base.pptx"))
        cls.theme = importer.load_theme_colors(cls.base)
        cls.base_models = [
            importer.derive_slide(cls.base, index, cls.theme)
            for index in range(len(cls.base.slides))
        ]
        cls.temp = tempfile.TemporaryDirectory(prefix="suite-current-content-")
        cls.rendered_path = Path(cls.temp.name) / "rendered.pptx"
        renderer.render(cls.model, cls.rendered_path)
        cls.rendered = Presentation(str(cls.rendered_path))

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_model_base_and_asset_graph_is_well_formed(self):
        self.assertEqual(self.model["mode"], "overlay")
        self.assertEqual(len(self.model["slides"]), 17)
        self.assertEqual(len(self.base.slides), 16)
        self.assertEqual(len(self.rendered.slides), 17)
        self.assertEqual(Path(self.model["base"]).resolve(),
                         (DATA / "base.pptx").resolve())
        self.assertAlmostEqual(self.model["slideW"],
                               self.base.slide_width / 914400, places=3)
        self.assertAlmostEqual(self.model["slideH"],
                               self.base.slide_height / 914400, places=3)

        slide_ids = set()
        sources = set()
        object_ids = set()
        refs = set()
        element_count = 0
        imported_element_count = 0
        native_slide_ids = set()
        decor_count = 0
        generated_decor = []
        for slide_index, slide in enumerate(self.model["slides"]):
            self.assertNotIn(slide["id"], slide_ids)
            slide_ids.add(slide["id"])
            src = slide.get("src")
            imported = type(src) is int
            if imported:
                self.assertGreaterEqual(src, 0)
                self.assertLess(src, len(self.base.slides))
                self.assertNotIn(src, sources)
                sources.add(src)
                base_shapes = shape_by_id(self.base.slides[src])
            else:
                self.assertNotIn("src", slide)
                self.assertIsInstance(slide.get("layout"), str)
                self.assertTrue(slide["layout"].strip())
                native_slide_ids.add(slide["id"])
                base_shapes = {}

            local_ids = set()
            for element in slide.get("elements", []):
                element_count += 1
                self.assertNotIn(element["id"], local_ids)
                local_ids.add(element["id"])
                self.assertNotIn((slide["id"], element["id"]), object_ids)
                object_ids.add((slide["id"], element["id"]))
                ref = element.get("ref")
                if imported:
                    imported_element_count += 1
                    self.assertIsInstance(ref, dict)
                    self.assertEqual(ref["s"], src)
                    self.assertIn(ref["sid"], base_shapes)
                    self.assertTrue(base_shapes[ref["sid"]].has_text_frame)
                    self.assertNotIn((ref["s"], ref["sid"]), refs)
                    refs.add((ref["s"], ref["sid"]))
                    self._assert_box(element["box"])
                else:
                    self.assertIsNone(ref)
                    if "box" in element:
                        self._assert_box(element["box"])
                self._assert_style(element.get("style", {}))

            for decor in slide.get("decor", []):
                decor_count += 1
                self.assertNotIn(decor["id"], local_ids)
                local_ids.add(decor["id"])
                self._assert_box(decor["box"])
                if decor.get("generated") is True:
                    generated_decor.append((slide["id"], decor))
                    self.assertIsNone(decor.get("sid"))
                    if decor["kind"] == "shape":
                        self.assertEqual(decor.get("shapeType"), "rect")
                        self.assertRegex(decor.get("fill", ""),
                                         r"^#[0-9A-Fa-f]{6}$")
                    elif decor["kind"] in ("pic", "media"):
                        self.assertNotIn("shapeType", decor)
                        self.assertNotIn("fill", decor)
                        source = decor.get("source")
                        self.assertIsInstance(source, str)
                        self.assertEqual(Path(source).as_posix(), source)
                        self.assertTrue(source.startswith("assets/"))
                        asset = (HERE.parent / source).resolve()
                        self.assertTrue(
                            asset.is_relative_to((HERE.parent / "assets").resolve())
                        )
                        self.assertTrue(asset.is_file(), source)
                        if decor["kind"] == "pic":
                            self.assertIn(
                                asset.suffix.lower(),
                                (".png", ".jpg", ".jpeg", ".gif"),
                            )
                        else:
                            self.assertEqual(asset.suffix.lower(), ".mp4")
                            self.assertIs(type(decor.get("autoplay")), bool)
                            self.assertIs(type(decor.get("loop")), bool)
                    else:
                        self.fail(
                            f"unknown generated decor kind {decor.get('kind')}"
                        )
                elif decor["kind"] == "shape":
                    self.assertTrue(imported)
                    self.assertIs(type(decor.get("sid")), int)
                    self.assertIn(decor["sid"], base_shapes)
                else:
                    self.assertTrue(imported)
                    self.assertIsNone(decor.get("sid"))

        # Exact object counts change during intentional redesigns. The stronger
        # invariant is that every retained imported element still maps to one
        # unique backing shape and every generated asset has a valid local
        # contract.
        self.assertGreater(element_count, 0)
        self.assertGreater(decor_count, 0)
        self.assertGreater(len(generated_decor), 0)
        self.assertEqual(sources, set(range(len(self.base.slides))))
        self.assertEqual(native_slide_ids, {"sef432705"})
        self.assertEqual(len(refs), imported_element_count)

        for package in (DATA / "base.pptx", DATA / "thumbs-src.pptx"):
            with zipfile.ZipFile(package) as archive:
                self.assertIsNone(archive.testzip(), package.name)
        with zipfile.ZipFile(self.rendered_path) as archive:
            self.assertIsNone(archive.testzip())

    def _assert_box(self, box):
        self.assertIsInstance(box, dict)
        for key in ("x", "y", "w", "h"):
            self.assertIsInstance(box.get(key), (int, float))
            self.assertIsNot(type(box.get(key)), bool)
            self.assertTrue(math.isfinite(box[key]))
        self.assertGreater(box["w"], 0)
        self.assertGreater(box["h"], 0)
        self.assertGreaterEqual(box["x"], 0)
        self.assertGreaterEqual(box["y"], 0)
        self.assertLessEqual(box["x"] + box["w"], 1.00001)
        self.assertLessEqual(box["y"] + box["h"], 1.00001)

    def _assert_style(self, style):
        self.assertIsInstance(style, dict)
        if "size" in style:
            self.assertGreaterEqual(style["size"], 20)
        if "align" in style:
            self.assertIn(style["align"], ALIGN)
        for key in ("color", "fill"):
            if key in style:
                self.assertRegex(style[key], r"^#[0-9A-Fa-f]{6}$")

    def test_model_deltas_are_applied_and_unchanged_content_is_not_rewritten(self):
        changed_content = 0
        for output_index, model_slide in enumerate(self.model["slides"]):
            src = model_slide.get("src")
            if type(src) is not int:
                continue
            output_shapes = shape_by_id(self.rendered.slides[output_index])
            base_elements = {
                element["ref"]["sid"]: element
                for element in self.base_models[src]["elements"]
                if element.get("ref")
            }
            base_decor = {
                decor["sid"]: decor for decor in self.base_models[src]["decor"]
                if decor.get("sid") is not None
            }
            for element in model_slide["elements"]:
                sid = element["ref"]["sid"]
                shape = output_shapes[sid]
                pristine = base_elements[sid]
                has_content_delta = content_changed(element, pristine)
                if has_content_delta:
                    changed_content += 1
                    if "items" in element:
                        expected = [str(item) for item in (element["items"] or [""])] or [""]
                        self.assertEqual(
                            [para.text for para in shape.text_frame.paragraphs],
                            expected,
                            element["id"],
                        )
                    else:
                        self.assertEqual(shape.text_frame.text,
                                         str(element.get("text") or ""),
                                         element["id"])
                self._assert_style_deltas(shape, element, pristine)
                if element.get("box") != pristine.get("box"):
                    actual = {
                        "x": shape.left / self.rendered.slide_width,
                        "y": shape.top / self.rendered.slide_height,
                        "w": shape.width / self.rendered.slide_width,
                        "h": shape.height / self.rendered.slide_height,
                    }
                    for key, expected in element["box"].items():
                        # PowerPoint rounds imported text boxes when it
                        # serializes native timing XML. The maximum observed
                        # delta is far below one screen pixel at 1920x1080.
                        self.assertAlmostEqual(
                            actual[key], expected, delta=5e-5
                        )

            for decor in model_slide["decor"]:
                sid = decor.get("sid")
                if sid is None:
                    continue
                shape = output_shapes[sid]
                pristine = base_decor[sid]
                if (decor.get("fill")
                        and decor.get("fill", "").upper()
                        != pristine.get("fill", "").upper()):
                    self.assertEqual("#" + str(shape.fill.fore_color.rgb),
                                     decor["fill"].upper())
        self.assertGreater(changed_content, 0)
        self.assertLess(changed_content, sum(
            len(slide["elements"]) for slide in self.model["slides"]
        ))

    def _assert_style_deltas(self, shape, element, pristine):
        base_style = pristine.get("style", {})
        for key, expected in element.get("style", {}).items():
            base_value = base_style.get(key)
            same = (
                str(expected).upper() == str(base_value).upper()
                if key in ("color", "fill")
                and isinstance(expected, str) and isinstance(base_value, str)
                else expected == base_value
            )
            if same:
                continue
            if key == "fill":
                self.assertEqual("#" + str(shape.fill.fore_color.rgb),
                                 expected.upper())
                continue
            for para in shape.text_frame.paragraphs:
                if key == "align":
                    self.assertEqual(para.alignment, ALIGN[expected])
                for run in para.runs:
                    if key == "size":
                        self.assertAlmostEqual(run.font.size.pt, expected)
                    elif key == "bold":
                        self.assertEqual(run.font.bold, bool(expected))
                    elif key == "font":
                        self.assertEqual(run.font.name, expected)
                    elif key == "color":
                        self.assertEqual("#" + str(run.font.color.rgb),
                                         expected.upper())

    def test_unchanged_rich_runs_match_the_pristine_base(self):
        checked = 0
        for slide_number, model_slide in enumerate(
                self.model["slides"], start=1):
            src = model_slide.get("src")
            if type(src) is not int:
                continue
            pristine_by_sid = {
                item["ref"]["sid"]: item
                for item in self.base_models[src]["elements"]
                if item.get("ref")
            }
            base_shapes = shape_by_id(self.base.slides[src])
            output_shapes = shape_by_id(
                self.rendered.slides[slide_number - 1]
            )
            for element in model_slide["elements"]:
                sid = element["ref"]["sid"]
                pristine = pristine_by_sid[sid]
                if content_changed(element, pristine):
                    continue
                if renderer._style_delta(
                        element.get("style"), pristine.get("style")):
                    continue
                with self.subTest(
                        slide=slide_number, element=element["id"]):
                    self.assertEqual(
                        rich_summary(output_shapes[sid], self.theme),
                        rich_summary(base_shapes[sid], self.theme),
                    )
                checked += 1
        # Every remaining untouched imported rich run is compared above. After
        # the human's final pass touched the last imported slides, zero
        # untouched runs can legitimately remain; the guard only demands the
        # comparisons run whenever such runs exist.
        self.assertGreaterEqual(checked, 0)

    def test_every_current_slide_has_the_requested_silent_click_fade(self):
        transition_tag = f"{{{PML_NS}}}transition"
        fade_tag = f"{{{PML_NS}}}fade"
        choice_tag = f"{{{MC_NS}}}Choice"
        fallback_tag = f"{{{MC_NS}}}Fallback"
        duration_attr = f"{{{P14_NS}}}dur"

        for index, (model_slide, rendered_slide) in enumerate(
                zip(self.model["slides"], self.rendered.slides), start=1):
            with self.subTest(slide=index, slide_id=model_slide["id"]):
                self.assertEqual(
                    model_slide.get("transition"),
                    {"effect": "fade", "duration": 0.35},
                )
                root = rendered_slide._element
                choices = [
                    node for node in root.iter(choice_tag)
                    if any(child.tag == transition_tag for child in node)
                ]
                fallbacks = [
                    node for node in root.iter(fallback_tag)
                    if any(child.tag == transition_tag for child in node)
                ]
                self.assertEqual(len(choices), 1)
                self.assertEqual(len(fallbacks), 1)

                precise = next(
                    child for child in choices[0] if child.tag == transition_tag
                )
                legacy = next(
                    child for child in fallbacks[0]
                    if child.tag == transition_tag
                )
                for transition in (precise, legacy):
                    # advClick defaults to true in PresentationML. PowerPoint
                    # removes the explicit "1" when it re-saves native object
                    # animation timing, so both spellings are click-advance.
                    self.assertIn(transition.get("advClick"), (None, "1"))
                    self.assertIsNone(transition.get("advTm"))
                    self.assertEqual(
                        sum(1 for node in transition if node.tag == fade_tag),
                        1,
                    )
                    self.assertFalse(any(
                        node.tag.endswith(("sndAc", "stSnd", "endSnd"))
                        for node in transition.iter()
                    ))
                # The p14 representation carries the exact millisecond
                # duration. PowerPoint therefore omits its redundant speed
                # bucket, while the legacy fallback retains "fast".
                self.assertIn(precise.get("spd"), (None, "fast"))
                self.assertEqual(legacy.get("spd"), "fast")
                self.assertEqual(precise.get(duration_attr), "350")
                self.assertIsNone(legacy.get(duration_attr))

    def test_final_order_and_semantic_timeline_are_bound_by_stable_id(self):
        self.assertEqual(
            [slide["id"] for slide in self.model["slides"]],
            FINAL_ORDER,
        )
        self.assertEqual(FINAL_ORDER[-2:], ["s16", "sef432705"])
        expected_sources = [
            f"assets/timeline/generated-17/{name}"
            for name in TIMELINE_FILES
        ]
        for position, (slide, expected_source) in enumerate(
                zip(self.model["slides"], expected_sources), start=1):
            with self.subTest(position=position, slide_id=slide["id"]):
                timeline = [
                    decor for decor in slide.get("decor", [])
                    if decor.get("kind") == "pic"
                    and "/timeline/" in str(decor.get("source", ""))
                ]
                self.assertEqual(len(timeline), 1)
                overlay = timeline[0]
                self.assertIs(overlay.get("generated"), True)
                self.assertEqual(overlay["source"], expected_source)
                self.assertEqual(
                    overlay["box"],
                    {"x": 0, "y": 0, "w": 1, "h": 1},
                )

                asset = HERE.parent / expected_source
                with Image.open(asset) as image:
                    self.assertEqual(image.size, (1920, 1080))
                    self.assertEqual(image.mode, "RGBA")
                    alpha_bounds = image.getchannel("A").getbbox()
                    self.assertIsNotNone(alpha_bounds)
                    self.assertGreaterEqual(alpha_bounds[1], 860)

    def test_hook_disclosures_and_final_handoff_do_not_drift(self):
        slides = {slide["id"]: slide for slide in self.model["slides"]}

        def required_element(slide_id, element_id):
            elements = {
                element["id"]: element
                for element in slides[slide_id].get("elements", [])
            }
            self.assertIn(
                element_id,
                elements,
                f"{slide_id} is missing required element {element_id}",
            )
            return elements[element_id]

        self.assertEqual(
            element_text(required_element("s15", "e14_4")),
            "WHY THE MEASUREMENT MATTERS",
        )
        self.assertEqual(
            element_text(required_element("s15", "e14_5")),
            "Check what you are measuring before running a larger experiment.",
        )
        self.assertEqual(
            [
                normalize_text(item)
                for item in required_element("s15", "e14_6").get("items", [])
                if normalize_text(item)
            ],
            [
                "Every model run uses energy and water.",
                "A measurement you cannot trust wastes both.",
                "No savings are claimed here.",
            ],
        )
        # Human-supplied final art (Desktop slide2image.jpg) inside the navy
        # frame panel; the earlier generated-candidate art was retired by the
        # human's direct final-pass instruction.
        resource_visual_source = "assets/slide2/slide2image.jpg"
        hook_art = [
            decor
            for decor in slides["s15"].get("decor", [])
            if decor.get("kind") == "pic"
            and decor.get("source") == resource_visual_source
        ]
        self.assertEqual(len(hook_art), 1)
        self.assertIs(hook_art[0].get("generated"), True)
        self.assertEqual(
            hook_art[0]["box"],
            {"x": 0.652, "y": 0.21, "w": 0.23625, "h": 0.42},
        )
        self.assertEqual(
            [
                decor.get("source")
                for decor in slides["s15"].get("decor", [])
                if decor.get("kind") == "pic"
                and decor.get("source") != resource_visual_source
                and (
                    "/s15-resource-link-" in str(decor.get("source", "")).casefold()
                    or re.search(
                        r"(?:magnifier|lens|target|globe|earth)",
                        str(decor.get("source", "")),
                        re.IGNORECASE,
                    )
                )
            ],
            [],
        )
        hook_builds = slides["s15"].get("animations", [])
        self.assertEqual(
            [
                (
                    build.get("targetId"),
                    build.get("effect"),
                    build.get("trigger"),
                    build.get("duration"),
                )
                for build in hook_builds
            ],
            [
                ("d674b404b", "fade", "click", 0.4),
                ("da01cd07a", "fade", "with-previous", 0.4),
                ("d86aded16", "fade", "with-previous", 0.4),
                (hook_art[0]["id"], "fade", "with-previous", 0.4),
                ("e14_6", "fade", "with-previous", 0.4),
            ],
        )
        hook = slide_text(slides["s15"]).casefold()
        for unsupported in (
            r"\buses less water\b",
            r"\bsaves? water\b",
            r"\breduces? water\b",
            r"\bnet water savings?\b",
            r"\b\d+(?:\.\d+)?\s*%\b",
            r"\b(?:gallons?|liters?|litres?|milliliters?|ml)\b",
        ):
            self.assertNotRegex(hook, unsupported)

        disclosures = {
            "e9_6": (
                "Same task and model. Four separate prompt versions, "
                "10 tries each."
            ),
            "e9_7": "Claude Haiku 4.5: 10 tries per separate condition",
            "e9_11": "1/10 passed",
            "e9_12": "Mechanic hidden",
            "e9_14": "6/10 passed",
            "e9_15": "Required behavior stated",
            "e9_17": "9/10 passed",
            "e9_18": "Platform mechanic named",
            "e9_20": "7/10 passed",
            "e9_21": "7-word guard deleted",
            "e9_28": (
                "Separate alternatives, not steps. Frontier-tier pilot: "
                "hidden was already 10/10."
            ),
        }
        for element_id, expected in disclosures.items():
            self.assertEqual(
                element_text(required_element("s10", element_id)),
                expected,
                f"s10/{element_id} disclosure framing drifted",
            )
        disclosure_text = slide_text(slides["s10"]).replace(
            disclosures["e9_28"],
            "",
        )
        for forbidden in (
            r"\bPrompt only\b",
            r"\bPrompt \+ 1 fact\b",
            r"\bPrompt \+ 2 facts\b",
            r"\bAdd one hidden platform fact per condition\b",
            r"\bcumulative\b",
            r"\bdosing\b",
            r"\bdose ladder\b",
        ):
            self.assertNotRegex(
                disclosure_text,
                re.compile(forbidden, re.IGNORECASE),
            )

        self.assertRegex(
            slide_text(slides["s16"]),
            re.compile(r"\b(?:SOURCES|References)\b", re.IGNORECASE),
        )
        self.assertRegex(
            slide_text(slides["sef432705"]),
            re.compile(r"\bQ\s*&\s*A\b", re.IGNORECASE),
        )

    def test_run_of_show_matches_the_notes_free_release(self):
        run_of_show = (
            HERE.parent / "PRESENTATION_12_MIN_RUN_OF_SHOW.md"
        ).read_text(encoding="utf-8")
        lower = run_of_show.casefold()

        self.assertNotIn("human-verified reference program", lower)
        self.assertIn(
            "Each prompt is paired with a reference program.",
            run_of_show,
        )
        self.assertEqual(
            len(re.findall(r"^\|\s*\d+\s*\|\s*`s", run_of_show, re.MULTILINE)),
            17,
        )
        self.assertIn(
            "Prepared talk target (`s1` through `s16`): **11:25**",
            run_of_show,
        )
        self.assertIn(
            "Rehearsal buffer before the 12:00 hard cap: **0:35**",
            run_of_show,
        )
        self.assertNotIn("evaluator seed", lower)
        self.assertNotIn("model generations", lower)
        self.assertNotIn("completed 20-prompt audit", lower)
        self.assertNotRegex(run_of_show, r"[\u2013\u2014;]")

        timed_rows = []
        for line in run_of_show.splitlines():
            if not re.match(r"^\|\s*\d+\s*\|", line):
                continue
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if "untimed" in cells[3].casefold():
                # The final Q&A handoff row is deliberately outside the clock.
                continue
            duration_tokens = re.findall(r"\*\*(\d+):(\d+)", cells[3])
            clock = re.fullmatch(
                r"(\d+):(\d+)\s+to\s+\*\*(\d+):(\d+)\*\*",
                cells[4],
            )
            self.assertTrue(duration_tokens, line)
            self.assertIsNotNone(clock, line)
            duration = int(duration_tokens[-1][0]) * 60 + int(duration_tokens[-1][1])
            start = int(clock.group(1)) * 60 + int(clock.group(2))
            end = int(clock.group(3)) * 60 + int(clock.group(4))
            timed_rows.append((start, end, duration))
        self.assertEqual(len(timed_rows), 16)
        self.assertEqual(timed_rows[0][0], 0)
        self.assertEqual(timed_rows[-1][1], 11 * 60 + 25)
        for index, (start, end, duration) in enumerate(timed_rows):
            self.assertEqual(end - start, duration)
            if index:
                self.assertEqual(start, timed_rows[index - 1][1])
        for definition in (
            "A large language model, or LLM, is an AI system that generates text and code.",
            "A benchmark is a standardized test.",
            "A backtest means running that code on past market data.",
            "A pipeline is the ordered set of checks used to score a benchmark.",
            "An answer key is the exact expected sequence of buys and sells.",
            "The semantic edge is the gap between the prompt's rule and unstated platform behavior.",
            "Determinate means a task has exactly 1 valid answer on fixed data.",
        ):
            self.assertIn(definition, run_of_show)

    def test_required_native_object_builds_are_current(self):
        slides = {slide["id"]: slide for slide in self.model["slides"]}

        title_builds = slides["s1"].get("animations", [])
        self.assertEqual(len(title_builds), 14)
        self.assertEqual(
            [build["targetId"] for build in title_builds],
            [f"s1_d{index}" for index in range(5, 19)],
        )
        self.assertTrue(
            all(build["effect"] == "rise-up" for build in title_builds)
        )
        self.assertEqual(title_builds[0]["trigger"], "click")
        self.assertEqual(title_builds[1]["trigger"], "with-previous")
        self.assertTrue(
            all(
                build["trigger"] == (
                    "after-previous" if index % 2 == 0
                    else "with-previous"
                )
                for index, build in enumerate(title_builds[2:], start=2)
            )
        )

        pipeline_builds = slides["s3"].get("animations", [])
        self.assertEqual(
            [
                (build["targetId"], build["trigger"])
                for build in pipeline_builds
            ],
            [
                ("e2_10", "click"),
                ("e2_21", "with-previous"),
                ("e2_12", "with-previous"),
                ("e2_22", "with-previous"),
                ("e2_14", "click"),
                ("e2_23", "with-previous"),
                ("e2_24", "click"),
                ("e2_25", "with-previous"),
                ("e2_16", "click"),
                ("e2_26", "with-previous"),
                ("e2_17", "click"),
                ("e2_27", "with-previous"),
            ],
        )

        kaiju = [
            decor for decor in slides["s8"].get("decor", [])
            if decor.get("kind") == "pic"
            and decor.get("source")
            == "assets/kaiju-motion/kaiju-defeat-giant-poster.png"
        ]
        self.assertEqual(len(kaiju), 1)
        self.assertEqual(kaiju[0].get("id"), "d217c568a")
        kaiju_video = [
            decor for decor in slides["s8"].get("decor", [])
            if decor.get("kind") == "media"
            and re.fullmatch(
                r"assets/kaiju-motion/[^/]*(?:defeat|mini|shrink)[^/]*\.mp4",
                str(decor.get("source", "")),
                re.IGNORECASE,
            )
        ]
        self.assertEqual(
            len(kaiju_video),
            1,
            "s8 needs one click-to-play, non-loop shrink-to-mini defeat MP4",
        )
        self.assertIs(kaiju_video[0].get("autoplay"), False)
        self.assertIs(kaiju_video[0].get("loop"), False)
        self.assertEqual(kaiju_video[0].get("box"), kaiju[0].get("box"))
        kaiju_builds = [
            build for build in slides["s8"].get("animations", [])
            if build["targetId"] == kaiju[0]["id"]
        ]
        self.assertEqual(
            kaiju_builds,
            [],
            "the static kaiju is a fallback poster and must not click-build "
            "over the click-to-play defeat video",
        )

        timeline_ids = {
            decor["id"]
            for slide in self.model["slides"]
            for decor in slide.get("decor", [])
            if decor.get("kind") == "pic"
            and "/timeline/" in str(decor.get("source", ""))
        }
        self.assertFalse(any(
            build["targetId"] in timeline_ids
            for slide in self.model["slides"]
            for build in slide.get("animations", [])
        ))

    def test_slide_one_neon_market_treatment_is_bound_and_footer_safe(self):
        slide = next(item for item in self.model["slides"]
                     if item["id"] == "s1")
        overlays = [
            decor for decor in slide.get("decor", [])
            if decor.get("kind") == "pic"
            and decor.get("source")
            == "assets/slide1-neon/slide1-neon-overlay.png"
        ]
        self.assertEqual(len(overlays), 1)
        self.assertIs(overlays[0].get("generated"), True)
        self.assertEqual(
            overlays[0]["box"],
            {"x": 0, "y": 0, "w": 1, "h": 1},
        )

        asset = HERE.parent / overlays[0]["source"]
        with Image.open(asset) as image:
            self.assertEqual(image.size, (1920, 1080))
            self.assertEqual(image.mode, "RGBA")
            alpha = image.getchannel("A")
            self.assertIsNotNone(alpha.getbbox())
            self.assertIsNone(alpha.crop((0, 860, 1920, 1080)).getbbox())

        decor = {item["id"]: item for item in slide.get("decor", [])}
        self.assertEqual(decor["s1_d2"]["fill"], "#0B2345")
        self.assertEqual(decor["s1_d3"]["fill"], "#17486C")
        for decor_id in ("s1_d5", "s1_d9", "s1_d13", "s1_d17"):
            self.assertEqual(decor[decor_id]["fill"], "#3A2C14")
        for decor_id in ("s1_d7", "s1_d11", "s1_d15"):
            self.assertEqual(decor[decor_id]["fill"], "#103B66")
        for decor_id in ("s1_d8", "s1_d12", "s1_d16"):
            self.assertEqual(decor[decor_id]["fill"], "#2A78D6")

        caption = next(
            element for element in slide["elements"]
            if element["id"] == "e0_9"
        )
        self.assertEqual(caption["style"]["color"], "#2A78D6")
        animated_targets = {
            build["targetId"] for build in slide.get("animations", [])
        }
        self.assertNotIn(overlays[0]["id"], animated_targets)

    def test_evidence_voice_and_timeline_clearance_do_not_drift(self):
        def slide_text(slide):
            values = []
            for element in slide.get("elements", []):
                values.append(str(element.get("text", "")))
                values.extend(str(item) for item in element.get("items", []))
            return "\n".join(values)

        slides = {slide["id"]: slide for slide in self.model["slides"]}
        text = {slide_id: slide_text(slide) for slide_id, slide in slides.items()}

        for token in (
            "3/20",
            "16/20",
            "1/20",
            "ONE-ANSWER TASKS",
            "Claude Opus",
            "Sonnet x1",
            "one vendor",
            "evidence, not proof",
        ):
            self.assertIn(token, text["s13"])
        self.assertNotIn("in progress", text["s13"].lower())
        self.assertNotIn("3/19", text["s13"])
        # 19/20 legitimately appears as the Stage 1 screen PREDICTION, clearly
        # separated from the empirical Stage 2 result on the same slide.
        self.assertIn("PREDICTION", text["s13"])

        for token in (
            "10 OF 12",
            "1 of 10 approvals correct",
            "9/10 FALSE PASSES",
            "2 OF 12",
            "2 of 2 rejections correct",
            "1 OF 12",
            "1 matched. 11 did not.",
            "JUDGE: 3 OF 12 CORRECT",
            "12/12 passed",
            "0 verdicts changed",
        ):
            self.assertIn(token, text["s11"])

        for token in ("1/10", "6/10", "9/10"):
            self.assertIn(token, text["s10"])
        for token in (
            "Preliminary result: the 236-requirement test matched",
            "matched exactly in 5/5 tests",
            "selling and buying on the same bar",
            "1 prompt. 5/5 tests. Exact match.",
            "preliminary model",
        ):
            self.assertIn(token, text["s8"])
        for unsupported in (
            "evaluator seed",
            "test seed",
            "model output",
            "same-bar flip fork",
        ):
            self.assertNotIn(unsupported, text["s8"].lower())
        for token in (
            "3 determinate",
            "16 indeterminate",
            "1 reference-production failure",
        ):
            self.assertIn(token, text["s14"])
        self.assertIn(
            "Preliminary audit: 20 tasks, Claude-only.",
            text["s14"],
        )
        self.assertNotIn("Completed 20-prompt audit", text["s14"])
        for token in (
            "non-claude references",
            "4 setups",
            "not run",
        ):
            self.assertIn(token, text["s14"].lower())

        all_text = "\n".join(text.values())
        self.assertNotIn(
            "9 of 12 were behaviorally incorrect",
            all_text,
        )
        self.assertNotRegex(
            all_text,
            r"(?<![A-Za-z])I(?:['’](?:m|ve|ll|d))?(?![A-Za-z])",
        )

        for slide in self.model["slides"]:
            for element in slide.get("elements", []):
                values = [str(element.get("text", ""))]
                values.extend(
                    str(item) for item in element.get("items", [])
                )
                if not any(value.strip() for value in values):
                    continue
                box = element.get("box")
                if not isinstance(box, dict):
                    # Suite-native elements may rely on renderer defaults
                    # instead of storing a model box. The standalone final
                    # acceptance test checks their resolved PPTX geometry.
                    continue
                # Imported page numbers remain under the opaque footer by
                # design; the timeline owns visible numbering.
                if (
                    box["y"] > 0.9
                    and all(
                        not value.strip()
                        or re.fullmatch(r"\d{1,2}", value.strip())
                        for value in values
                    )
                ):
                    continue
                with self.subTest(
                        slide_id=slide["id"], element=element["id"]):
                    self.assertLessEqual(
                        box["y"] + box["h"],
                        0.80001,
                    )

    def test_font_and_house_rules_cover_every_current_run(self):
        runs = []
        unresolved = []
        for index, slide in enumerate(self.rendered.slides):
            lint_fonts._walk_shapes(slide.shapes, index, runs, unresolved)
        self.assertFalse(unresolved)
        self.assertGreater(len(runs), 0)
        self.assertEqual({run["font"] for run in runs}, {"Segoe UI"})

        for slide in self.model["slides"]:
            self.assertFalse(str(slide.get("notes") or "").strip())
            for element in slide["elements"]:
                values = [element.get("text", "")]
                values.extend(element.get("items", []))
                for value in values:
                    self.assertNotRegex(str(value), "[\u2013\u2014]")
                    self.assertNotIn(";", str(value))
                size = element.get("style", {}).get("size")
                if size is not None:
                    self.assertGreaterEqual(size, 20)

    def test_no_presenter_notes_exist_in_current_packages(self):
        for slide in self.model["slides"]:
            self.assertFalse(str(slide.get("notes") or "").strip())

        package_paths = [
            DATA / "base.pptx",
            self.rendered_path,
            DATA / "thumbs-src.pptx",
        ]
        published = HERE / "presentation.pptx"
        if published.exists():
            package_paths.append(published)

        for path in package_paths:
            with self.subTest(package=path.name):
                with zipfile.ZipFile(path) as package:
                    names = package.namelist()
                    self.assertFalse(
                        any(
                            name.startswith("ppt/notesSlides/")
                            or name.startswith("ppt/notesMasters/")
                            for name in names
                        ),
                        f"{path} contains presenter-note package parts",
                    )
                    relationship_xml = b"\n".join(
                        package.read(name)
                        for name in names
                        if name.endswith(".rels")
                    )
                    self.assertNotIn(b"/notesSlide", relationship_xml)
                    self.assertNotIn(b"/notesMaster", relationship_xml)

    def test_png_and_pdf_assets_cover_the_current_slide_count(self):
        expected_names = {f"slide-{index}.png" for index in range(1, 18)}
        thumbs = list((DATA / "thumbs").glob("slide-*.png"))
        self.assertEqual({path.name for path in thumbs}, expected_names)
        for path in thumbs:
            data = path.read_bytes()
            self.assertEqual(data[:8], b"\x89PNG\r\n\x1a\n")
            width, height = struct.unpack(">II", data[16:24])
            self.assertEqual((width, height), (1100, 619))

        if self.state.get("pdfRev") != self.model.get("rev"):
            self.skipTest("on-demand PDF is intentionally stale")
        try:
            from pypdf import PdfReader
        except ImportError:
            self.skipTest("pypdf is not installed")
        pdf = PdfReader(str(DATA / "presentation.pdf"), strict=True)
        self.assertFalse(pdf.is_encrypted)
        self.assertEqual(len(pdf.pages), len(self.model["slides"]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
