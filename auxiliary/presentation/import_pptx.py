#!/usr/bin/env python
"""Import an existing .pptx into the suite as an editable OVERLAY.

    python import_pptx.py <source.pptx>

This does NOT flatten the deck. It:
  1. copies the source to data/base.pptx  (a pristine base — never edited)
  2. copies the source to ../presentation.pptx  (the initial working copy)
  3. writes data/model.json in "overlay" mode: every editable text shape becomes
     an element carrying a ref {s, sid} back to the real shape (so edits are
     applied in place on the base — images / tables / layouts all survive), plus
     the shape's real geometry + text style so the dashboard can render a true
     thumbnail. Non-text shapes with a solid colour become read-only "decor" so
     the watch view resembles the actual slide.
"""
import json
import os
import shutil
import stat
import sys
import tempfile

from lxml import etree

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL, MSO_COLOR_TYPE, MSO_THEME_COLOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
BASE = os.path.join(DATA, "base.pptx")
MODEL = os.path.join(DATA, "model.json")
WORKING = os.path.join(os.path.dirname(HERE), "presentation.pptx")

ALIGN = {1: "left", 2: "center", 3: "right", 4: "justify"}
BOX_SHAPE_PRESETS = frozenset({
    "rect",
    "roundRect",
    "round1Rect",
    "round2SameRect",
    "round2DiagRect",
    "snip1Rect",
    "snip2SameRect",
    "snip2DiagRect",
    "snipRoundRect",
})
_DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _clear_readonly(path):
    if os.path.exists(path):
        os.chmod(path, os.stat(path).st_mode | stat.S_IWRITE)


# map MSO_THEME_COLOR members -> a:clrScheme element names
_THEME_KEY = {
    MSO_THEME_COLOR.DARK_1: "dk1", MSO_THEME_COLOR.LIGHT_1: "lt1",
    MSO_THEME_COLOR.DARK_2: "dk2", MSO_THEME_COLOR.LIGHT_2: "lt2",
    MSO_THEME_COLOR.TEXT_1: "dk1", MSO_THEME_COLOR.BACKGROUND_1: "lt1",
    MSO_THEME_COLOR.TEXT_2: "dk2", MSO_THEME_COLOR.BACKGROUND_2: "lt2",
    MSO_THEME_COLOR.ACCENT_1: "accent1", MSO_THEME_COLOR.ACCENT_2: "accent2",
    MSO_THEME_COLOR.ACCENT_3: "accent3", MSO_THEME_COLOR.ACCENT_4: "accent4",
    MSO_THEME_COLOR.ACCENT_5: "accent5", MSO_THEME_COLOR.ACCENT_6: "accent6",
    MSO_THEME_COLOR.HYPERLINK: "hlink", MSO_THEME_COLOR.FOLLOWED_HYPERLINK: "folHlink",
}
_A_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def load_theme_colors(prs):
    """Return {'dk1': '#RRGGBB', ...} from the first slide master's theme."""
    colors = {}
    try:
        theme_part = prs.slide_masters[0].part.part_related_by(RT.THEME)
        root = etree.fromstring(theme_part.blob)
        scheme = root.find(".//a:clrScheme", _A_NS)
        for child in scheme:
            tag = etree.QName(child).localname
            srgb = child.find("a:srgbClr", _A_NS)
            sysc = child.find("a:sysClr", _A_NS)
            if srgb is not None:
                colors[tag] = "#" + srgb.get("val").upper()
            elif sysc is not None:
                colors[tag] = "#" + sysc.get("lastClr", "000000").upper()
    except Exception:
        pass
    return colors


def _brighten(hexv, b):
    """Apply a -1..1 brightness (shade/tint) to a #RRGGBB colour."""
    try:
        r, g, bl = int(hexv[1:3], 16), int(hexv[3:5], 16), int(hexv[5:7], 16)
        if b > 0:
            r, g, bl = (int(c + (255 - c) * b) for c in (r, g, bl))
        else:
            r, g, bl = (int(c * (1 + b)) for c in (r, g, bl))
        return "#%02X%02X%02X" % (r, g, bl)
    except Exception:
        return hexv


def _hex(color, theme=None):
    """Resolve a ColorFormat to #RRGGBB, following theme colours + brightness."""
    try:
        ctype = color.type
    except Exception:
        return None
    hexv = None
    try:
        if ctype == MSO_COLOR_TYPE.SCHEME and theme:
            hexv = theme.get(_THEME_KEY.get(color.theme_color))
        elif ctype is not None:
            hexv = "#" + str(color.rgb)   # RGBColor is a tuple subclass; str() gives 'RRGGBB'
    except Exception:
        return None
    if hexv is None:
        return None
    try:
        b = color.brightness
    except Exception:
        b = 0
    return _brighten(hexv, b) if b else hexv


def geom(shape, W, H):
    try:
        if None in (shape.left, shape.top, shape.width, shape.height):
            return None
        return {"x": round(shape.left / W, 5), "y": round(shape.top / H, 5),
                "w": round(shape.width / W, 5), "h": round(shape.height / H, 5)}
    except Exception:
        return None


def text_style(shape, theme=None):
    tf = shape.text_frame
    size = None; bold = None; color = None
    for para in tf.paragraphs:
        for run in para.runs:
            if run.font.size is not None:
                size = max(size or 0, run.font.size.pt)
            if bold is None and run.font.bold is not None:
                bold = run.font.bold
            if color is None:
                c = _hex(run.font.color, theme)
                if c:
                    color = c
    p0 = tf.paragraphs[0]
    align = ALIGN.get(int(p0.alignment)) if p0.alignment is not None else None
    st = {}
    if size:
        st["size"] = round(size, 1)
    if bold:
        st["bold"] = True
    if color:
        st["color"] = color
    if align:
        st["align"] = align
    return st


def shape_fill_hex(shape, theme=None):
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return _hex(shape.fill.fore_color, theme)
    except Exception:
        pass
    return None


def box_shape_preset(shape):
    """Return the OOXML preset for an eligible imported box/card AutoShape.

    A preset alone is not enough: PowerPoint text boxes and placeholders can
    also carry rect geometry, but converting those would change layout
    semantics rather than merely remove a card's rounded corners. Keep the
    import metadata deliberately narrow so every later caller can fail closed.
    """
    try:
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            return None
        if shape.is_placeholder:
            return None
        preset_geometry = shape._element.find(
            f".//{{{_DRAWINGML_NS}}}prstGeom"
        )
        if preset_geometry is None:
            return None
        preset = preset_geometry.get("prst")
        return preset if preset in BOX_SHAPE_PRESETS else None
    except Exception:
        return None


def derive_slide(prs, si, theme=None):
    """Derive one overlay-model slide using the canonical import rules.

    render_pptx.py and tools_decor_sids.py deliberately call this helper too.
    Keeping one classifier prevents legacy migrations from drifting away from
    the importer (theme-color fills previously exposed exactly that bug).
    """
    slide = prs.slides[si]
    W, H = prs.slide_width, prs.slide_height
    if theme is None:
        theme = load_theme_colors(prs)
    try:
        title_shape = slide.shapes.title
    except Exception:
        title_shape = None
    title_sid = title_shape.shape_id if title_shape is not None else None

    elements = []
    decor = []
    for shape in slide.shapes:
        sid = shape.shape_id
        try:
            stype = shape.shape_type
        except Exception:
            stype = None
        box = geom(shape, W, H)
        has_text = getattr(shape, "has_text_frame", False) and shape.has_text_frame
        text_content = shape.text_frame.text.strip() if has_text else ""
        source_preset = box_shape_preset(shape)

        if stype == MSO_SHAPE_TYPE.PICTURE:
            if box:
                decor.append({"id": f"s{si + 1}_d{len(decor)}", "kind": "pic",
                              "sid": None, "box": box})
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            tbl = shape.table
            elements.append({
                "id": f"e{si}_{sid}", "type": "table", "readonly": True,
                "text": f"\u25A6 table {len(tbl.rows)}\u00D7{len(tbl.columns)}",
                "box": box,
            })
            continue
        if has_text and text_content:
            nonempty = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            style = text_style(shape, theme)
            fill = shape_fill_hex(shape, theme)
            if fill:
                style["fill"] = fill
            element = {
                "id": f"e{si}_{sid}",
                "ref": {"s": si, "sid": sid},
                "box": box,
                "style": style,
            }
            if source_preset:
                element["sourcePreset"] = source_preset
            if title_sid is not None and sid == title_sid:
                element.update({"type": "title", "text": " ".join(nonempty)})
            elif len(nonempty) > 1:
                element.update({"type": "bullets", "items": nonempty})
            else:
                element.update({"type": "body", "text": nonempty[0]})
            elements.append(element)
            continue

        fill = shape_fill_hex(shape, theme)
        if fill and box:
            decor_item = {
                "id": f"s{si + 1}_d{len(decor)}", "kind": "shape",
                "sid": sid, "box": box, "fill": fill,
            }
            if source_preset:
                decor_item["sourcePreset"] = source_preset
            decor.append(decor_item)

    elements.sort(key=lambda e: 0 if e.get("type") == "title" else 1)
    notes = ""
    try:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text or ""
    except Exception:
        notes = ""
    try:
        layout_name = slide.slide_layout.name
    except Exception:
        layout_name = "imported"
    return {
        "id": f"s{si + 1}", "src": si, "layout": layout_name,
        "notes": notes, "decor": decor, "elements": elements,
    }


def derive_model(prs, src, base_path=BASE):
    """Build a complete overlay model without touching destination files."""
    W, H = prs.slide_width, prs.slide_height
    theme = load_theme_colors(prs)
    return {
        "title": os.path.splitext(os.path.basename(src))[0],
        "mode": "overlay",
        "base": os.path.abspath(base_path),
        "slideW": round(W / 914400, 3),
        "slideH": round(H / 914400, 3),
        "theme": {"accent": "#2563eb"},
        "shapePresetSchema": 1,
        "rev": 0,
        "updatedAt": None,
        "slides": [derive_slide(prs, si, theme) for si in range(len(prs.slides))],
    }


def _stage_copy(src, destination):
    os.makedirs(os.path.dirname(destination), exist_ok=True)
    fd, staged = tempfile.mkstemp(
        prefix=os.path.basename(destination) + ".import-",
        suffix=".tmp", dir=os.path.dirname(destination),
    )
    os.close(fd)
    try:
        shutil.copyfile(src, staged)
        return staged
    except Exception:
        try:
            os.unlink(staged)
        except OSError:
            pass
        raise


def _stage_model(model, destination):
    os.makedirs(os.path.dirname(destination), exist_ok=True)
    fd, staged = tempfile.mkstemp(
        prefix=os.path.basename(destination) + ".import-",
        suffix=".tmp", dir=os.path.dirname(destination),
    )
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as f:
            json.dump(model, f, indent=2, ensure_ascii=False)
            f.flush()
            os.fsync(f.fileno())
        return staged
    except Exception:
        try:
            os.close(fd)
        except OSError:
            pass
        try:
            os.unlink(staged)
        except OSError:
            pass
        raise


def _commit_staged(staged_pairs):
    """Replace destinations in order and restore every old file on failure."""
    backups = {}
    committed = []
    try:
        for staged, destination in staged_pairs:
            backup = None
            if os.path.exists(destination):
                _clear_readonly(destination)
                fd, backup = tempfile.mkstemp(
                    prefix=os.path.basename(destination) + ".backup-",
                    suffix=".tmp", dir=os.path.dirname(destination),
                )
                os.close(fd)
                os.unlink(backup)
                os.replace(destination, backup)
                backups[destination] = backup
            try:
                os.replace(staged, destination)
                committed.append(destination)
            except Exception:
                if backup and os.path.exists(backup):
                    os.replace(backup, destination)
                    backups.pop(destination, None)
                raise
    except Exception:
        for destination in reversed(committed):
            if os.path.exists(destination):
                _clear_readonly(destination)
                os.unlink(destination)
            backup = backups.pop(destination, None)
            if backup and os.path.exists(backup):
                os.replace(backup, destination)
        raise
    finally:
        for staged, _ in staged_pairs:
            if os.path.exists(staged):
                try:
                    os.unlink(staged)
                except OSError:
                    pass
    for backup in backups.values():
        try:
            os.unlink(backup)
        except OSError:
            pass


def import_pptx(src, base_path=None, model_path=None, working_path=None):
    """Validate, derive, and stage an import before replacing any live file."""
    src = os.path.abspath(src)
    base_path = os.path.abspath(base_path or BASE)
    model_path = os.path.abspath(model_path or MODEL)
    working_path = os.path.abspath(working_path or WORKING)

    # Validate and derive first. The former implementation copied a corrupt
    # source over both decks before Presentation() discovered the bad package.
    prs = Presentation(src)
    model = derive_model(prs, src, base_path)
    if not model["slides"]:
        raise ValueError("cannot import a presentation with zero slides")

    staged_base = staged_working = staged_model = None
    try:
        staged_base = _stage_copy(src, base_path)
        staged_working = _stage_copy(src, working_path)
        staged_model = _stage_model(model, model_path)
        Presentation(staged_base)
        Presentation(staged_working)
        # The model is deliberately last. Exceptions roll every earlier file
        # back, and a sudden process loss never exposes a new model first.
        _commit_staged([
            (staged_base, base_path),
            (staged_working, working_path),
            (staged_model, model_path),
        ])
    finally:
        for staged in (staged_base, staged_working, staged_model):
            if staged and os.path.exists(staged):
                try:
                    os.unlink(staged)
                except OSError:
                    pass

    slides = model["slides"]
    ntext = sum(len([e for e in s["elements"] if not e.get("readonly")]) for s in slides)
    ndecor = sum(len(s["decor"]) for s in slides)
    print(f"imported {len(slides)} slides ({ntext} editable text, {ndecor} decor) from {os.path.basename(src)}")
    print(f"  base    : {base_path}")
    print(f"  model   : {model_path}")
    print(f"  working : {working_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: python import_pptx.py <source.pptx>"); sys.exit(2)
    import_pptx(sys.argv[1])
