#!/usr/bin/env python
"""Render the JSON presentation model into a real .pptx file.

    python render_pptx.py [model.json | -] [output.pptx]

Two modes, chosen by model["mode"]:

  "native"  (default) — build the deck from scratch from the model. Used for new
                        blank decks created inside the suite.

  "overlay"           — open a pristine base .pptx (model["base"]) and apply text
                        edits IN PLACE onto the real shapes referenced by each
                        element's ref {s, sid}. Images, tables, decorative shapes
                        and formatting are preserved. Supports edit, notes,
                        delete, reorder, and appending new native slides.

"-" as the model path reads the model JSON from stdin (the server uses this to
avoid a Windows file-handle race with its own atomic writes to model.json).
"""
import copy
import ctypes
import hashlib
import json
import math
import re
import sys
import os
import zipfile
from ctypes import wintypes
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

from import_pptx import BOX_SHAPE_PRESETS, derive_slide, load_theme_colors

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = Path(HERE).parent.resolve()
PROJECT_ASSETS = (PROJECT_ROOT / "assets").resolve()
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif"}
VIDEO_EXTENSIONS = {".mp4"}

LAYOUT_MAP = {
    "title": 0, "content": 1, "bullets": 1, "section": 2,
    "two": 3, "compare": 4, "titleonly": 5, "blank": 6,
}

PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def _xml_name(namespace, local_name):
    return f"{{{namespace}}}{local_name}"


def _resolve_image_source(source):
    if not isinstance(source, str) or not source:
        raise ValueError("generated picture source must be a non-empty string")
    if os.path.isabs(source) or "\\" in source:
        raise ValueError("generated picture source must be a normalized project-relative path")
    parts = source.split("/")
    if (len(parts) < 2 or parts[0].lower() != "assets"
            or any(part in ("", ".", "..") for part in parts)):
        raise ValueError("generated picture source must be inside the project assets folder")
    candidate = (PROJECT_ROOT / Path(*parts)).resolve(strict=True)
    try:
        candidate.relative_to(PROJECT_ASSETS)
    except ValueError as error:
        raise ValueError(
            "generated picture source must stay inside the project assets folder"
        ) from error
    if not candidate.is_file():
        raise ValueError("generated picture source must be a file")
    if candidate.suffix.lower() not in IMAGE_EXTENSIONS:
        raise ValueError("generated picture source must be a PNG, JPEG, or GIF image")
    return candidate


def _resolve_video_source(source):
    if not isinstance(source, str) or not source:
        raise ValueError("generated media source must be a non-empty string")
    if os.path.isabs(source) or "\\" in source:
        raise ValueError("generated media source must be a normalized project-relative path")
    parts = source.split("/")
    if (len(parts) < 2 or parts[0].lower() != "assets"
            or any(part in ("", ".", "..") for part in parts)):
        raise ValueError("generated media source must be inside the project assets folder")
    candidate = (PROJECT_ROOT / Path(*parts)).resolve(strict=True)
    try:
        candidate.relative_to(PROJECT_ASSETS)
    except ValueError as error:
        raise ValueError(
            "generated media source must stay inside the project assets folder"
        ) from error
    if not candidate.is_file():
        raise ValueError("generated media source must be a file")
    if candidate.suffix.lower() not in VIDEO_EXTENSIONS:
        raise ValueError("generated media source must be an MP4 video")
    return candidate


def _hex_to_rgb(value, default=(37, 99, 235)):
    try:
        value = value.lstrip("#")
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
    except Exception:
        return RGBColor(*default)


def _same_value(key, left, right):
    if key in ("color", "fill", "outline") and isinstance(left, str) and isinstance(right, str):
        return left.upper() == right.upper()
    return left == right


def _style_delta(style, base_style):
    """Return only explicit style values that differ from the pristine base."""
    style = style if isinstance(style, dict) else {}
    base_style = base_style if isinstance(base_style, dict) else {}
    allowed = {"color", "size", "bold", "font", "align", "fill", "outline"}
    return {
        key: value for key, value in style.items()
        if key in allowed and not _same_value(key, value, base_style.get(key))
    }


def _box_differs(box, base_box):
    if not isinstance(box, dict):
        return False
    base_box = base_box if isinstance(base_box, dict) else {}
    return any(base_box.get(key) != value for key, value in box.items()
               if key in ("x", "y", "w", "h"))


def _content_differs(element, base_element):
    """Compare the model abstraction, including text-vs-bullets representation."""
    base_element = base_element if isinstance(base_element, dict) else {}
    if "items" in element:
        return "items" not in base_element or element.get("items") != base_element.get("items")
    if "text" in element:
        return "text" not in base_element or element.get("text") != base_element.get("text")
    return False


def _validate_model(model):
    if not isinstance(model, dict):
        raise ValueError("model must be a JSON object")
    slides = model.get("slides")
    if not isinstance(slides, list):
        raise ValueError("model.slides must be an array")
    if not slides:
        raise ValueError("model must contain at least one slide")
    mode = model.get("mode") or "native"
    if "shapePresetSchema" in model and model["shapePresetSchema"] != 1:
        raise ValueError("shapePresetSchema must be 1 when present")
    seen_ids = set()
    for index, slide in enumerate(slides):
        if not isinstance(slide, dict):
            raise ValueError(f"slide {index} must be an object")
        slide_id = slide.get("id")
        if slide_id is not None:
            if not isinstance(slide_id, str) or not slide_id:
                raise ValueError(f"slide {index} id must be a non-empty string")
            if slide_id in seen_ids:
                raise ValueError(f"duplicate slide id: {slide_id}")
            seen_ids.add(slide_id)
        if not isinstance(slide.get("elements", []), list):
            raise ValueError(f"slide {index} elements must be an array")
        if not isinstance(slide.get("decor", []), list):
            raise ValueError(f"slide {index} decor must be an array")
        for decor_index, decor in enumerate(slide.get("decor", [])):
            label = f"slide {index} decor {decor_index}"
            if not isinstance(decor, dict):
                raise ValueError(f"{label} must be an object")
            if "generated" in decor and decor["generated"] is not True:
                raise ValueError(f"{label}.generated must be true when present")
            if "sourcePreset" in decor:
                if decor["sourcePreset"] not in BOX_SHAPE_PRESETS:
                    raise ValueError(
                        f"{label}.sourcePreset is not an eligible box/card preset"
                    )
                if (
                    mode != "overlay"
                    or type(slide.get("src")) is not int
                    or decor.get("generated") is True
                    or decor.get("kind") != "shape"
                    or type(decor.get("sid")) is not int
                ):
                    raise ValueError(
                        f"{label}.sourcePreset is valid only for an imported AutoShape"
                    )
            if "corners" in decor:
                if decor["corners"] != "sharp":
                    raise ValueError(f"{label}.corners must be sharp")
                if decor.get("sourcePreset") not in BOX_SHAPE_PRESETS:
                    raise ValueError(
                        f"{label}.corners needs eligible imported sourcePreset metadata"
                    )
            if "shapeType" in decor and decor["shapeType"] != "rect":
                raise ValueError(f"{label} has an unknown shapeType")
            if decor.get("generated") is True:
                is_native_slide = slide.get("src") is None
                is_imported_overlay_slide = (
                    mode == "overlay" and type(slide.get("src")) is int
                )
                if not is_native_slide and not is_imported_overlay_slide:
                    raise ValueError(
                        f"{label} is allowed only on a suite-native or "
                        "imported overlay slide"
                    )
                if not isinstance(decor.get("id"), str) or not decor["id"]:
                    raise ValueError(f"{label} id must be a non-empty string")
                if "sid" not in decor or decor["sid"] is not None:
                    raise ValueError(f"{label}.sid must be null")
                box = decor.get("box")
                if not isinstance(box, dict):
                    raise ValueError(f"{label}.box must be an object")
                for key in ("x", "y", "w", "h"):
                    value = box.get(key)
                    if (isinstance(value, bool)
                            or not isinstance(value, (int, float))
                            or not math.isfinite(value)):
                        raise ValueError(
                            f"{label}.box.{key} must be a finite number"
                        )
                    if not 0 <= value <= 1:
                        raise ValueError(
                            f"{label}.box.{key} must be between 0 and 1"
                        )
                    if key in ("w", "h") and value <= 0:
                        raise ValueError(f"{label}.box.{key} must be positive")
                if box["x"] + box["w"] > 1 or box["y"] + box["h"] > 1:
                    raise ValueError(f"{label}.box must fit within the slide")
                if decor.get("kind") == "shape":
                    if decor.get("shapeType") != "rect":
                        raise ValueError(f"{label}.shapeType must be rect")
                    fill = decor.get("fill")
                    if not isinstance(fill, str) or re.fullmatch(
                            r"#[0-9a-fA-F]{6}", fill) is None:
                        raise ValueError(f"{label}.fill must be a #RRGGBB color")
                    if "source" in decor:
                        raise ValueError(f"{label} generated shape cannot have source")
                elif decor.get("kind") == "pic":
                    if "shapeType" in decor:
                        raise ValueError(f"{label} generated picture cannot have shapeType")
                    if "fill" in decor:
                        raise ValueError(f"{label} generated picture cannot have fill")
                    _resolve_image_source(decor.get("source"))
                elif decor.get("kind") == "media":
                    if "shapeType" in decor:
                        raise ValueError(f"{label} generated media cannot have shapeType")
                    if "fill" in decor:
                        raise ValueError(f"{label} generated media cannot have fill")
                    _resolve_video_source(decor.get("source"))
                    if type(decor.get("autoplay")) is not bool:
                        raise ValueError(f"{label}.autoplay must be a boolean")
                    if type(decor.get("loop")) is not bool:
                        raise ValueError(f"{label}.loop must be a boolean")
                else:
                    raise ValueError(f"{label} kind must be shape, pic, or media")
            elif decor.get("kind") == "media":
                raise ValueError(
                    f"{label} media is valid only for generated decor"
                )
            elif "shapeType" in decor:
                raise ValueError(
                    f"{label}.shapeType is valid only for generated decor"
                )
        layout = slide.get("layout", "content")
        if not isinstance(layout, str):
            raise ValueError(f"slide {index} layout must be a string")
        if "transition" in slide:
            transition = slide["transition"]
            if not isinstance(transition, dict):
                raise ValueError(f"slide {index} transition must be an object")
            unknown = set(transition) - {"effect", "duration"}
            if unknown:
                raise ValueError(
                    f"slide {index} transition has unknown field: "
                    f"{sorted(str(field) for field in unknown)[0]}"
                )
            effect = transition.get("effect")
            if effect not in ("fade", "none"):
                raise ValueError(
                    f"slide {index} transition.effect must be fade or none"
                )
            if effect == "none":
                if "duration" in transition:
                    raise ValueError(
                        f"slide {index} none transition cannot have a duration"
                    )
            else:
                duration = transition.get("duration")
                if (isinstance(duration, bool)
                        or not isinstance(duration, (int, float))
                        or not math.isfinite(duration)
                        or not 0.1 <= duration <= 10):
                    raise ValueError(
                        f"slide {index} fade transition duration must be "
                        "from 0.1 through 10 seconds"
                    )
        for element in slide.get("elements", []):
            if not isinstance(element, dict):
                raise ValueError(f"slide {index} contains a non-object element")
            if "items" in element and not isinstance(element["items"], list):
                raise ValueError(f"slide {index} element items must be an array")
            if "style" in element and not isinstance(element["style"], dict):
                raise ValueError(f"slide {index} element style must be an object")
            if "box" in element and element["box"] is not None and not isinstance(element["box"], dict):
                raise ValueError(f"slide {index} element box must be an object")
            if element.get("box") is not None and slide.get("src") is None:
                box = element["box"]
                for key in ("x", "y", "w", "h"):
                    value = box.get(key)
                    if (isinstance(value, bool)
                            or not isinstance(value, (int, float))
                            or not math.isfinite(value)):
                        raise ValueError(
                            f"slide {index} native element box.{key} must be a finite number"
                        )
                    if not 0 <= value <= 1:
                        raise ValueError(
                            f"slide {index} native element box.{key} must be between 0 and 1"
                        )
                    if key in ("w", "h") and value <= 0:
                        raise ValueError(
                            f"slide {index} native element box.{key} must be positive"
                        )
                if box["x"] + box["w"] > 1 or box["y"] + box["h"] > 1:
                    raise ValueError(
                        f"slide {index} native element box must fit within the slide"
                    )
            if "sourcePreset" in element:
                if element["sourcePreset"] not in BOX_SHAPE_PRESETS:
                    raise ValueError(
                        f"slide {index} element sourcePreset is not an "
                        "eligible box/card preset"
                    )
                if (
                    mode != "overlay"
                    or type(slide.get("src")) is not int
                    or not isinstance(element.get("ref"), dict)
                    or bool(element.get("readonly"))
                ):
                    raise ValueError(
                        f"slide {index} element sourcePreset is valid only "
                        "for an editable imported AutoShape"
                    )
            if "corners" in element:
                if element["corners"] != "sharp":
                    raise ValueError(
                        f"slide {index} element corners must be sharp"
                    )
                if element.get("sourcePreset") not in BOX_SHAPE_PRESETS:
                    raise ValueError(
                        f"slide {index} element corners needs eligible "
                        "imported sourcePreset metadata"
                    )
        animations = slide.get("animations", [])
        if not isinstance(animations, list):
            raise ValueError(f"slide {index} animations must be an array")
        local_ids = {
            item.get("id")
            for item in list(slide.get("elements", [])) + list(slide.get("decor", []))
            if isinstance(item, dict) and isinstance(item.get("id"), str)
        }
        media_ids = {
            item.get("id")
            for item in slide.get("decor", [])
            if isinstance(item, dict) and item.get("kind") == "media"
        }
        animated_ids = set()
        for animation_index, animation in enumerate(animations):
            label = f"slide {index} animation {animation_index}"
            if not isinstance(animation, dict):
                raise ValueError(f"{label} must be an object")
            unknown = set(animation) - {
                "targetId", "effect", "trigger", "duration", "delay",
            }
            if unknown:
                raise ValueError(
                    f"{label} has unknown field: "
                    f"{sorted(str(field) for field in unknown)[0]}"
                )
            target_id = animation.get("targetId")
            if target_id not in local_ids:
                raise ValueError(
                    f"{label}.targetId must name an object on the same slide"
                )
            if target_id in media_ids:
                raise ValueError(f"{label}.targetId cannot name a media item")
            if target_id in animated_ids:
                raise ValueError(f"{label}.targetId is duplicated")
            animated_ids.add(target_id)
            if animation.get("effect") not in ("appear", "fade", "wipe", "rise-up"):
                raise ValueError(
                    f"{label}.effect must be appear, fade, wipe, or rise-up"
                )
            trigger = animation.get("trigger")
            if trigger not in ("click", "with-previous", "after-previous"):
                raise ValueError(
                    f"{label}.trigger must be click, with-previous, or after-previous"
                )
            if animation_index == 0 and trigger != "click":
                raise ValueError(f"{label}.trigger must be click for the first animation")
            duration = animation.get("duration")
            if (isinstance(duration, bool)
                    or not isinstance(duration, (int, float))
                    or not math.isfinite(duration)
                    or not 0.1 <= duration <= 10):
                raise ValueError(
                    f"{label}.duration must be from 0.1 through 10 seconds"
                )
            delay = animation.get("delay")
            if (isinstance(delay, bool)
                    or not isinstance(delay, (int, float))
                    or not math.isfinite(delay)
                    or not 0 <= delay <= 30):
                raise ValueError(
                    f"{label}.delay must be from 0 through 30 seconds"
                )


# ---------------------------------------------------------------- text helpers
def apply_text(tf, value):
    """Replace a text frame's content with a single string, preserving the
    first run's formatting where possible."""
    value = "" if value is None else str(value)
    paras = tf.paragraphs
    p0 = paras[0]
    for p in list(paras[1:]):
        p._p.getparent().remove(p._p)
    if p0.runs:
        p0.runs[0].text = value
        for r in list(p0.runs[1:]):
            r._r.getparent().remove(r._r)
    else:
        p0.text = value


def _copy_para_format(src_para, dst_para):
    """Clone src_para's paragraph properties and first-run formatting (size,
    color, bold, font) onto dst_para. Used only for lines the model adds beyond
    what the base deck had, so they match instead of rendering bare."""
    if src_para is None or dst_para is None:
        return
    src_pPr = src_para._p.find(qn("a:pPr"))
    if src_pPr is not None:
        old_pPr = dst_para._p.find(qn("a:pPr"))
        if old_pPr is not None:
            dst_para._p.remove(old_pPr)
        dst_para._p.insert(0, copy.deepcopy(src_pPr))
    if not src_para.runs or not dst_para.runs:
        return
    src_rPr = src_para.runs[0]._r.find(qn("a:rPr"))
    if src_rPr is None:
        return
    dst_r = dst_para.runs[0]._r
    old_rPr = dst_r.find(qn("a:rPr"))
    if old_rPr is not None:
        dst_r.remove(old_rPr)
    dst_r.insert(0, copy.deepcopy(src_rPr))


def _set_para_text(p, text):
    """Set one paragraph's text, keeping that paragraph's own run formatting."""
    if p.runs:
        p.runs[0].text = text
        for r in list(p.runs[1:]):
            r._r.getparent().remove(r._r)
        return
    # An empty base paragraph carries its formatting in <a:endParaRPr>. Setting
    # .text creates a bare run, so lift that formatting onto the new run.
    end_rpr = p._p.find(qn("a:endParaRPr"))
    p.text = text
    if end_rpr is not None and p.runs:
        r = p.runs[0]._r
        old = r.find(qn("a:rPr"))
        if old is not None:
            r.remove(old)
        rpr = copy.deepcopy(end_rpr)
        rpr.tag = qn("a:rPr")
        r.insert(0, rpr)


def apply_bullets(tf, items):
    """Replace a text frame's paragraphs with `items`, preserving each base
    paragraph's own formatting.

    python-pptx's add_paragraph() creates a bare <a:p> with no run properties,
    so rebuilding every line after the first that way silently dropped the base
    deck's explicit run color and 20pt sizing: captions rendered near-white on
    pale fills (illegible), and sub-labels fell back to the ~18pt inherited
    size, under the 20pt house-rule floor. Reusing each paragraph in place keeps
    its own color and size; only genuinely new lines are cloned from the last
    base paragraph so they are never bare.
    """
    items = [str(x) for x in (items or [""])] or [""]
    paras = list(tf.paragraphs)
    template = paras[-1] if paras else None
    for i, it in enumerate(items):
        if i < len(paras):
            _set_para_text(paras[i], it)
        else:
            new_p = tf.add_paragraph()
            _set_para_text(new_p, it)
            _copy_para_format(template, new_p)
    for p in paras[len(items):]:
        p._p.getparent().remove(p._p)


_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}


def apply_style(shape, style):
    """Apply an element's style dict (color/bold/size/align/fill/outline/font) onto a
    real shape. Runs every render, so it must be idempotent: style already
    mirrors the base deck's own formatting until a set-style/set-fill edit
    changes it, so re-applying the same values each pass is a harmless
    no-op. font is the one field with no base.pptx equivalent to mirror,
    model.json simply never carries a run's font name until a set-style
    --font edit adds one (see lint_fonts.py, which reads the real rendered
    run names since this was previously a read-only lint-only gap)."""
    if not style:
        return
    if getattr(shape, "has_text_frame", False):
        tf = shape.text_frame
        color = style.get("color")
        size = style.get("size")
        bold = style.get("bold")
        font = style.get("font")
        align = _ALIGN.get(style.get("align"))
        for para in tf.paragraphs:
            if align is not None:
                para.alignment = align
            for run in para.runs:
                if color:
                    run.font.color.rgb = _hex_to_rgb(color)
                if size:
                    run.font.size = Pt(size)
                if bold is not None:
                    run.font.bold = bool(bold)
                if font:
                    run.font.name = font
    fill = style.get("fill")
    if fill:
        try:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _hex_to_rgb(fill)
            # Imported text cards can retain a source outline even after their
            # fill is moved onto the approved deck palette. Carry the requested
            # fill through to any explicitly prohibited red, green, or orange
            # outline so a fill-only cleanup is visually complete. Also remove
            # the legacy navy outline when a card is explicitly moved to UCR
            # blue, which otherwise renders darker than matching blue peers.
            try:
                line_rgb = shape.line.fill.fore_color.rgb
                line_hex = None if line_rgb is None else str(line_rgb).upper()
                fill_hex = str(fill).upper()
                if (
                    line_hex in {"CE3B3B", "0A8A0A", "E9A400"}
                    or (line_hex == "14315A" and fill_hex == "#2A78D6")
                ):
                    shape.line.color.rgb = _hex_to_rgb(fill)
            except Exception:
                pass
        except Exception:
            pass
    outline = style.get("outline")
    if outline:
        try:
            shape.line.color.rgb = _hex_to_rgb(outline)
        except Exception:
            pass


def apply_box(shape, box, slide_w, slide_h):
    """Reposition/resize a real shape from a model box dict {x,y,w,h}, each a
    fraction of slide width/height (the same units import_pptx.py's geom()
    captured at import). Runs every render like apply_style, so a shape with
    no set-box edit keeps whatever base.pptx already gave it (box mirrors the
    base geometry until an edit changes it, harmless no-op otherwise). This is
    the missing lever several polish passes flagged: overlay elements/decor
    carried box geometry as read-only metadata, so a box too short for its
    text was only fixable by shortening the text forever, never by growing
    the box itself."""
    if not box:
        return
    try:
        if "x" in box:
            shape.left = Emu(round(box["x"] * slide_w))
        if "y" in box:
            shape.top = Emu(round(box["y"] * slide_h))
        if "w" in box:
            shape.width = Emu(round(box["w"] * slide_w))
        if "h" in box:
            shape.height = Emu(round(box["h"] * slide_h))
    except Exception:
        pass


def apply_corners(shape, item, label):
    """Apply a sharp-corner intent to one imported backing AutoShape in place.

    Changing only the existing p:spPr/a:prstGeom node retains the p:cNvPr
    shape id, text body, transform, fill/line/effects, z-order, and animation
    target identity. The model's immutable-source metadata and the live shape
    are both checked so a forged target can never turn another object class
    into a rectangle.
    """
    if item.get("corners") != "sharp":
        return
    expected = item.get("sourcePreset")
    if expected not in BOX_SHAPE_PRESETS:
        raise ValueError(f"{label} has no eligible imported sourcePreset")
    try:
        is_auto_shape = shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        is_placeholder = shape.is_placeholder
    except Exception as error:
        raise ValueError(f"{label} is not an inspectable imported AutoShape") from error
    if not is_auto_shape or is_placeholder:
        raise ValueError(f"{label} is not an eligible imported box/card AutoShape")
    preset_geometry = shape._element.find(f".//{qn('a:prstGeom')}")
    if preset_geometry is None or preset_geometry.get("prst") != expected:
        actual = None if preset_geometry is None else preset_geometry.get("prst")
        raise ValueError(
            f"{label} backing geometry mismatch: expected {expected}, found {actual}"
        )
    preset_geometry.set("prst", "rect")
    # Rectangle geometry has no adjustment handles. Replace only avLst so an
    # unrelated extension node, if PowerPoint stored one on this geometry,
    # remains intact.
    for adjustment_list in preset_geometry.findall(qn("a:avLst")):
        preset_geometry.remove(adjustment_list)
    preset_geometry.insert(0, etree.Element(qn("a:avLst")))


def _is_full_slide_box(box, tolerance=0.002):
    """Return whether a normalized box covers the slide.

    Imported decks commonly use a full-slide rectangle as their background.
    If that rectangle is moved aside to fix an incorrect top-layer stacking
    order, preserve its color as the actual slide background so the visual
    theme survives without obscuring any content.
    """
    if not isinstance(box, dict):
        return False
    try:
        return (
            abs(float(box.get("x", 0))) <= tolerance
            and abs(float(box.get("y", 0))) <= tolerance
            and abs(float(box.get("w", 0)) - 1) <= tolerance
            and abs(float(box.get("h", 0)) - 1) <= tolerance
        )
    except (TypeError, ValueError):
        return False


def preserve_hidden_full_slide_decor_as_background(slide, decor, base_decor):
    """Move a hidden full-slide decor's color into the slide background."""
    base_box = base_decor.get("box")
    current_box = decor.get("box")
    was_full_slide = _is_full_slide_box(base_box)
    try:
        base_area = float(base_box.get("w", 0)) * float(base_box.get("h", 0))
        current_area = (
            float(current_box.get("w", 0)) * float(current_box.get("h", 0))
        )
    except (AttributeError, TypeError, ValueError):
        return
    # Also support the explicit hide-to-background repair used for large
    # imported panels that were expanded into a full-slide overlay in an
    # earlier edit. The pristine base no longer records that intermediate
    # geometry, but a large source panel collapsed to a tiny corner and given
    # a new fill is an unambiguous background request.
    hidden_large_panel = (
        base_area >= 0.20
        and current_area <= 0.0002
        and decor.get("fill") != base_decor.get("fill")
    )
    if not was_full_slide and not hidden_large_panel:
        return
    if _is_full_slide_box(current_box):
        return
    fill = decor.get("fill") or base_decor.get("fill")
    if not fill:
        return
    background_fill = slide.background.fill
    background_fill.solid()
    background_fill.fore_color.rgb = _hex_to_rgb(fill)


def add_generated_rect(slide, decor, slide_w, slide_h):
    """Append one validated solid rectangle to a rendered slide.

    Rendering starts from either the pristine imported base or a fresh native
    slide, so appending model-generated rectangles is naturally idempotent:
    prior generated output is never the next render's input.
    """
    box = decor["box"]
    slide_w = int(slide_w)
    slide_h = int(slide_h)
    # A legal but extremely small positive normalized width can round to zero
    # EMUs. Convert edges and clamp the start one EMU inside the slide so every
    # w/h > 0 model rectangle remains a real, in-bounds PowerPoint shape.
    left = min(slide_w - 1, max(0, round(box["x"] * slide_w)))
    top = min(slide_h - 1, max(0, round(box["y"] * slide_h)))
    right = min(slide_w, max(left + 1, round((box["x"] + box["w"]) * slide_w)))
    bottom = min(slide_h, max(top + 1, round((box["y"] + box["h"]) * slide_h)))
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Emu(left),
        Emu(top),
        Emu(right - left),
        Emu(bottom - top),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex_to_rgb(decor["fill"])
    # PowerPoint's default auto-shape outline otherwise adds a theme-colored
    # one-pixel border the model never asked for.
    shape.line.fill.background()
    shape.name = f"Suite shape {decor['id']}"
    return shape


def add_generated_picture(slide, decor, slide_w, slide_h):
    """Append one validated project image to a rendered slide."""
    box = decor["box"]
    slide_w = int(slide_w)
    slide_h = int(slide_h)
    left = min(slide_w - 1, max(0, round(box["x"] * slide_w)))
    top = min(slide_h - 1, max(0, round(box["y"] * slide_h)))
    right = min(slide_w, max(left + 1, round((box["x"] + box["w"]) * slide_w)))
    bottom = min(slide_h, max(top + 1, round((box["y"] + box["h"]) * slide_h)))
    source = _resolve_image_source(decor["source"])
    picture = slide.shapes.add_picture(
        str(source),
        Emu(left),
        Emu(top),
        Emu(right - left),
        Emu(bottom - top),
    )
    picture.name = f"Suite image {decor['id']}"
    return picture


def add_native_generated_decor(slide, decor_items, slide_w, slide_h):
    """Render suite-generated native decor behind native text.

    PowerPoint's shape-tree order is back-to-front. A newly-added slide may
    already contain layout placeholders, and add_shape/add_picture append in
    front of them. Render generated items in model order, then move their XML
    nodes as one ordered block immediately after the shape-tree metadata.
    Media remains a post-save PowerPoint COM operation.
    """
    targets = {}
    rendered = []
    for decor in decor_items:
        if decor.get("generated") is not True:
            continue
        if decor.get("kind") == "media":
            continue
        if decor.get("kind") == "pic":
            shape = add_generated_picture(
                slide, decor, slide_w, slide_h
            )
        else:
            shape = add_generated_rect(
                slide, decor, slide_w, slide_h
            )
        rendered.append(shape)
        targets[decor["id"]] = shape.shape_id

    if rendered:
        shape_tree = slide.shapes._spTree
        elements = [shape._element for shape in rendered]
        for element in elements:
            shape_tree.remove(element)
        # p:spTree begins with p:nvGrpSpPr and p:grpSpPr. Inserting after
        # those nodes puts decor behind every placeholder/text shape while
        # retaining the model's back-to-front decor order.
        for offset, element in enumerate(elements):
            shape_tree.insert(2 + offset, element)
    return targets


def apply_notes(slide, text):
    if text is None:
        return
    if text == "" and not slide.has_notes_slide:
        return
    slide.notes_slide.notes_text_frame.text = text


# ----------------------------------------------------------- slide transitions
def _transition_speed(duration):
    """Return a coarse legacy fallback bucket for a precise p14 duration."""
    if duration <= 0.5:
        return "fast"
    if duration <= 1.0:
        return "med"
    return "slow"


def _contains_slide_transition(element):
    transition_tag = _xml_name(PML_NS, "transition")
    return any(node.tag == transition_tag for node in element.iter())


def _remove_slide_transition_nodes(slide):
    """Remove only slide-transition representations, preserving unrelated MC."""
    root = slide._element
    transition_tag = _xml_name(PML_NS, "transition")
    alternate_content_tag = _xml_name(MC_NS, "AlternateContent")
    for child in list(root):
        if child.tag == transition_tag:
            root.remove(child)
        elif child.tag == alternate_content_tag and _contains_slide_transition(child):
            root.remove(child)


def _make_fade_transition(duration):
    """Build a precise Office 2010+ Fade plus a legacy-compatible fallback."""
    alternate = etree.Element(
        _xml_name(MC_NS, "AlternateContent"),
        nsmap={"mc": MC_NS, "p14": P14_NS},
    )
    choice = etree.SubElement(alternate, _xml_name(MC_NS, "Choice"))
    choice.set("Requires", "p14")
    fallback = etree.SubElement(alternate, _xml_name(MC_NS, "Fallback"))
    duration_ms = str(int(round(duration * 1000)))
    speed = _transition_speed(duration)

    for parent, precise in ((choice, True), (fallback, False)):
        transition = etree.SubElement(
            parent, _xml_name(PML_NS, "transition")
        )
        transition.set("spd", speed)
        transition.set("advClick", "1")
        if precise:
            transition.set(_xml_name(P14_NS, "dur"), duration_ms)
        fade = etree.SubElement(transition, _xml_name(PML_NS, "fade"))
        fade.set("thruBlk", "0")
    return alternate


def apply_slide_transition(slide, transition):
    """Apply one explicit transition policy; None means preserve imported XML."""
    if transition is None:
        return
    _remove_slide_transition_nodes(slide)
    if transition["effect"] == "none":
        return

    root = slide._element
    alternate = _make_fade_transition(transition["duration"])
    timing_tag = _xml_name(PML_NS, "timing")
    ext_list_tag = _xml_name(PML_NS, "extLst")
    alternate_content_tag = _xml_name(MC_NS, "AlternateContent")
    insert_at = len(root)
    for index, child in enumerate(root):
        if child.tag in (timing_tag, ext_list_tag):
            insert_at = index
            break
        # Markup-Compatibility wrappers occupy the schema position of their
        # selected child. A timing/extLst wrapper is therefore also an anchor.
        if child.tag == alternate_content_tag and any(
                node.tag in (timing_tag, ext_list_tag)
                for node in child.iter()):
            insert_at = index
            break
    root.insert(insert_at, alternate)


# ---------------------------------------------------------- embedded media
def _requested_media(model):
    return [
        (slide_index, decor)
        for slide_index, slide in enumerate(model.get("slides", []), 1)
        for decor in slide.get("decor", [])
        if decor.get("generated") is True and decor.get("kind") == "media"
    ]


def _stream_sha256(stream):
    digest = hashlib.sha256()
    while True:
        chunk = stream.read(1024 * 1024)
        if not chunk:
            return digest.hexdigest()
        digest.update(chunk)


def _file_sha256(path):
    with open(path, "rb") as stream:
        return _stream_sha256(stream)


def _verify_embedded_media(out_path, sources):
    """Fail unless every requested MP4 is stored internally in the PPTX."""
    relationship_ns = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    source_hashes = {_file_sha256(source) for source in sources}
    embedded_hashes = set()
    internal_media_relationships = 0
    with zipfile.ZipFile(out_path) as package:
        media_names = [
            name for name in package.namelist()
            if name.lower().startswith("ppt/media/")
            and name.lower().endswith(".mp4")
        ]
        if not media_names:
            raise RuntimeError(
                "PowerPoint did not embed the requested MP4 under ppt/media"
            )
        for name in media_names:
            with package.open(name) as stream:
                embedded_hashes.add(_stream_sha256(stream))
        missing_hashes = source_hashes - embedded_hashes
        if missing_hashes:
            raise RuntimeError(
                "PowerPoint output does not contain the requested MP4 bytes"
            )

        for name in package.namelist():
            if not name.lower().endswith(".rels"):
                continue
            try:
                root = etree.fromstring(package.read(name))
            except Exception as error:
                raise RuntimeError(
                    f"PowerPoint wrote an unreadable relationship part: {name}"
                ) from error
            for relationship in root.findall(
                    f"{{{relationship_ns}}}Relationship"):
                rel_type = str(relationship.get("Type") or "").lower()
                target = str(relationship.get("Target") or "")
                media_relation = (
                    rel_type.endswith("/video")
                    or rel_type.endswith("/media")
                    or target.lower().endswith(".mp4")
                )
                if not media_relation:
                    continue
                if str(relationship.get("TargetMode") or "").lower() == "external":
                    raise RuntimeError(
                        "PowerPoint created an external media relationship"
                    )
                if target.lower().endswith(".mp4"):
                    internal_media_relationships += 1
        if internal_media_relationships < len(sources):
            raise RuntimeError(
                "PowerPoint did not relate every video shape to embedded MP4 media"
            )


def _find_com_shape_by_name(slide, name):
    for shape_index in range(1, slide.Shapes.Count + 1):
        shape = slide.Shapes.Item(shape_index)
        if str(shape.Name) == name:
            return shape
    return None


def _generated_decor_shape_name(decor):
    if decor.get("kind") == "media":
        return f"Suite media {decor['id']}"
    if decor.get("kind") == "pic":
        return f"Suite image {decor['id']}"
    return f"Suite shape {decor['id']}"


def _send_native_decor_behind_content(presentation, model):
    """Restore native decor model order after COM inserts embedded media."""
    for slide_index, slide_model in enumerate(model.get("slides", []), 1):
        if slide_model.get("src") is not None:
            continue
        generated = [
            decor for decor in slide_model.get("decor", [])
            if decor.get("generated") is True
        ]
        if not any(decor.get("kind") == "media" for decor in generated):
            continue
        slide = presentation.Slides.Item(slide_index)
        # Repeated send-to-back calls reverse the requested sequence, so walk
        # it backwards to leave model order as the final back-to-front order.
        for decor in reversed(generated):
            shape = _find_com_shape_by_name(
                slide, _generated_decor_shape_name(decor)
            )
            if shape is None:
                raise RuntimeError(
                    f"PowerPoint could not resolve generated decor "
                    f"{decor['id']} on slide {slide_index}"
                )
            shape.ZOrder(1)  # msoSendToBack


def _current_powerpoint_pids():
    """Return the running POWERPNT.EXE process IDs without creating a process.

    A native Toolhelp snapshot avoids invoking PowerShell from the renderer.
    Snapshot failures are intentionally surfaced to the caller: treating an
    unreadable snapshot as an empty set could falsely claim ownership of a
    PowerPoint instance that was already running.
    """
    if os.name != "nt":
        raise OSError("PowerPoint process discovery requires Windows")

    class PROCESSENTRY32W(ctypes.Structure):
        _fields_ = [
            ("dwSize", wintypes.DWORD),
            ("cntUsage", wintypes.DWORD),
            ("th32ProcessID", wintypes.DWORD),
            ("th32DefaultHeapID", ctypes.c_size_t),
            ("th32ModuleID", wintypes.DWORD),
            ("cntThreads", wintypes.DWORD),
            ("th32ParentProcessID", wintypes.DWORD),
            ("pcPriClassBase", wintypes.LONG),
            ("dwFlags", wintypes.DWORD),
            ("szExeFile", wintypes.WCHAR * 260),
        ]

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_snapshot = kernel32.CreateToolhelp32Snapshot
    create_snapshot.argtypes = [wintypes.DWORD, wintypes.DWORD]
    create_snapshot.restype = wintypes.HANDLE
    process_first = kernel32.Process32FirstW
    process_first.argtypes = [
        wintypes.HANDLE, ctypes.POINTER(PROCESSENTRY32W),
    ]
    process_first.restype = wintypes.BOOL
    process_next = kernel32.Process32NextW
    process_next.argtypes = [
        wintypes.HANDLE, ctypes.POINTER(PROCESSENTRY32W),
    ]
    process_next.restype = wintypes.BOOL
    close_handle = kernel32.CloseHandle
    close_handle.argtypes = [wintypes.HANDLE]
    close_handle.restype = wintypes.BOOL

    snapshot = create_snapshot(0x00000002, 0)  # TH32CS_SNAPPROCESS
    invalid_handle = ctypes.c_void_p(-1).value
    if snapshot == invalid_handle:
        raise ctypes.WinError(ctypes.get_last_error())

    try:
        entry = PROCESSENTRY32W()
        entry.dwSize = ctypes.sizeof(entry)
        if not process_first(snapshot, ctypes.byref(entry)):
            error_code = ctypes.get_last_error()
            if error_code == 18:  # ERROR_NO_MORE_FILES
                return set()
            raise ctypes.WinError(error_code)
        result = set()
        while True:
            if str(entry.szExeFile).casefold() == "powerpnt.exe":
                result.add(int(entry.th32ProcessID))
            if not process_next(snapshot, ctypes.byref(entry)):
                error_code = ctypes.get_last_error()
                if error_code != 18:  # ERROR_NO_MORE_FILES
                    raise ctypes.WinError(error_code)
                return result
    finally:
        close_handle(snapshot)


def _powerpoint_application_pid(application):
    """Resolve the process that owns a PowerPoint Application COM object."""
    if os.name != "nt":
        raise OSError("PowerPoint window discovery requires Windows")
    hwnd = int(application.HWND)
    if hwnd <= 0:
        raise OSError("PowerPoint returned an invalid application window handle")
    process_id = wintypes.DWORD()
    user32 = ctypes.WinDLL("user32", use_last_error=True)
    get_window_process = user32.GetWindowThreadProcessId
    get_window_process.argtypes = [wintypes.HWND, ctypes.POINTER(wintypes.DWORD)]
    get_window_process.restype = wintypes.DWORD
    if not get_window_process(wintypes.HWND(hwnd), ctypes.byref(process_id)):
        raise ctypes.WinError(ctypes.get_last_error())
    if not process_id.value:
        raise OSError("PowerPoint application window has no owning process")
    return int(process_id.value)


def _dispatch_powerpoint_application(dispatch_ex):
    """Activate PowerPoint and conservatively determine process ownership.

    PowerPoint behaves as a single-instance COM server even when called via
    DispatchEx. We may call Application.Quit only when all three signals agree:
    exactly one POWERPNT PID appeared, the returned Application belongs to it,
    and neither process snapshot failed. Any uncertainty deliberately yields
    ``owns_application = False``.
    """
    before = None
    try:
        before = _current_powerpoint_pids()
    except Exception:
        pass

    application = dispatch_ex("PowerPoint.Application")

    after = None
    application_pid = None
    try:
        after = _current_powerpoint_pids()
    except Exception:
        pass
    try:
        application_pid = _powerpoint_application_pid(application)
    except Exception:
        pass

    owns_application = False
    if before is not None and after is not None and application_pid is not None:
        new_pids = after - before
        owns_application = new_pids == {application_pid}
    return application, owns_application


def _quit_powerpoint_if_owned(application, owns_application):
    """Quit only an application whose process this activation proved it made."""
    if application is None or not owns_application:
        return
    try:
        application.Quit()
    except Exception:
        pass


def _assert_media_playback(shape, decor, slide_index):
    if shape is None:
        raise RuntimeError(
            f"PowerPoint could not reopen media {decor['id']} on slide {slide_index}"
        )
    settings = shape.AnimationSettings.PlaySettings
    expected_play = -1 if decor["autoplay"] else 0
    expected_loop = -1 if decor["loop"] else 0
    if int(shape.Visible) != -1:
        raise RuntimeError(
            f"PowerPoint did not keep media {decor['id']} visible"
        )
    if int(settings.HideWhileNotPlaying) != 0:
        raise RuntimeError(
            f"PowerPoint hid media {decor['id']} while it was not playing"
        )
    if int(settings.PlayOnEntry) != expected_play:
        raise RuntimeError(
            f"PowerPoint did not persist autoplay for media {decor['id']}"
        )
    if int(settings.LoopUntilStopped) != expected_loop:
        raise RuntimeError(
            f"PowerPoint did not persist looping for media {decor['id']}"
        )


def apply_powerpoint_media(out_path, model):
    """Embed MP4s with PowerPoint's native media API and verify persistence."""
    requested = _requested_media(model)
    if not requested:
        return
    if os.name != "nt":
        raise RuntimeError("embedded MP4 media requires PowerPoint on Windows")

    resolved = [
        (slide_index, decor, _resolve_video_source(decor["source"]))
        for slide_index, decor in requested
    ]
    try:
        import pythoncom
        import win32com.client
    except ImportError as error:
        raise RuntimeError(
            "embedded MP4 media requires the installed pywin32 package"
        ) from error

    application = None
    owns_application = False
    presentation = None
    slide = None
    shape = None
    settings = None
    absolute_output = str(Path(out_path).resolve())
    pythoncom.CoInitialize()
    try:
        application, owns_application = _dispatch_powerpoint_application(
            win32com.client.DispatchEx
        )
        presentation = application.Presentations.Open(
            absolute_output, False, False, False
        )
        if int(presentation.Slides.Count) != len(model.get("slides", [])):
            raise RuntimeError(
                "PowerPoint slide count does not match the media model"
            )
        slide_width = float(presentation.PageSetup.SlideWidth)
        slide_height = float(presentation.PageSetup.SlideHeight)
        for slide_index, decor, source in resolved:
            box = decor["box"]
            slide = presentation.Slides.Item(slide_index)
            left = float(box["x"]) * slide_width
            top = float(box["y"]) * slide_height
            width = float(box["w"]) * slide_width
            height = float(box["h"]) * slide_height
            shape = slide.Shapes.AddMediaObject2(
                str(source),
                0,   # LinkToFile = msoFalse
                -1,  # SaveWithDocument = msoTrue
                left,
                top,
                width,
                height,
            )
            # PowerPoint may preserve the video's source aspect ratio during
            # AddMediaObject2 even when explicit dimensions were supplied.
            # The persisted model box is authoritative, so unlock and apply
            # each edge again after insertion.
            shape.LockAspectRatio = 0
            shape.Left = left
            shape.Top = top
            shape.Width = width
            shape.Height = height
            shape.Name = f"Suite media {decor['id']}"
            shape.Visible = -1
            settings = shape.AnimationSettings.PlaySettings
            settings.PlayOnEntry = -1 if decor["autoplay"] else 0
            settings.LoopUntilStopped = -1 if decor["loop"] else 0
            settings.HideWhileNotPlaying = 0
            _assert_media_playback(shape, decor, slide_index)
        _send_native_decor_behind_content(presentation, model)
        presentation.Save()
        presentation.Close()
        presentation = None
        presentation = application.Presentations.Open(
            absolute_output, False, False, False
        )
        for slide_index, decor, _ in resolved:
            slide = presentation.Slides.Item(slide_index)
            shape = _find_com_shape_by_name(
                slide, f"Suite media {decor['id']}"
            )
            _assert_media_playback(shape, decor, slide_index)
        presentation.Close()
        presentation = None
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        presentation = None
        settings = None
        shape = None
        slide = None
        _quit_powerpoint_if_owned(application, owns_application)
        application = None
        pythoncom.CoUninitialize()

    _verify_embedded_media(
        out_path, [source for _, _, source in resolved]
    )


# ---------------------------------------------------------- object animations
def apply_powerpoint_animations(out_path, model, animation_targets):
    """Use PowerPoint's native sequence API to author real entrance builds.

    python-pptx intentionally has no animation API. The suite runs on Windows
    with PowerPoint already required for PDF/thumbnail export, so using the
    same installed application here gives us PowerPoint-authored timing XML
    instead of hand-maintaining a fragile subset of the schema.
    """
    requested = [
        (slide_index, animation)
        for slide_index, slide in enumerate(model.get("slides", []))
        for animation in slide.get("animations", [])
    ]
    if not requested:
        return
    if os.name != "nt":
        raise RuntimeError("native object animations require PowerPoint on Windows")
    if len(animation_targets) != len(model.get("slides", [])):
        raise RuntimeError("animation target map does not match the rendered slide count")
    for slide_index, animation in requested:
        if animation["targetId"] not in animation_targets[slide_index]:
            raise ValueError(
                f"animation target {animation['targetId']} has no rendered PowerPoint shape"
            )

    try:
        import pythoncom
        import win32com.client
    except ImportError as error:
        raise RuntimeError(
            "native object animations require the installed pywin32 package"
        ) from error

    effect_ids = {"appear": 1, "fade": 10, "wipe": 22, "rise-up": 34}
    trigger_ids = {"click": 1, "with-previous": 2, "after-previous": 3}
    application = None
    owns_application = False
    presentation = None
    slide = None
    shape = None
    sequence = None
    effect = None
    pythoncom.CoInitialize()
    try:
        application, owns_application = _dispatch_powerpoint_application(
            win32com.client.DispatchEx
        )
        presentation = application.Presentations.Open(
            str(Path(out_path).resolve()), False, False, False
        )
        for slide_index, model_slide in enumerate(model.get("slides", []), 1):
            if not model_slide.get("animations"):
                continue
            slide = presentation.Slides.Item(slide_index)
            by_id = {}
            for shape_index in range(1, slide.Shapes.Count + 1):
                shape = slide.Shapes.Item(shape_index)
                by_id[int(shape.Id)] = shape
            sequence = slide.TimeLine.MainSequence
            for animation in model_slide["animations"]:
                shape_id = animation_targets[slide_index - 1][animation["targetId"]]
                shape = by_id.get(int(shape_id))
                if shape is None:
                    raise RuntimeError(
                        f"PowerPoint could not find rendered shape {shape_id} "
                        f"for {animation['targetId']} on slide {slide_index}"
                    )
                effect = sequence.AddEffect(
                    shape,
                    effect_ids[animation["effect"]],
                    0,
                    trigger_ids[animation["trigger"]],
                )
                effect.Timing.Duration = float(animation["duration"])
                effect.Timing.TriggerDelayTime = float(animation["delay"])
        presentation.Save()
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        presentation = None
        effect = None
        sequence = None
        shape = None
        slide = None
        _quit_powerpoint_if_owned(application, owns_application)
        application = None
        pythoncom.CoUninitialize()


# --------------------------------------------------------------- native filler
def _find_placeholder(slide, idxs):
    by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    for idx in idxs:
        if idx in by_idx:
            return by_idx[idx]
    return None


def fill_native(slide, elements, accent, slide_width, slide_height):
    """Place the model's elements onto a freshly-added slide (native mode)."""
    stack_top = 1.6
    used_body = False
    targets = {}

    def remember(element, shape):
        if isinstance(element.get("id"), str):
            targets[element["id"]] = shape.shape_id
            shape.name = f"Suite element {element['id']}"

    for el in elements:
        etype = el.get("type", "body")
        text = el.get("text", "")
        items = el.get("items")

        if etype == "title":
            ph = _find_placeholder(slide, [0]) or (slide.shapes.title if slide.shapes.title else None)
            if ph is not None:
                apply_text(ph.text_frame, text)
                apply_style(ph, el.get("style"))
                apply_box(ph, el.get("box"), slide_width, slide_height)
                remember(el, ph)
                continue
        if etype == "subtitle":
            ph = _find_placeholder(slide, [1])
            if ph is not None and ph.placeholder_format.idx != 0:
                apply_text(ph.text_frame, text)
                apply_style(ph, el.get("style"))
                apply_box(ph, el.get("box"), slide_width, slide_height)
                remember(el, ph)
                continue
        if etype in ("bullets",) or (etype == "body" and items):
            ph = None if used_body else _find_placeholder(slide, [1, 2, 13, 14])
            if ph is not None and ph.placeholder_format.idx != 0:
                apply_bullets(ph.text_frame, items or [text])
                apply_style(ph, el.get("style"))
                apply_box(ph, el.get("box"), slide_width, slide_height)
                used_body = True
                remember(el, ph)
                continue
        if etype == "body":
            ph = None if used_body else _find_placeholder(slide, [1, 2, 13, 14])
            if ph is not None and ph.placeholder_format.idx != 0:
                apply_text(ph.text_frame, text)
                apply_style(ph, el.get("style"))
                apply_box(ph, el.get("box"), slide_width, slide_height)
                used_body = True
                remember(el, ph)
                continue

        box = slide.shapes.add_textbox(Inches(0.7), Inches(stack_top), Inches(11.9), Inches(1.0))
        tf = box.text_frame; tf.word_wrap = True
        if items:
            apply_bullets(tf, items)
        elif etype == "heading":
            apply_text(tf, text)
            if tf.paragraphs[0].runs:
                r = tf.paragraphs[0].runs[0]; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = accent
        else:
            apply_text(tf, text)
        # Explicit element style/fill and box geometry are meaningful on
        # native slides; an absent box preserves the legacy auto-stack.
        apply_style(box, el.get("style"))
        apply_box(box, el.get("box"), slide_width, slide_height)
        remember(el, box)
        stack_top += 1.1
    return targets


def _layout_for(prs, key):
    if key in LAYOUT_MAP and LAYOUT_MAP[key] < len(prs.slide_layouts):
        return prs.slide_layouts[LAYOUT_MAP[key]]
    for lay in prs.slide_layouts:
        if getattr(lay, "name", "") == key:
            return lay
    return prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]


# -------------------------------------------------------------------- renderers
def render_native(model, out_path):
    accent = _hex_to_rgb((model.get("theme") or {}).get("accent", "#2563eb"))
    prs = Presentation()
    slide_w = model.get("slideW", 13.333)
    slide_h = model.get("slideH", 7.5)
    if (isinstance(slide_w, bool) or not isinstance(slide_w, (int, float))
            or slide_w <= 0):
        raise ValueError("slideW must be a positive number")
    if (isinstance(slide_h, bool) or not isinstance(slide_h, (int, float))
            or slide_h <= 0):
        raise ValueError("slideH must be a positive number")
    prs.slide_width = Inches(slide_w)
    prs.slide_height = Inches(slide_h)
    animation_targets = []
    for ms in model.get("slides", []):
        slide = prs.slides.add_slide(_layout_for(prs, ms.get("layout", "content")))
        targets = fill_native(
            slide, ms.get("elements", []), accent,
            prs.slide_width, prs.slide_height,
        )
        targets.update(
            add_native_generated_decor(
                slide,
                ms.get("decor", []),
                prs.slide_width,
                prs.slide_height,
            )
        )
        animation_targets.append(targets)
        apply_notes(slide, ms.get("notes", ""))
        apply_slide_transition(slide, ms.get("transition"))
    prs.save(out_path)
    apply_powerpoint_animations(out_path, model, animation_targets)
    apply_powerpoint_media(out_path, model)


def render_overlay(model, out_path):
    accent = _hex_to_rgb((model.get("theme") or {}).get("accent", "#2563eb"))
    base = model.get("base")
    if not isinstance(base, str) or not base or not os.path.exists(base):
        raise SystemExit(f"overlay base not found: {base}")
    prs = Presentation(base)
    sldIdLst = prs.slides._sldIdLst
    orig_sldIds = list(sldIdLst)          # <p:sldId> elements, in source order
    base_slides = list(prs.slides)        # slide objects, same order
    n = len(base_slides)
    theme = load_theme_colors(prs)
    base_models = [derive_slide(prs, i, theme) for i in range(n)]

    # A p:sldId XML node cannot occur twice: appending the same source node a
    # second time moves it, silently collapsing two model slides into one.
    # Validate all imported sources/refs before mutating any package XML.
    used_model_sources = set()
    for index, model_slide in enumerate(model.get("slides", [])):
        src = model_slide.get("src")
        if src is None:
            continue
        if type(src) is not int:  # bool is an int subclass; it is not a slide index
            raise ValueError(f"overlay slide {index} src must be an integer or null")
        if not 0 <= src < n:
            raise ValueError(f"overlay slide {index} src is out of range: {src}")
        if src in used_model_sources:
            raise ValueError(f"duplicate overlay src: {src}")
        used_model_sources.add(src)
        shape_by_id = {shape.shape_id: shape for shape in base_slides[src].shapes}
        for element in model_slide.get("elements", []):
            ref = element.get("ref")
            if ref is None:
                continue
            if not isinstance(ref, dict):
                raise ValueError(f"overlay slide {index} element ref must be an object")
            if type(ref.get("s")) is not int or ref.get("s") != src:
                raise ValueError(f"overlay slide {index} element ref.s must equal src")
            sid = ref.get("sid")
            if type(sid) is not int:
                raise ValueError(f"overlay slide {index} element ref.sid must be an integer")
            shape = shape_by_id.get(sid)
            if shape is None or not getattr(shape, "has_text_frame", False):
                raise ValueError(f"overlay slide {index} element references no text shape: {sid}")
        for decor in model_slide.get("decor", []):
            if not isinstance(decor, dict):
                raise ValueError(f"overlay slide {index} contains a non-object decor entry")
            sid = decor.get("sid")
            if sid is None:
                continue
            if type(sid) is not int or sid not in shape_by_id:
                raise ValueError(f"overlay slide {index} decor references no shape: {sid}")

    desired = []                          # sldId elements in final model order
    used_src = set()
    animation_targets = []

    for ms in model.get("slides", []):
        src = ms.get("src")
        if src is not None:
            used_src.add(src)
            slide = base_slides[src]
            shape_by_id = {sh.shape_id: sh for sh in slide.shapes}
            base_element_by_sid = {
                element["ref"]["sid"]: element
                for element in base_models[src].get("elements", [])
                if isinstance(element.get("ref"), dict)
            }
            # Omitting an imported element from the model is the persisted
            # meaning of delete-element. Previously the loop below merely
            # stopped editing that shape, leaving the pristine base shape
            # visible and turning a reported-success delete into a no-op.
            kept_element_sids = {
                element["ref"]["sid"]
                for element in ms.get("elements", [])
                if isinstance(element.get("ref"), dict)
            }
            for sid in set(base_element_by_sid) - kept_element_sids:
                shape = shape_by_id.pop(sid, None)
                if shape is not None:
                    slide.shapes._spTree.remove(shape._element)
            for el in ms.get("elements", []):
                ref = el.get("ref")
                if not ref:
                    continue
                sid = ref.get("sid")
                sh = shape_by_id[sid]
                base_el = base_element_by_sid.get(sid, {})
                # The model stores a coarse preview style and flattened text
                # representation. Reapplying those unchanged values destroyed
                # mixed rich runs on every no-op render. Only true deltas are
                # edits; pristine values leave the base XML entirely alone.
                if _content_differs(el, base_el):
                    if "items" in el:
                        apply_bullets(sh.text_frame, el.get("items"))
                    elif "text" in el:
                        apply_text(sh.text_frame, el.get("text"))
                style_delta = _style_delta(el.get("style"), base_el.get("style"))
                if style_delta:
                    apply_style(sh, style_delta)
                if _box_differs(el.get("box"), base_el.get("box")):
                    apply_box(sh, el.get("box"), prs.slide_width, prs.slide_height)
                apply_corners(
                    sh,
                    el,
                    f"overlay slide {src} element {el.get('id')}",
                )
            targets = {
                element["id"]: element["ref"]["sid"]
                for element in ms.get("elements", [])
                if isinstance(element.get("id"), str)
                and isinstance(element.get("ref"), dict)
                and element["ref"].get("sid") in shape_by_id
            }
            base_decor_by_sid = {
                decor.get("sid"): decor for decor in base_models[src].get("decor", [])
                if decor.get("sid") is not None
            }
            for d in ms.get("decor", []):
                if d.get("generated") is True:
                    if d.get("kind") == "media":
                        # python-pptx has no media API. This entry is inserted
                        # after the package is saved by apply_powerpoint_media.
                        continue
                    if d.get("kind") == "pic":
                        generated = add_generated_picture(
                            slide, d, prs.slide_width, prs.slide_height
                        )
                    else:
                        generated = add_generated_rect(
                            slide, d, prs.slide_width, prs.slide_height
                        )
                    targets[d["id"]] = generated.shape_id
                    continue
                sid = d.get("sid")
                if sid is None:
                    continue
                sh = shape_by_id.get(sid)
                if sh is None:
                    continue
                base_decor = base_decor_by_sid.get(sid, {})
                preserve_hidden_full_slide_decor_as_background(
                    slide, d, base_decor
                )
                fill = d.get("fill")
                if fill and not _same_value("fill", fill, base_decor.get("fill")):
                    try:
                        sh.fill.solid()
                        sh.fill.fore_color.rgb = _hex_to_rgb(fill)
                        # A fill-only palette cleanup can otherwise leave the
                        # imported outline visibly red, green, or orange. When
                        # the source border is one of those prohibited accent
                        # colors, carry the requested UCR fill color through
                        # to the outline as well.
                        try:
                            line_rgb = sh.line.fill.fore_color.rgb
                            if (
                                line_rgb is not None
                                and str(line_rgb).upper()
                                in {"CE3B3B", "0A8A0A", "E9A400"}
                            ):
                                sh.line.color.rgb = _hex_to_rgb(fill)
                        except Exception:
                            pass
                    except Exception:
                        pass
                if _box_differs(d.get("box"), base_decor.get("box")):
                    apply_box(sh, d.get("box"), prs.slide_width, prs.slide_height)
                apply_corners(
                    sh,
                    d,
                    f"overlay slide {src} decor {d.get('id')}",
                )
                if isinstance(d.get("id"), str):
                    targets[d["id"]] = sid
            if ms.get("notes") != base_models[src].get("notes"):
                apply_notes(slide, ms.get("notes"))
            apply_slide_transition(slide, ms.get("transition"))
            desired.append(orig_sldIds[src])
            animation_targets.append(targets)
        else:
            # a brand-new slide added inside the suite — build it natively
            prs.slides.add_slide(_layout_for(prs, ms.get("layout", "content")))
            new_slide = base_slides_len_slide(prs)
            targets = fill_native(
                new_slide, ms.get("elements", []), accent,
                prs.slide_width, prs.slide_height,
            )
            targets.update(
                add_native_generated_decor(
                    new_slide,
                    ms.get("decor", []),
                    prs.slide_width,
                    prs.slide_height,
                )
            )
            apply_notes(new_slide, ms.get("notes", ""))
            apply_slide_transition(new_slide, ms.get("transition"))
            desired.append(list(sldIdLst)[-1])
            animation_targets.append(targets)

    # rebuild the slide-id list to match the model (handles reorder + delete)
    for sid in list(sldIdLst):
        sldIdLst.remove(sid)
    for i, sid in enumerate(orig_sldIds):
        if i not in used_src:
            try:
                prs.part.drop_rel(sid.get(qn("r:id")))
            except Exception:
                pass
    for el in desired:
        sldIdLst.append(el)

    prs.save(out_path)
    apply_powerpoint_animations(out_path, model, animation_targets)
    apply_powerpoint_media(out_path, model)


def base_slides_len_slide(prs):
    """Return the slide object most recently added."""
    return list(prs.slides)[-1]


def render(model, out_path):
    _validate_model(model)
    if (model.get("mode") or "native") == "overlay":
        render_overlay(model, out_path)
    else:
        render_native(model, out_path)
    return out_path


if __name__ == "__main__":
    model_path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "data", "model.json")
    out_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(HERE), "presentation.pptx")
    if model_path == "-":
        # Read raw bytes and decode UTF-8 explicitly: on Windows, sys.stdin may
        # default to cp1252, which corrupts curly quotes/dashes into lone
        # surrogates and crashes lxml when saving.
        model = json.loads(sys.stdin.buffer.read().decode("utf-8"))
    else:
        with open(model_path, "r", encoding="utf-8") as f:
            model = json.load(f)
    render(model, out_path)
    print("rendered", out_path)
