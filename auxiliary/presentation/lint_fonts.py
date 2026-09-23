#!/usr/bin/env python
"""Font-family scanner for `ppt lint`.

model.json's overlay elements never carry a font name (overlay rendering only
ever writes text/color/bold/size/align, see render_pptx.py apply_style), so
the pinned house rule "one single font family across the whole deck" has
nothing to check in the model itself, only in the actual rendered pptx (each
run's font comes from base.pptx and is untouched by the renderer). This walks
presentation.pptx's real runs and reports every run with an explicit,
non-inherited font name, so server.js can flag whichever slides use a
  minority font, the same majority-wins pattern lintModel() already uses for
  header alignment. Paragraph/list-style defaults are resolved from OOXML.
  Runs still inherited from a placeholder/master/theme are reported as
  unresolved coverage, never silently omitted as though the scan were clean.

    python lint_fonts.py <presentation.pptx>

Prints {"runs": [...], "coverage": {"total": N, "resolved": N,
"unresolved": N, "complete": bool, "unresolvedRuns": [...]}}
as JSON on stdout. Slide index is 0-based and matches model.slides order
exactly (render_pptx.py rebuilds the output slide order to equal model order,
regardless of each slide's base.pptx source index).
"""
import json
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def _latin_typeface(parent):
    if parent is None:
        return None
    latin = parent.find(qn("a:latin"))
    if latin is None:
        return None
    value = latin.get("typeface")
    return value.strip() if value and value.strip() else None


def _effective_font_name(run, para):
    """Resolve the font levels python-pptx does not expose on Run.font.

    Full placeholder/master inheritance is not available through python-pptx,
    but paragraph defRPr and text-frame list-style defaults are common and
    deterministic OOXML levels. Anything still inherited is reported as
    unresolved coverage rather than silently omitted and mistaken for clean.
    """
    name = run.font.name
    if name:
        return (None, "theme-token") if name.startswith("+") else (name, "run")

    ppr = para._p.find(qn("a:pPr"))
    if ppr is not None:
        name = _latin_typeface(ppr.find(qn("a:defRPr")))
        if name:
            return (None, "theme-token") if name.startswith("+") else (name, "paragraph")

    txbody = para._p.getparent()
    if txbody is not None:
        list_style = txbody.find(qn("a:lstStyle"))
        if list_style is not None:
            level = getattr(para, "level", 0) + 1
            level_ppr = list_style.find(qn(f"a:lvl{level}pPr"))
            if level_ppr is not None:
                name = _latin_typeface(level_ppr.find(qn("a:defRPr")))
                if name:
                    return (None, "theme-token") if name.startswith("+") else (name, "list-style")
    return None, "inherited-unresolved"


def _runs_in_text_frame(tf, slide_idx, out, unresolved):
    for para in tf.paragraphs:
        for run in para.runs:
            text = run.text.strip()
            if not text:
                continue
            name, source = _effective_font_name(run, para)
            if name:
                out.append({"slide": slide_idx, "font": name, "text": text[:60],
                            "source": source})
            else:
                unresolved.append({"slide": slide_idx, "text": text[:60],
                                   "reason": source})


def _walk_shapes(shapes, slide_idx, out, unresolved=None):
    if unresolved is None:
        unresolved = []
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                _walk_shapes(shape.shapes, slide_idx, out, unresolved)
                continue
        except Exception:
            pass
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            _runs_in_text_frame(shape.text_frame, slide_idx, out, unresolved)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    _runs_in_text_frame(cell.text_frame, slide_idx, out, unresolved)
    return unresolved


def main(pptx_path):
    prs = Presentation(pptx_path)
    out = []
    unresolved = []
    for i, slide in enumerate(prs.slides):
        _walk_shapes(slide.shapes, i, out, unresolved)
    total = len(out) + len(unresolved)
    print(json.dumps({
        "runs": out,
        "coverage": {
            "total": total,
            "resolved": len(out),
            "unresolved": len(unresolved),
            "complete": not unresolved,
            "unresolvedRuns": unresolved,
        },
    }))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: python lint_fonts.py <presentation.pptx>", file=sys.stderr)
        sys.exit(2)
    main(sys.argv[1])
