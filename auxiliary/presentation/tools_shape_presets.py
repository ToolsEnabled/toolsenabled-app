#!/usr/bin/env python
"""Startup migration helper for imported box/card shape metadata.

Older overlay models predate ``sourcePreset``. This reads the immutable
base.pptx through import_pptx.py's canonical classifier and prints only
eligible, non-placeholder rectangle-family AutoShapes keyed by slide and
PowerPoint shape id:

    python tools_shape_presets.py <base.pptx>

Output:
    {"<slide index>": {"<shape id>": "<preset>", ...}, ...}
"""
import json
import sys

from pptx import Presentation

from import_pptx import box_shape_preset


def main(base_path):
    presentation = Presentation(base_path)
    result = {}
    for slide_index, slide in enumerate(presentation.slides):
        shapes = {}
        for shape in slide.shapes:
            preset = box_shape_preset(shape)
            if preset:
                shapes[str(shape.shape_id)] = preset
        result[str(slide_index)] = shapes
    print(json.dumps(result, sort_keys=True))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(
            "usage: python tools_shape_presets.py <base.pptx>",
            file=sys.stderr,
        )
        sys.exit(2)
    main(sys.argv[1])
