#!/usr/bin/env python
"""Isolated regressions for the PPTX import/render/lint helper programs."""
import base64
import contextlib
import io
import json
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

sys.dont_write_bytecode = True

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

import import_pptx as importer
import lint_fonts
import render_pptx as renderer
import tools_decor_sids
import tools_shape_presets

PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
XML_NS = {"p": PML_NS, "mc": MC_NS, "p14": P14_NS}
H264_FIXTURE = (
    Path(__file__).parent
    / "test-fixtures"
    / "h264-blue-32x32.mp4.base64"
)


def xml_name(namespace, local_name):
    return f"{{{namespace}}}{local_name}"


def blank_deck(path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(path)
    return path


def destinations(root):
    return (
        root / "data" / "base.pptx",
        root / "data" / "model.json",
        root / "presentation.pptx",
    )


def run_summary(shape):
    result = []
    for para in shape.text_frame.paragraphs:
        runs = []
        for run in para.runs:
            try:
                color = None if run.font.color.type is None else str(run.font.color.rgb)
            except Exception:
                color = "non-rgb"
            runs.append((
                run.text,
                None if run.font.size is None else run.font.size.pt,
                run.font.bold,
                color,
                run.font.name,
            ))
        result.append((para.text, para.alignment, runs))
    return result


def slide_xml(path, number):
    with zipfile.ZipFile(path) as package:
        return etree.fromstring(
            package.read(f"ppt/slides/slide{number}.xml")
        )


def insert_before_timing(slide, element):
    root = slide._element
    anchors = {
        xml_name(PML_NS, "timing"),
        xml_name(PML_NS, "extLst"),
    }
    index = next(
        (i for i, child in enumerate(root) if child.tag in anchors),
        len(root),
    )
    root.insert(index, element)


def direct_transition(effect="push", with_sound=False):
    transition = etree.Element(xml_name(PML_NS, "transition"))
    transition.set("spd", "slow")
    transition.set("advClick", "0")
    transition.set("advTm", "777")
    etree.SubElement(transition, xml_name(PML_NS, effect))
    if with_sound:
        sound_action = etree.SubElement(
            transition, xml_name(PML_NS, "sndAc")
        )
        etree.SubElement(sound_action, xml_name(PML_NS, "endSnd"))
    return transition


def timing_alternate_content():
    alternate = etree.Element(
        xml_name(MC_NS, "AlternateContent"),
        nsmap={"mc": MC_NS, "p14": P14_NS},
    )
    choice = etree.SubElement(alternate, xml_name(MC_NS, "Choice"))
    choice.set("Requires", "p14")
    etree.SubElement(choice, xml_name(PML_NS, "timing"))
    fallback = etree.SubElement(alternate, xml_name(MC_NS, "Fallback"))
    etree.SubElement(fallback, xml_name(PML_NS, "timing"))
    return alternate


class ImportSafetyTests(unittest.TestCase):
    def test_corrupt_source_does_not_replace_any_destination(self):
        with tempfile.TemporaryDirectory(prefix="suite-import-invalid-") as td:
            root = Path(td)
            base, model, working = destinations(root)
            base.parent.mkdir(parents=True)
            base.write_bytes(b"old base")
            working.write_bytes(b"old working")
            model.write_bytes(b'{"old":true}')
            bad = root / "bad.pptx"
            bad.write_bytes(b"not a package")

            with self.assertRaises(Exception):
                importer.import_pptx(bad, base, model, working)

            self.assertEqual(base.read_bytes(), b"old base")
            self.assertEqual(working.read_bytes(), b"old working")
            self.assertEqual(model.read_bytes(), b'{"old":true}')

    def test_commit_failure_rolls_all_destinations_back(self):
        with tempfile.TemporaryDirectory(prefix="suite-import-rollback-") as td:
            root = Path(td)
            source = blank_deck(root / "source.pptx")
            base, model, working = destinations(root)
            base.parent.mkdir(parents=True)
            base.write_bytes(b"old base")
            working.write_bytes(b"old working")
            model.write_bytes(b"old model")
            real_replace = os.replace
            failed = False

            def replace_once(src, dst):
                nonlocal failed
                if (not failed and Path(dst) == model
                        and ".import-" in Path(src).name):
                    failed = True
                    raise OSError("forced model commit failure")
                return real_replace(src, dst)

            with mock.patch.object(importer.os, "replace", side_effect=replace_once):
                with self.assertRaisesRegex(OSError, "forced model"):
                    importer.import_pptx(source, base, model, working)

            self.assertTrue(failed)
            self.assertEqual(base.read_bytes(), b"old base")
            self.assertEqual(working.read_bytes(), b"old working")
            self.assertEqual(model.read_bytes(), b"old model")

    def test_new_import_emits_decor_ids_and_shape_sids(self):
        with tempfile.TemporaryDirectory(prefix="suite-import-decor-") as td:
            root = Path(td)
            source = root / "theme.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_shape(
                1, Inches(1), Inches(1), Inches(2), Inches(1),
            )
            shape.text = ""
            shape.fill.solid()
            shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
            expected_sid = shape.shape_id
            prs.save(source)
            base, model_path, working = destinations(root)

            with contextlib.redirect_stdout(io.StringIO()):
                importer.import_pptx(source, base, model_path, working)
            model = json.loads(model_path.read_text(encoding="utf-8"))
            decor = model["slides"][0]["decor"]

            self.assertEqual(len(decor), 1)
            self.assertEqual(decor[0]["id"], "s1_d0")
            self.assertEqual(decor[0]["sid"], expected_sid)
            self.assertEqual(decor[0]["fill"], "#4F81BD")

            capture = io.StringIO()
            with contextlib.redirect_stdout(capture):
                tools_decor_sids.main(str(base))
            self.assertEqual(json.loads(capture.getvalue())["0"], [expected_sid])

    def test_import_marks_only_non_placeholder_rectangle_family_autoshapes(self):
        with tempfile.TemporaryDirectory(prefix="suite-import-box-presets-") as td:
            root = Path(td)
            source = root / "cards.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            rounded = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(1), Inches(1), Inches(3), Inches(1),
            )
            rounded.text = "Rounded text card"
            rectangle = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(1), Inches(2.5), Inches(3), Inches(1),
            )
            rectangle.fill.solid()
            rectangle.fill.fore_color.rgb = RGBColor(17, 34, 51)
            oval = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(5), Inches(1), Inches(1), Inches(1),
            )
            oval.fill.solid()
            oval.fill.fore_color.rgb = RGBColor(68, 85, 102)
            textbox = slide.shapes.add_textbox(
                Inches(5), Inches(2.5), Inches(2), Inches(1),
            )
            textbox.text = "Not a card AutoShape"
            title_slide = prs.slides.add_slide(prs.slide_layouts[0])
            title_slide.shapes.title.text = "Placeholder, not a card"
            placeholder_sid = title_slide.shapes.title.shape_id
            prs.save(source)
            base, model_path, working = destinations(root)

            with contextlib.redirect_stdout(io.StringIO()):
                importer.import_pptx(source, base, model_path, working)
            model = json.loads(model_path.read_text(encoding="utf-8"))
            imported = model["slides"][0]
            by_sid = {}
            for item in imported["elements"] + imported["decor"]:
                sid = item.get("sid")
                if sid is None and isinstance(item.get("ref"), dict):
                    sid = item["ref"].get("sid")
                if sid is not None:
                    by_sid[sid] = item

            self.assertEqual(model["shapePresetSchema"], 1)
            self.assertEqual(
                by_sid[rounded.shape_id]["sourcePreset"], "roundRect"
            )
            self.assertEqual(
                by_sid[rectangle.shape_id]["sourcePreset"], "rect"
            )
            self.assertNotIn("sourcePreset", by_sid[oval.shape_id])
            self.assertNotIn("sourcePreset", by_sid[textbox.shape_id])
            placeholder = next(
                item for item in model["slides"][1]["elements"]
                if item.get("ref", {}).get("sid") == placeholder_sid
            )
            self.assertNotIn("sourcePreset", placeholder)

            capture = io.StringIO()
            with contextlib.redirect_stdout(capture):
                tools_shape_presets.main(str(base))
            presets = json.loads(capture.getvalue())["0"]
            self.assertEqual(
                presets,
                {
                    str(rounded.shape_id): "roundRect",
                    str(rectangle.shape_id): "rect",
                },
            )
            capture = io.StringIO()
            with contextlib.redirect_stdout(capture):
                tools_shape_presets.main(str(base))
            self.assertEqual(json.loads(capture.getvalue())["1"], {})


class PowerPointOwnershipTests(unittest.TestCase):
    def test_dispatch_owns_only_the_exact_new_application_pid(self):
        application = mock.Mock(name="powerpoint-application")
        dispatch = mock.Mock(return_value=application)
        with (
            mock.patch.object(
                renderer, "_current_powerpoint_pids",
                side_effect=({41}, {41, 73}),
            ),
            mock.patch.object(
                renderer, "_powerpoint_application_pid", return_value=73,
            ),
        ):
            actual, owns = renderer._dispatch_powerpoint_application(dispatch)

        self.assertIs(actual, application)
        self.assertTrue(owns)
        dispatch.assert_called_once_with("PowerPoint.Application")

    def test_dispatch_does_not_own_an_attached_existing_application(self):
        application = mock.Mock(name="powerpoint-application")
        with (
            mock.patch.object(
                renderer, "_current_powerpoint_pids",
                side_effect=({41}, {41}),
            ),
            mock.patch.object(
                renderer, "_powerpoint_application_pid", return_value=41,
            ),
        ):
            _, owns = renderer._dispatch_powerpoint_application(
                mock.Mock(return_value=application)
            )

        self.assertFalse(owns)

    def test_dispatch_fails_safe_for_ambiguous_or_unreadable_pid_evidence(self):
        cases = [
            ((set(), {73, 74}), 73),
            (({41}, {41, 73}), 41),
            ((OSError("snapshot unavailable"), {73}), 73),
            ((set(), {73}), OSError("window PID unavailable")),
        ]
        for snapshots, application_pid in cases:
            with self.subTest(
                    snapshots=snapshots, application_pid=application_pid):
                application = mock.Mock(name="powerpoint-application")
                with (
                    mock.patch.object(
                        renderer, "_current_powerpoint_pids",
                        side_effect=snapshots,
                    ),
                    mock.patch.object(
                        renderer, "_powerpoint_application_pid",
                        side_effect=(
                            application_pid
                            if isinstance(application_pid, Exception)
                            else None
                        ),
                        return_value=(
                            None
                            if isinstance(application_pid, Exception)
                            else application_pid
                        ),
                    ),
                ):
                    _, owns = renderer._dispatch_powerpoint_application(
                        mock.Mock(return_value=application)
                    )
                self.assertFalse(owns)

    def test_quit_is_called_only_for_a_proven_owned_application(self):
        attached = mock.Mock(name="attached-application")
        owned = mock.Mock(name="owned-application")

        renderer._quit_powerpoint_if_owned(attached, False)
        renderer._quit_powerpoint_if_owned(owned, True)

        attached.Quit.assert_not_called()
        owned.Quit.assert_called_once_with()

    @unittest.skipUnless(os.name == "nt", "PowerPoint COM is Windows-only")
    def test_animation_postprocessor_closes_deck_but_not_attached_app(self):
        import pythoncom

        application = mock.Mock(name="attached-application")
        presentation = mock.Mock(name="presentation")
        application.Presentations.Open.return_value = presentation
        presentation.Slides.Item.side_effect = RuntimeError("forced failure")
        model = {
            "slides": [{
                "animations": [{
                    "targetId": "target-1",
                    "effect": "appear",
                    "trigger": "click",
                    "duration": 0.5,
                    "delay": 0,
                }],
            }],
        }
        with (
            mock.patch.object(
                renderer, "_dispatch_powerpoint_application",
                return_value=(application, False),
            ),
            mock.patch.object(
                renderer, "_quit_powerpoint_if_owned",
                wraps=renderer._quit_powerpoint_if_owned,
            ) as quit_if_owned,
            mock.patch.object(pythoncom, "CoInitialize"),
            mock.patch.object(pythoncom, "CoUninitialize"),
        ):
            with self.assertRaisesRegex(RuntimeError, "forced failure"):
                renderer.apply_powerpoint_animations(
                    "unused.pptx", model, [{"target-1": 7}]
                )

        presentation.Close.assert_called_once_with()
        quit_if_owned.assert_called_once_with(application, False)
        application.Quit.assert_not_called()

    @unittest.skipUnless(os.name == "nt", "PowerPoint COM is Windows-only")
    def test_media_postprocessor_closes_deck_but_not_attached_app(self):
        import pythoncom

        application = mock.Mock(name="attached-application")
        presentation = mock.Mock(name="presentation")
        presentation.Slides.Count = 0
        application.Presentations.Open.return_value = presentation
        media = {
            "id": "media-1",
            "kind": "media",
            "generated": True,
            "source": "assets/clip.mp4",
            "box": {"x": 0.1, "y": 0.1, "w": 0.2, "h": 0.2},
            "autoplay": True,
            "loop": True,
        }
        model = {"slides": [{"decor": [media]}]}
        with (
            mock.patch.object(
                renderer, "_resolve_video_source",
                return_value=Path("unused.mp4"),
            ),
            mock.patch.object(
                renderer, "_dispatch_powerpoint_application",
                return_value=(application, False),
            ),
            mock.patch.object(
                renderer, "_quit_powerpoint_if_owned",
                wraps=renderer._quit_powerpoint_if_owned,
            ) as quit_if_owned,
            mock.patch.object(pythoncom, "CoInitialize"),
            mock.patch.object(pythoncom, "CoUninitialize"),
        ):
            with self.assertRaisesRegex(
                    RuntimeError, "slide count does not match"):
                renderer.apply_powerpoint_media("unused.pptx", model)

        presentation.Close.assert_called_once_with()
        quit_if_owned.assert_called_once_with(application, False)
        application.Quit.assert_not_called()

    def test_native_media_z_order_preserves_model_order_and_skips_imported_slides(self):
        calls = []

        def fake_shape(name):
            shape = mock.Mock()
            shape.Name = name
            shape.ZOrder.side_effect = lambda command: calls.append(
                (name, command)
            )
            return shape

        class Shapes:
            def __init__(self, items):
                self.items = items
                self.Count = len(items)

            def Item(self, index):
                return self.items[index - 1]

        imported_shapes = Shapes([
            fake_shape("Suite media imported-media"),
        ])
        native_shapes = Shapes([
            fake_shape("Suite shape native-bg"),
            fake_shape("Suite media native-video"),
            fake_shape("Suite image native-art"),
            fake_shape("Suite element native-copy"),
        ])
        slides = mock.Mock()
        slides.Item.side_effect = lambda index: mock.Mock(
            Shapes=imported_shapes if index == 1 else native_shapes
        )
        presentation = mock.Mock(Slides=slides)
        model = {
            "slides": [
                {
                    "src": 0,
                    "decor": [{
                        "id": "imported-media", "kind": "media",
                        "generated": True,
                    }],
                },
                {
                    "decor": [
                        {
                            "id": "native-bg", "kind": "shape",
                            "generated": True,
                        },
                        {
                            "id": "native-video", "kind": "media",
                            "generated": True,
                        },
                        {
                            "id": "native-art", "kind": "pic",
                            "generated": True,
                        },
                    ],
                },
            ],
        }

        renderer._send_native_decor_behind_content(presentation, model)

        self.assertEqual(
            calls,
            [
                ("Suite image native-art", 1),
                ("Suite media native-video", 1),
                ("Suite shape native-bg", 1),
            ],
        )
        slides.Item.assert_called_once_with(2)


class OverlayRenderingTests(unittest.TestCase):
    def test_unchanged_rich_runs_survive_noop_render(self):
        with tempfile.TemporaryDirectory(prefix="suite-rich-noop-") as td:
            root = Path(td)
            source = root / "rich.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(6), Inches(2),
            )
            para = shape.text_frame.paragraphs[0]
            run = para.add_run()
            run.text = "small"
            run.font.size = Pt(12)
            run.font.bold = False
            run.font.color.rgb = RGBColor(255, 0, 0)
            run = para.add_run()
            run.text = "BIG"
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 255)
            para = shape.text_frame.add_paragraph()
            para.alignment = PP_ALIGN.RIGHT
            run = para.add_run()
            run.text = "second"
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0, 128, 0)
            prs.save(source)
            base, model_path, working = destinations(root)
            with contextlib.redirect_stdout(io.StringIO()):
                importer.import_pptx(source, base, model_path, working)
            model = json.loads(model_path.read_text(encoding="utf-8"))
            output = root / "rendered.pptx"

            renderer.render(model, output)

            before = Presentation(source).slides[0].shapes[0]
            after = Presentation(output).slides[0].shapes[0]
            self.assertEqual(run_summary(after), run_summary(before))

    def test_actual_overlay_deltas_still_apply(self):
        with tempfile.TemporaryDirectory(prefix="suite-overlay-delta-") as td:
            root = Path(td)
            source = root / "source.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(4), Inches(1),
            )
            shape.text = "Original"
            run = shape.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(255, 0, 0)
            prs.save(source)
            base, model_path, working = destinations(root)
            with contextlib.redirect_stdout(io.StringIO()):
                importer.import_pptx(source, base, model_path, working)
            model = json.loads(model_path.read_text(encoding="utf-8"))
            element = model["slides"][0]["elements"][0]
            element["text"] = "Changed"
            element["style"].update({
                "color": "#0000FF", "size": 30, "bold": True,
                "font": "Arial", "align": "right", "fill": "#00FF00",
            })
            element["box"]["x"] = 0.25
            output = root / "rendered.pptx"

            renderer.render(model, output)

            rendered = Presentation(output)
            shape = rendered.slides[0].shapes[0]
            run = shape.text_frame.paragraphs[0].runs[0]
            self.assertEqual(shape.text, "Changed")
            self.assertEqual(run.font.size.pt, 30)
            self.assertTrue(run.font.bold)
            self.assertEqual(run.font.name, "Arial")
            self.assertEqual(str(run.font.color.rgb), "0000FF")
            self.assertEqual(shape.text_frame.paragraphs[0].alignment, PP_ALIGN.RIGHT)
            self.assertEqual(str(shape.fill.fore_color.rgb), "00FF00")
            self.assertAlmostEqual(shape.left / rendered.slide_width, 0.25, places=5)

    @unittest.skipUnless(os.name == "nt", "native PowerPoint builds require Windows")
    def test_sharp_corner_roundtrip_preserves_shape_style_order_and_animation_identity(self):
        with tempfile.TemporaryDirectory(prefix="suite-sharp-corners-") as td:
            root = Path(td)
            source = root / "cards.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            anchor = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(0.4), Inches(0.4), Inches(0.5), Inches(0.5),
            )
            anchor.name = "Z order anchor"
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(1.25), Inches(1.5), Inches(4.5), Inches(1.4),
            )
            card.name = "Animated text card"
            card.text = "Behavioral correctness"
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(17, 34, 51)
            card.line.color.rgb = RGBColor(68, 85, 102)
            card.line.width = Pt(2.25)
            card.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
            card_sid = card.shape_id
            card_box = (card.left, card.top, card.width, card.height)
            decor = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(6.2), Inches(1.5), Inches(1.5), Inches(1.4),
            )
            decor.name = "Rounded decor card"
            decor.fill.solid()
            decor.fill.fore_color.rgb = RGBColor(241, 184, 45)
            decor_sid = decor.shape_id
            prs.save(source)

            base, model_path, working = destinations(root)
            with contextlib.redirect_stdout(io.StringIO()):
                importer.import_pptx(source, base, model_path, working)
            model = json.loads(model_path.read_text(encoding="utf-8"))
            model_slide = model["slides"][0]
            element = next(
                item for item in model_slide["elements"]
                if item.get("ref", {}).get("sid") == card_sid
            )
            decor_item = next(
                item for item in model_slide["decor"]
                if item.get("sid") == decor_sid
            )
            self.assertEqual(element["sourcePreset"], "roundRect")
            self.assertEqual(decor_item["sourcePreset"], "roundRect")
            element["corners"] = "sharp"
            decor_item["corners"] = "sharp"
            model_slide["animations"] = [{
                "targetId": element["id"],
                "effect": "fade",
                "trigger": "click",
                "duration": 0.45,
                "delay": 0,
            }]

            output = root / "rendered.pptx"
            renderer.render(model, output)
            roundtrip = root / "roundtrip.pptx"
            Presentation(output).save(roundtrip)

            for candidate in (output, roundtrip):
                rendered = Presentation(candidate)
                shapes = rendered.slides[0].shapes
                self.assertEqual(
                    [shape.name for shape in shapes],
                    ["Z order anchor", "Animated text card", "Rounded decor card"],
                )
                rendered_card = next(
                    shape for shape in shapes if shape.shape_id == card_sid
                )
                rendered_decor = next(
                    shape for shape in shapes if shape.shape_id == decor_sid
                )
                self.assertEqual(
                    rendered_card.auto_shape_type, MSO_SHAPE.RECTANGLE
                )
                self.assertEqual(
                    rendered_decor.auto_shape_type, MSO_SHAPE.RECTANGLE
                )
                self.assertEqual(rendered_card.text, "Behavioral correctness")
                self.assertEqual(
                    (
                        rendered_card.left,
                        rendered_card.top,
                        rendered_card.width,
                        rendered_card.height,
                    ),
                    card_box,
                )
                self.assertEqual(
                    str(rendered_card.fill.fore_color.rgb), "112233"
                )
                self.assertEqual(
                    str(rendered_card.line.color.rgb), "445566"
                )
                self.assertAlmostEqual(
                    rendered_card.line.width.pt, 2.25, places=2
                )
                root_xml = slide_xml(candidate, 1)
                geometry = root_xml.xpath(
                    ".//p:sp[p:nvSpPr/p:cNvPr[@id=$sid]]"
                    "/p:spPr/a:prstGeom",
                    namespaces={
                        "p": PML_NS,
                        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                    },
                    sid=str(card_sid),
                )
                self.assertEqual(
                    [node.get("prst") for node in geometry], ["rect"]
                )
                targets = root_xml.findall(
                    f".//{xml_name(PML_NS, 'spTgt')}"
                )
                self.assertTrue(
                    any(target.get("spid") == str(card_sid) for target in targets)
                )

            # A second render starts from the same pristine rounded base and
            # reaches the identical sharp result without accumulating shapes.
            second = root / "second-render.pptx"
            renderer.render(model, second)
            second_shapes = Presentation(second).slides[0].shapes
            self.assertEqual(len(second_shapes), 3)
            self.assertEqual(
                next(
                    shape for shape in second_shapes
                    if shape.shape_id == card_sid
                ).auto_shape_type,
                MSO_SHAPE.RECTANGLE,
            )

    def test_sharp_corner_renderer_rejects_forged_backing_geometry(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        oval = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(1), Inches(1), Inches(2), Inches(1),
        )
        with self.assertRaisesRegex(ValueError, "backing geometry mismatch"):
            renderer.apply_corners(
                oval,
                {"sourcePreset": "roundRect", "corners": "sharp"},
                "forged target",
            )

    def test_generated_rectangle_renders_idempotently_and_tracks_model_edits(self):
        with tempfile.TemporaryDirectory(prefix="suite-generated-rect-") as td:
            root = Path(td)
            base = blank_deck(root / "base.pptx")
            decor = {
                "id": "generated-1",
                "kind": "shape",
                "sid": None,
                "generated": True,
                "shapeType": "rect",
                "box": {"x": 0.1, "y": 0.2, "w": 0.3, "h": 0.4},
                "fill": "#112233",
            }
            model = {
                "mode": "overlay",
                "base": str(base),
                "slides": [{
                    "id": "s1", "src": 0, "layout": "blank",
                    "elements": [], "decor": [decor],
                }],
            }

            first = root / "first.pptx"
            second = root / "second.pptx"
            renderer.render(model, first)
            renderer.render(model, second)

            for output in (first, second):
                rendered = Presentation(output)
                self.assertEqual(len(rendered.slides[0].shapes), 1)
                shape = rendered.slides[0].shapes[0]
                self.assertEqual(shape.auto_shape_type, MSO_SHAPE.RECTANGLE)
                self.assertEqual(str(shape.fill.fore_color.rgb), "112233")
                self.assertEqual(shape.line.fill.type, MSO_FILL.BACKGROUND)
                self.assertAlmostEqual(
                    shape.left / rendered.slide_width, 0.1, places=5
                )
                self.assertAlmostEqual(
                    shape.top / rendered.slide_height, 0.2, places=5
                )
                self.assertAlmostEqual(
                    shape.width / rendered.slide_width, 0.3, places=5
                )
                self.assertAlmostEqual(
                    shape.height / rendered.slide_height, 0.4, places=5
                )

            decor["fill"] = "#ABCDEF"
            decor["box"] = {"x": 0.2, "y": 0.1, "w": 0.25, "h": 0.5}
            edited = root / "edited.pptx"
            renderer.render(model, edited)
            shape = Presentation(edited).slides[0].shapes[0]
            self.assertEqual(str(shape.fill.fore_color.rgb), "ABCDEF")
            self.assertAlmostEqual(
                shape.left / Presentation(edited).slide_width, 0.2, places=5
            )
            self.assertAlmostEqual(
                shape.width / Presentation(edited).slide_width, 0.25, places=5
            )

            decor["box"] = {
                "x": 0.5, "y": 0.5, "w": 5e-324, "h": 5e-324,
            }
            tiny = root / "tiny.pptx"
            renderer.render(model, tiny)
            tiny_shape = Presentation(tiny).slides[0].shapes[0]
            self.assertEqual(tiny_shape.width, 1)
            self.assertEqual(tiny_shape.height, 1)

            model["slides"][0]["decor"] = []
            deleted = root / "deleted.pptx"
            renderer.render(model, deleted)
            self.assertEqual(
                len(Presentation(deleted).slides[0].shapes), 0
            )

    @unittest.skipUnless(os.name == "nt", "native PowerPoint builds require Windows")
    def test_generated_picture_and_native_animation_render_into_pptx(self):
        with tempfile.TemporaryDirectory(prefix="suite-generated-picture-") as td:
            root = Path(td)
            assets = root / "assets"
            assets.mkdir()
            image_path = assets / "pixel.png"
            Image.new("RGBA", (32, 32), (38, 230, 255, 200)).save(image_path)
            base = blank_deck(root / "base.pptx")
            picture = {
                "id": "image-1",
                "kind": "pic",
                "sid": None,
                "generated": True,
                "source": "assets/pixel.png",
                "box": {"x": 0.1, "y": 0.2, "w": 0.3, "h": 0.4},
            }
            model = {
                "mode": "overlay",
                "base": str(base),
                "slides": [{
                    "id": "s1", "src": 0, "layout": "blank",
                    "elements": [], "decor": [picture],
                    "animations": [{
                        "targetId": "image-1",
                        "effect": "rise-up",
                        "trigger": "click",
                        "duration": 0.6,
                        "delay": 0,
                    }],
                }],
            }
            output = root / "rendered.pptx"
            with (
                mock.patch.object(renderer, "PROJECT_ROOT", root.resolve()),
                mock.patch.object(renderer, "PROJECT_ASSETS", assets.resolve()),
            ):
                renderer.render(model, output)

            rendered = Presentation(output)
            self.assertEqual(len(rendered.slides[0].shapes), 1)
            shape = rendered.slides[0].shapes[0]
            self.assertEqual(shape.shape_type, 13)  # MSO_SHAPE_TYPE.PICTURE
            self.assertAlmostEqual(
                shape.left / rendered.slide_width, 0.1, places=5
            )
            root_xml = slide_xml(output, 1)
            timing = root_xml.find(xml_name(PML_NS, "timing"))
            self.assertIsNotNone(timing)
            targets = timing.findall(
                f".//{xml_name(PML_NS, 'spTgt')}"
            )
            self.assertTrue(
                any(target.get("spid") == str(shape.shape_id) for target in targets)
            )
            effects = timing.findall(
                f".//{xml_name(PML_NS, 'animEffect')}"
            )
            self.assertTrue(effects)

    @unittest.skipUnless(os.name == "nt", "embedded MP4 media requires Windows")
    def test_generated_video_is_embedded_with_persistent_playback_settings(self):
        import pythoncom
        import win32com.client

        with tempfile.TemporaryDirectory(prefix="suite-generated-video-") as td:
            root = Path(td)
            assets = root / "assets"
            assets.mkdir()
            video_path = assets / "clip.mp4"
            video_path.write_bytes(base64.b64decode(H264_FIXTURE.read_text()))
            base = blank_deck(root / "base.pptx")
            media = {
                "id": "media-1",
                "kind": "media",
                "sid": None,
                "generated": True,
                "source": "assets/clip.mp4",
                "box": {"x": 0.15, "y": 0.2, "w": 0.4, "h": 0.3},
                "autoplay": True,
                "loop": True,
            }
            model = {
                "mode": "overlay",
                "base": str(base),
                "slides": [{
                    "id": "s1", "src": 0, "layout": "blank",
                    "elements": [], "decor": [media],
                }],
            }
            output = root / "rendered.pptx"
            with (
                mock.patch.object(renderer, "PROJECT_ROOT", root.resolve()),
                mock.patch.object(renderer, "PROJECT_ASSETS", assets.resolve()),
            ):
                renderer.render(model, output)

            with zipfile.ZipFile(output) as package:
                embedded = [
                    name for name in package.namelist()
                    if name.lower().startswith("ppt/media/")
                    and name.lower().endswith(".mp4")
                ]
                self.assertTrue(embedded)
                self.assertIn(
                    video_path.read_bytes(),
                    [package.read(name) for name in embedded],
                )
                relationship_ns = (
                    "http://schemas.openxmlformats.org/package/2006/relationships"
                )
                media_relationships = []
                for name in package.namelist():
                    if not name.lower().endswith(".rels"):
                        continue
                    root_xml = etree.fromstring(package.read(name))
                    for relationship in root_xml.findall(
                            f"{{{relationship_ns}}}Relationship"):
                        rel_type = str(relationship.get("Type") or "").lower()
                        target = str(relationship.get("Target") or "")
                        if (
                            rel_type.endswith("/video")
                            or rel_type.endswith("/media")
                            or target.lower().endswith(".mp4")
                        ):
                            media_relationships.append(relationship)
                self.assertTrue(media_relationships)
                self.assertTrue(all(
                    str(relationship.get("TargetMode") or "").lower()
                    != "external"
                    for relationship in media_relationships
                ))

            application = None
            presentation = None
            pythoncom.CoInitialize()
            try:
                application = win32com.client.DispatchEx(
                    "PowerPoint.Application"
                )
                presentation = application.Presentations.Open(
                    str(output.resolve()), False, False, False
                )
                slide = presentation.Slides.Item(1)
                shape = renderer._find_com_shape_by_name(
                    slide, "Suite media media-1"
                )
                self.assertIsNotNone(shape)
                settings = shape.AnimationSettings.PlaySettings
                self.assertEqual(int(shape.Visible), -1)
                self.assertEqual(int(settings.HideWhileNotPlaying), 0)
                self.assertEqual(int(settings.PlayOnEntry), -1)
                self.assertEqual(int(settings.LoopUntilStopped), -1)
                self.assertAlmostEqual(
                    float(shape.Left) / float(presentation.PageSetup.SlideWidth),
                    0.15,
                    places=4,
                )
                self.assertAlmostEqual(
                    float(shape.Width) / float(presentation.PageSetup.SlideWidth),
                    0.4,
                    places=4,
                )
            finally:
                if presentation is not None:
                    presentation.Close()
                if application is not None:
                    application.Quit()
                pythoncom.CoUninitialize()

    @unittest.skipUnless(os.name == "nt", "embedded MP4 media requires Windows")
    def test_native_generated_video_embeds_in_model_order_behind_text(self):
        import pythoncom
        import win32com.client

        with tempfile.TemporaryDirectory(prefix="suite-native-video-") as td:
            root = Path(td)
            assets = root / "assets"
            assets.mkdir()
            (assets / "clip.mp4").write_bytes(
                base64.b64decode(H264_FIXTURE.read_text())
            )
            model = {
                "mode": "native",
                "slides": [{
                    "id": "s1",
                    "layout": "blank",
                    "elements": [{
                        "id": "e-copy",
                        "type": "body",
                        "text": "Text remains above generated media",
                    }],
                    "decor": [
                        {
                            "id": "d-bg",
                            "kind": "shape",
                            "sid": None,
                            "generated": True,
                            "shapeType": "rect",
                            "box": {
                                "x": 0, "y": 0, "w": 1, "h": 1,
                            },
                            "fill": "#003DA5",
                        },
                        {
                            "id": "d-video",
                            "kind": "media",
                            "sid": None,
                            "generated": True,
                            "source": "assets/clip.mp4",
                            "box": {
                                "x": 0.1, "y": 0.2,
                                "w": 0.4, "h": 0.3,
                            },
                            "autoplay": True,
                            "loop": True,
                        },
                    ],
                }],
            }
            output = root / "native-video.pptx"
            with (
                mock.patch.object(renderer, "PROJECT_ROOT", root.resolve()),
                mock.patch.object(
                    renderer, "PROJECT_ASSETS", assets.resolve()
                ),
            ):
                renderer.render(model, output)

            application = None
            presentation = None
            pythoncom.CoInitialize()
            try:
                application = win32com.client.DispatchEx(
                    "PowerPoint.Application"
                )
                presentation = application.Presentations.Open(
                    str(output.resolve()), False, False, False
                )
                slide = presentation.Slides.Item(1)
                background = renderer._find_com_shape_by_name(
                    slide, "Suite shape d-bg"
                )
                media = renderer._find_com_shape_by_name(
                    slide, "Suite media d-video"
                )
                text = renderer._find_com_shape_by_name(
                    slide, "Suite element e-copy"
                )
                self.assertIsNotNone(background)
                self.assertIsNotNone(media)
                self.assertIsNotNone(text)
                self.assertLess(
                    int(background.ZOrderPosition),
                    int(media.ZOrderPosition),
                )
                self.assertLess(
                    int(media.ZOrderPosition),
                    int(text.ZOrderPosition),
                )
                settings = media.AnimationSettings.PlaySettings
                self.assertEqual(int(settings.PlayOnEntry), -1)
                self.assertEqual(int(settings.LoopUntilStopped), -1)
            finally:
                if presentation is not None:
                    presentation.Close()
                if application is not None:
                    application.Quit()
                pythoncom.CoUninitialize()

    def test_malformed_generated_media_fails_before_rendering(self):
        with tempfile.TemporaryDirectory(prefix="suite-generated-media-invalid-") as td:
            root = Path(td)
            assets = root / "assets"
            assets.mkdir()
            video_path = assets / "clip.mp4"
            video_path.write_bytes(base64.b64decode(H264_FIXTURE.read_text()))
            (assets / "clip.mov").write_bytes(video_path.read_bytes())
            base = blank_deck(root / "base.pptx")
            valid = {
                "id": "media-1",
                "kind": "media",
                "sid": None,
                "generated": True,
                "source": "assets/clip.mp4",
                "box": {"x": 0.1, "y": 0.2, "w": 0.3, "h": 0.4},
                "autoplay": True,
                "loop": True,
            }
            variants = []

            def variant(name, update):
                decor = json.loads(json.dumps(valid))
                update(decor)
                variants.append((name, decor, None))

            variant("missing-source", lambda d: d.pop("source"))
            variant("absolute-source", lambda d: d.update(source=str(video_path)))
            variant("traversal-source", lambda d: d.update(source="assets/../clip.mp4"))
            variant("wrong-extension", lambda d: d.update(source="assets/clip.mov"))
            variant("missing-autoplay", lambda d: d.pop("autoplay"))
            variant("integer-autoplay", lambda d: d.update(autoplay=1))
            variant("string-loop", lambda d: d.update(loop="true"))
            variant("shape-type", lambda d: d.update(shapeType="rect"))
            variant("fill", lambda d: d.update(fill="#112233"))
            variant("not-generated", lambda d: d.pop("generated"))
            variants.append((
                "animation-target",
                json.loads(json.dumps(valid)),
                [{
                    "targetId": "media-1",
                    "effect": "fade",
                    "trigger": "click",
                    "duration": 0.5,
                    "delay": 0,
                }],
            ))

            with (
                mock.patch.object(renderer, "PROJECT_ROOT", root.resolve()),
                mock.patch.object(renderer, "PROJECT_ASSETS", assets.resolve()),
            ):
                for name, decor, animations in variants:
                    with self.subTest(name=name):
                        slide = {
                            "id": "s1", "src": 0, "layout": "blank",
                            "elements": [], "decor": [decor],
                        }
                        if animations is not None:
                            slide["animations"] = animations
                        model = {
                            "mode": "overlay",
                            "base": str(base),
                            "slides": [slide],
                        }
                        output = root / f"{name}.pptx"
                        with self.assertRaises((ValueError, FileNotFoundError)):
                            renderer.render(model, output)
                        self.assertFalse(output.exists())

    def test_malformed_generated_rectangles_fail_closed(self):
        with tempfile.TemporaryDirectory(prefix="suite-generated-invalid-") as td:
            root = Path(td)
            base = blank_deck(root / "base.pptx")
            valid = {
                "id": "generated-1",
                "kind": "shape",
                "sid": None,
                "generated": True,
                "shapeType": "rect",
                "box": {"x": 0.1, "y": 0.2, "w": 0.3, "h": 0.4},
                "fill": "#112233",
            }
            variants = []

            def variant(name, update):
                decor = json.loads(json.dumps(valid))
                update(decor)
                variants.append((name, decor))

            variant("unknown-shape", lambda d: d.update(shapeType="ellipse"))
            variant("missing-shape-type", lambda d: d.pop("shapeType"))
            variant("non-shape-kind", lambda d: d.update(kind="pic"))
            variant("backed-sid", lambda d: d.update(sid=7))
            variant("missing-sid", lambda d: d.pop("sid"))
            variant("missing-coordinate", lambda d: d["box"].pop("w"))
            variant("boolean-coordinate", lambda d: d["box"].update(x=True))
            variant("zero-width", lambda d: d["box"].update(w=0))
            variant("horizontal-overflow", lambda d: d["box"].update(x=0.8))
            variant("vertical-overflow", lambda d: d["box"].update(y=0.7))
            variant("invalid-fill", lambda d: d.update(fill="navy"))
            variant("false-generated", lambda d: d.update(generated=False))

            for name, decor in variants:
                with self.subTest(name=name):
                    model = {
                        "mode": "overlay",
                        "base": str(base),
                        "slides": [{
                            "id": "s1", "src": 0, "layout": "blank",
                            "elements": [], "decor": [decor],
                        }],
                    }
                    output = root / f"{name}.pptx"
                    with self.assertRaises(ValueError):
                        renderer.render(model, output)
                    self.assertFalse(output.exists())

            invalid_native_source = {
                "mode": "native",
                "slides": [{
                    "id": "s1", "src": 0, "layout": "blank",
                    "elements": [], "decor": [valid],
                }],
            }
            output = root / "native-slide-with-imported-source.pptx"
            with self.assertRaisesRegex(
                    ValueError, "suite-native or imported overlay slide"):
                renderer.render(invalid_native_source, output)
            self.assertFalse(output.exists())

    def test_deleted_imported_element_removes_its_backing_shape(self):
        with tempfile.TemporaryDirectory(prefix="suite-overlay-delete-") as td:
            root = Path(td)
            source = root / "source.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(4), Inches(1),
            )
            shape.text = "Delete me"
            deleted_sid = shape.shape_id
            prs.save(source)
            base, model_path, working = destinations(root)
            with contextlib.redirect_stdout(io.StringIO()):
                importer.import_pptx(source, base, model_path, working)
            model = json.loads(model_path.read_text(encoding="utf-8"))
            model["slides"][0]["elements"] = []
            output = root / "rendered.pptx"

            renderer.render(model, output)

            rendered_ids = {
                candidate.shape_id
                for candidate in Presentation(output).slides[0].shapes
            }
            self.assertNotIn(deleted_sid, rendered_ids)

    def test_duplicate_and_invalid_overlay_sources_are_rejected(self):
        with tempfile.TemporaryDirectory(prefix="suite-overlay-invalid-") as td:
            root = Path(td)
            base = blank_deck(root / "base.pptx")
            duplicate = {
                "mode": "overlay", "base": str(base),
                "slides": [
                    {"id": "s1", "src": 0, "elements": []},
                    {"id": "s2", "src": 0, "elements": []},
                ],
            }
            output = root / "duplicate.pptx"
            with self.assertRaisesRegex(ValueError, "duplicate overlay src"):
                renderer.render(duplicate, output)
            self.assertFalse(output.exists())

            for invalid in (True, 0.0, "0", -1, 1):
                with self.subTest(src=invalid):
                    model = {
                        "mode": "overlay", "base": str(base),
                        "slides": [{"id": "s1", "src": invalid, "elements": []}],
                    }
                    output = root / ("invalid-" + str(invalid) + ".pptx")
                    with self.assertRaises((ValueError, SystemExit)):
                        renderer.render(model, output)
                    self.assertFalse(output.exists())


class SlideTransitionTests(unittest.TestCase):
    def assert_no_transition(self, root):
        self.assertFalse(root.findall("./p:transition", XML_NS))
        for alternate in root.findall("./mc:AlternateContent", XML_NS):
            self.assertFalse(alternate.findall(".//p:transition", XML_NS))

    def assert_fade_transition(self, root, duration_ms, speed):
        self.assertFalse(root.findall("./p:transition", XML_NS))
        alternates = [
            child for child in root
            if (child.tag == xml_name(MC_NS, "AlternateContent")
                and child.findall(".//p:transition", XML_NS))
        ]
        self.assertEqual(len(alternates), 1)
        alternate = alternates[0]
        choice = alternate.find("./mc:Choice", XML_NS)
        fallback = alternate.find("./mc:Fallback", XML_NS)
        self.assertIsNotNone(choice)
        self.assertIsNotNone(fallback)
        self.assertEqual(choice.get("Requires"), "p14")
        self.assertEqual(choice.nsmap.get("p14"), P14_NS)

        precise = choice.find("./p:transition", XML_NS)
        legacy = fallback.find("./p:transition", XML_NS)
        self.assertIsNotNone(precise)
        self.assertIsNotNone(legacy)
        self.assertEqual(
            precise.attrib,
            {
                "spd": speed,
                "advClick": "1",
                xml_name(P14_NS, "dur"): str(duration_ms),
            },
        )
        self.assertEqual(
            legacy.attrib,
            {"spd": speed, "advClick": "1"},
        )
        for transition in (precise, legacy):
            self.assertIsNone(transition.get("advTm"))
            self.assertFalse(transition.findall("./p:sndAc", XML_NS))
            children = list(transition)
            self.assertEqual(len(children), 1)
            self.assertEqual(children[0].tag, xml_name(PML_NS, "fade"))
            self.assertEqual(children[0].attrib, {"thruBlk": "0"})

        alternate_index = list(root).index(alternate)
        for anchor_name in ("timing", "extLst"):
            anchor = root.find(f"./p:{anchor_name}", XML_NS)
            if anchor is not None:
                self.assertLess(alternate_index, list(root).index(anchor))
        for child in root:
            if (child.tag == xml_name(MC_NS, "AlternateContent")
                    and child is not alternate
                    and (child.findall(".//p:timing", XML_NS)
                         or child.findall(".//p:extLst", XML_NS))):
                self.assertLess(alternate_index, list(root).index(child))

    def test_transition_model_validation_is_strict_and_bounded(self):
        invalid = [
            None,
            "fade",
            {},
            {"effect": "wipe", "duration": 0.35},
            {"effect": "fade"},
            {"effect": "fade", "duration": True},
            {"effect": "fade", "duration": "0.35"},
            {"effect": "fade", "duration": float("nan")},
            {"effect": "fade", "duration": float("inf")},
            {"effect": "fade", "duration": 0.099},
            {"effect": "fade", "duration": 10.001},
            {"effect": "none", "duration": 0.35},
            {"effect": "none", "extra": 1},
            {"effect": "fade", "duration": 0.35, "speed": "fast"},
        ]
        with tempfile.TemporaryDirectory(prefix="suite-transition-invalid-") as td:
            root = Path(td)
            for index, transition in enumerate(invalid):
                with self.subTest(transition=transition):
                    output = root / f"invalid-{index}.pptx"
                    model = {
                        "mode": "native",
                        "slides": [{
                            "id": "s1",
                            "layout": "blank",
                            "elements": [],
                            "transition": transition,
                        }],
                    }
                    with self.assertRaises(ValueError):
                        renderer.render(model, output)
                    self.assertFalse(output.exists())

    def test_native_fades_emit_precise_and_legacy_ooxml(self):
        cases = [
            (0.1, 100, "fast"),
            (0.35, 350, "fast"),
            (0.75, 750, "med"),
            (2, 2000, "slow"),
            (10, 10000, "slow"),
        ]
        with tempfile.TemporaryDirectory(prefix="suite-transition-native-") as td:
            root = Path(td)
            output = root / "native.pptx"
            slides = [
                {
                    "id": f"s{index}",
                    "layout": "blank",
                    "elements": [],
                    "transition": {"effect": "fade", "duration": duration},
                }
                for index, (duration, _, _) in enumerate(cases, 1)
            ]
            slides.extend([
                {
                    "id": "none",
                    "layout": "blank",
                    "elements": [],
                    "transition": {"effect": "none"},
                },
                {"id": "absent", "layout": "blank", "elements": []},
            ])
            renderer.render({"mode": "native", "slides": slides}, output)

            for index, (_, duration_ms, speed) in enumerate(cases, 1):
                self.assert_fade_transition(
                    slide_xml(output, index), duration_ms, speed
                )
            self.assert_no_transition(slide_xml(output, len(cases) + 1))
            self.assert_no_transition(slide_xml(output, len(cases) + 2))

            # python-pptx does not expose transitions in its public API. It
            # must nevertheless preserve the raw MC/p14 XML on load/save.
            roundtrip = root / "roundtrip.pptx"
            reloaded = Presentation(output)
            renderer.apply_slide_transition(
                reloaded.slides[1], {"effect": "fade", "duration": 0.35}
            )
            renderer.apply_slide_transition(
                reloaded.slides[1], {"effect": "fade", "duration": 0.35}
            )
            reloaded.save(roundtrip)
            self.assert_fade_transition(slide_xml(roundtrip, 2), 350, "fast")

    def test_overlay_preserves_absent_removes_none_and_replaces_fade(self):
        with tempfile.TemporaryDirectory(prefix="suite-transition-overlay-") as td:
            root = Path(td)
            base = root / "base.pptx"
            prs = Presentation()
            slides = [
                prs.slides.add_slide(prs.slide_layouts[6])
                for _ in range(3)
            ]
            insert_before_timing(
                slides[0], direct_transition("push", with_sound=True)
            )
            slides[1]._element.append(renderer._make_fade_transition(0.8))
            slides[1]._element.append(timing_alternate_content())
            insert_before_timing(
                slides[2], direct_transition("wipe", with_sound=True)
            )
            slides[2]._element.append(renderer._make_fade_transition(1.5))
            slides[2]._element.append(timing_alternate_content())
            prs.save(base)

            model = {
                "mode": "overlay",
                "base": str(base),
                "slides": [
                    {
                        "id": "preserve",
                        "src": 0,
                        "elements": [],
                        "decor": [],
                    },
                    {
                        "id": "remove",
                        "src": 1,
                        "elements": [],
                        "decor": [],
                        "transition": {"effect": "none"},
                    },
                    {
                        "id": "replace",
                        "src": 2,
                        "elements": [],
                        "decor": [],
                        "transition": {"effect": "fade", "duration": 0.35},
                    },
                    {
                        "id": "new",
                        "layout": "blank",
                        "elements": [],
                        "decor": [],
                        "transition": {"effect": "fade", "duration": 0.75},
                    },
                ],
            }
            output = root / "rendered.pptx"
            renderer.render(model, output)

            preserved = slide_xml(output, 1)
            original = preserved.find("./p:transition", XML_NS)
            self.assertIsNotNone(original)
            self.assertEqual(original.get("advClick"), "0")
            self.assertEqual(original.get("advTm"), "777")
            self.assertIsNotNone(original.find("./p:push", XML_NS))
            self.assertIsNotNone(original.find("./p:sndAc", XML_NS))

            removed = slide_xml(output, 2)
            self.assert_no_transition(removed)
            removed_timing_wrappers = [
                child for child in removed
                if (child.tag == xml_name(MC_NS, "AlternateContent")
                    and child.findall(".//p:timing", XML_NS))
            ]
            self.assertEqual(len(removed_timing_wrappers), 1)

            replaced = slide_xml(output, 3)
            self.assert_fade_transition(replaced, 350, "fast")
            self.assertFalse(replaced.findall(".//p:sndAc", XML_NS))
            replaced_timing_wrappers = [
                child for child in replaced
                if (child.tag == xml_name(MC_NS, "AlternateContent")
                    and child.findall(".//p:timing", XML_NS))
            ]
            self.assertEqual(len(replaced_timing_wrappers), 1)
            self.assert_fade_transition(slide_xml(output, 4), 750, "med")

            # A second overlay pass must remain idempotent and must not grow
            # duplicate AlternateContent transition blocks.
            second_model = json.loads(json.dumps(model))
            second_model["base"] = str(output)
            second_model["slides"][3]["src"] = 3
            second = root / "rendered-twice.pptx"
            renderer.render(second_model, second)
            self.assert_fade_transition(slide_xml(second, 3), 350, "fast")
            self.assert_fade_transition(slide_xml(second, 4), 750, "med")


class NativeAndFontTests(unittest.TestCase):
    def test_generated_decor_renders_on_both_native_slide_paths(self):
        with tempfile.TemporaryDirectory(prefix="suite-native-decor-") as td:
            root = Path(td)
            assets = root / "assets"
            assets.mkdir()
            Image.new(
                "RGBA", (32, 32), (42, 120, 214, 255)
            ).save(assets / "art.png")
            (assets / "clip.mp4").write_bytes(
                base64.b64decode(H264_FIXTURE.read_text())
            )
            base = blank_deck(root / "base.pptx")

            for mode in ("native", "overlay"):
                with self.subTest(mode=mode):
                    slide = {
                        "id": "s1",
                        "layout": "blank",
                        "elements": [{
                            "id": "e-copy",
                            "type": "body",
                            "text": "Native copy stays above generated decor",
                        }],
                        "decor": [
                            {
                                "id": "d-bg",
                                "kind": "shape",
                                "sid": None,
                                "generated": True,
                                "shapeType": "rect",
                                "box": {
                                    "x": 0, "y": 0, "w": 1, "h": 1,
                                },
                                "fill": "#003DA5",
                            },
                            {
                                "id": "d-art",
                                "kind": "pic",
                                "sid": None,
                                "generated": True,
                                "source": "assets/art.png",
                                "box": {
                                    "x": 0.55, "y": 0.15,
                                    "w": 0.3, "h": 0.4,
                                },
                            },
                            {
                                "id": "d-video",
                                "kind": "media",
                                "sid": None,
                                "generated": True,
                                "source": "assets/clip.mp4",
                                "box": {
                                    "x": 0.1, "y": 0.6,
                                    "w": 0.3, "h": 0.2,
                                },
                                "autoplay": False,
                                "loop": True,
                            },
                            {
                                "id": "d-rail",
                                "kind": "shape",
                                "sid": None,
                                "generated": True,
                                "shapeType": "rect",
                                "box": {
                                    "x": 0.05, "y": 0.9,
                                    "w": 0.9, "h": 0.03,
                                },
                                "fill": "#F1B82D",
                            },
                        ],
                        "animations": [
                            {
                                "targetId": "d-art",
                                "effect": "fade",
                                "trigger": "click",
                                "duration": 0.4,
                                "delay": 0,
                            },
                            {
                                "targetId": "e-copy",
                                "effect": "rise-up",
                                "trigger": "after-previous",
                                "duration": 0.4,
                                "delay": 0.1,
                            },
                        ],
                    }
                    model = {"mode": mode, "slides": [slide]}
                    if mode == "overlay":
                        model["base"] = str(base)
                    output = root / f"{mode}.pptx"

                    with (
                        mock.patch.object(
                            renderer, "PROJECT_ROOT", root.resolve()
                        ),
                        mock.patch.object(
                            renderer, "PROJECT_ASSETS", assets.resolve()
                        ),
                        mock.patch.object(
                            renderer, "apply_powerpoint_animations"
                        ) as animations,
                        mock.patch.object(
                            renderer, "apply_powerpoint_media"
                        ) as media,
                    ):
                        renderer.render(model, output)

                    rendered = Presentation(output)
                    names = [
                        shape.name for shape in rendered.slides[0].shapes
                    ]
                    self.assertEqual(
                        names,
                        [
                            "Suite shape d-bg",
                            "Suite image d-art",
                            "Suite shape d-rail",
                            "Suite element e-copy",
                        ],
                    )
                    targets = animations.call_args.args[2]
                    self.assertEqual(
                        set(targets[0]),
                        {"d-bg", "d-art", "d-rail", "e-copy"},
                    )
                    self.assertNotIn("d-video", targets[0])
                    media.assert_called_once_with(output, model)

    def test_native_styles_fill_and_explicit_box_apply(self):
        with tempfile.TemporaryDirectory(prefix="suite-native-style-") as td:
            output = Path(td) / "native.pptx"
            model = {
                "mode": "native",
                "slides": [{
                    "id": "s1", "layout": "blank",
                    "elements": [{
                        "id": "e1", "type": "body", "text": "Styled",
                        "box": {"x": 0.3, "y": 0.4, "w": 0.2, "h": 0.1},
                        "style": {
                            "color": "#FF0000", "size": 40, "bold": True,
                            "font": "Arial", "align": "right", "fill": "#00FF00",
                            "outline": "#112233",
                        },
                    }],
                }],
            }

            renderer.render(model, output)

            prs = Presentation(output)
            shape = prs.slides[0].shapes[0]
            run = shape.text_frame.paragraphs[0].runs[0]
            self.assertEqual(run.font.size.pt, 40)
            self.assertTrue(run.font.bold)
            self.assertEqual(run.font.name, "Arial")
            self.assertEqual(str(run.font.color.rgb), "FF0000")
            self.assertEqual(str(shape.fill.fore_color.rgb), "00FF00")
            self.assertEqual(str(shape.line.fill.fore_color.rgb), "112233")
            self.assertEqual(shape.text_frame.paragraphs[0].alignment, PP_ALIGN.RIGHT)
            self.assertAlmostEqual(shape.left / prs.slide_width, 0.3, places=5)
            self.assertAlmostEqual(shape.top / prs.slide_height, 0.4, places=5)
            self.assertAlmostEqual(shape.width / prs.slide_width, 0.2, places=5)
            self.assertAlmostEqual(shape.height / prs.slide_height, 0.1, places=5)

    def test_native_placeholder_honors_explicit_box(self):
        with tempfile.TemporaryDirectory(prefix="suite-native-placeholder-box-") as td:
            output = Path(td) / "native-placeholder.pptx"
            explicit = {"x": 0.12, "y": 0.08, "w": 0.76, "h": 0.14}
            model = {
                "mode": "native",
                "slides": [{
                    "id": "s1", "layout": "content",
                    "elements": [{
                        "id": "e-title", "type": "title",
                        "text": "Placed title", "box": explicit,
                    }],
                }],
            }

            renderer.render(model, output)

            prs = Presentation(output)
            shape = next(
                item for item in prs.slides[0].shapes
                if item.name == "Suite element e-title"
            )
            self.assertAlmostEqual(shape.left / prs.slide_width, explicit["x"], places=5)
            self.assertAlmostEqual(shape.top / prs.slide_height, explicit["y"], places=5)
            self.assertAlmostEqual(shape.width / prs.slide_width, explicit["w"], places=5)
            self.assertAlmostEqual(shape.height / prs.slide_height, explicit["h"], places=5)

    def test_font_lint_resolves_paragraph_defaults_and_reports_unknowns(self):
        with tempfile.TemporaryDirectory(prefix="suite-font-lint-") as td:
            path = Path(td) / "fonts.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for index, font in enumerate(("Arial", "Calibri")):
                shape = slide.shapes.add_textbox(
                    Inches(1), Inches(1 + index), Inches(3), Inches(0.5),
                )
                para = shape.text_frame.paragraphs[0]
                para.text = "font " + font
                ppr = para._p.get_or_add_pPr()
                default = OxmlElement("a:defRPr")
                latin = OxmlElement("a:latin")
                latin.set("typeface", font)
                default.append(latin)
                ppr.append(default)
            slide.shapes.add_textbox(
                Inches(1), Inches(3), Inches(3), Inches(0.5),
            ).text = "theme inherited"
            prs.save(path)

            capture = io.StringIO()
            with contextlib.redirect_stdout(capture):
                lint_fonts.main(str(path))
            result = json.loads(capture.getvalue())

            self.assertEqual({run["font"] for run in result["runs"]},
                             {"Arial", "Calibri"})
            self.assertEqual(result["coverage"]["total"], 3)
            self.assertEqual(result["coverage"]["resolved"], 2)
            self.assertEqual(result["coverage"]["unresolved"], 1)
            self.assertFalse(result["coverage"]["complete"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
