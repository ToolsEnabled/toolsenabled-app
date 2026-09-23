"""Scribe document engine: run-aware, formatting-preserving .docx editing.

The one rule this file exists to enforce: never touch `paragraph.text = ...`.
python-docx implements that setter by clearing every run and writing one bare
run, which on the real corpus collapses a 13-run paragraph to 1 and destroys the
colored audit markup. Every mutation here works at run level and splits runs
rather than replacing them.

Addressing is by `w14:paraId`, Word's own 8-hex-digit stable paragraph id, which
is stamped on open if absent and survives insert, delete, and reorder. Indexes
are never used as identity.

Measured against LEAN-Bench_McNair_2026.docx: 333 paragraphs, 404 runs, worst
paragraph 13 runs with 8 rPr blocks, 10 paragraphs with a run boundary
mid-sentence.
"""

import copy
import difflib
import hashlib
import io
import json
import math
import os
import re
import secrets
import shutil
import tempfile
import time
import zipfile
from datetime import datetime

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from lxml import etree

try:
    import regex as timeout_regex
except ImportError:
    timeout_regex = None

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

UNDERLINE_VALUES = {
    "single", "words", "double", "thick", "dotted", "dottedHeavy",
    "dash", "dashedHeavy", "dashLong", "dashLongHeavy", "dotDash",
    "dashDotHeavy", "dotDotDash", "dashDotDotHeavy", "wave",
    "wavyHeavy", "wavyDouble",
}
HIGHLIGHT_VALUES = {
    "black", "blue", "cyan", "green", "magenta", "red", "yellow",
    "white", "darkBlue", "darkCyan", "darkGreen", "darkMagenta",
    "darkRed", "darkYellow", "darkGray", "lightGray",
}
MAX_DOCX_PARTS = 10000
MAX_DOCX_PART_BYTES = 64 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 200
MAX_FIND_QUERY_CHARS = 10000
REGEX_SEARCH_TIMEOUT_S = 1.0
VISUAL_RPR_TAGS = (
    "b", "bCs", "i", "iCs", "u", "color", "highlight", "sz", "szCs",
)
CLEARABLE_UNIFORM_DIRECT = {
    "b": ("b", "bCs"),
    "i": ("i", "iCs"),
    "u": ("u",),
    "size": ("sz", "szCs"),
}
OOXML_FALSE_VALUES = {"0", "false", "off", "no", "none"}


def file_sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as package:
        for chunk in iter(lambda: package.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def w(tag):
    return "{%s}%s" % (W, tag)


def w14(tag):
    return "{%s}%s" % (W14, tag)


class DocError(Exception):
    """A refused operation. The message is written for the agent to read and act on."""

    def __init__(self, message, **detail):
        super().__init__(message)
        self.message = message
        self.detail = detail


# --------------------------------------------------------------------------
# run-level primitives
# --------------------------------------------------------------------------

def run_children_ok(r):
    """True if a run holds only rPr and w:t.

    A run containing w:br, w:tab, w:drawing, or a footnote reference cannot be
    safely rebuilt by concatenating its text, because rebuilding would drop or
    reorder the non-text child. The corpus contains exactly one such run, so
    this is rare, but silently eating a line break is exactly the class of
    damage this engine exists to prevent. Refuse instead.
    """
    for ch in r:
        tag = ch.tag
        if tag != w("rPr") and tag != w("t"):
            return False
    return True


def run_text(r):
    return "".join(t.text or "" for t in r.findall(w("t")))


def validate_xml_text(value, name="text"):
    """Refuse characters that XML 1.0 cannot store before touching a run."""
    if not isinstance(value, str):
        raise DocError("%s must be a string" % name)
    for character in value:
        codepoint = ord(character)
        if not (
            codepoint in (0x09, 0x0A, 0x0D)
            or 0x20 <= codepoint <= 0xD7FF
            or 0xE000 <= codepoint <= 0xFFFD
            or 0x10000 <= codepoint <= 0x10FFFF
        ):
            raise DocError(
                "%s contains a character that cannot be stored in a Word document"
                % name
            )
    return value


def set_run_text(r, s):
    """Replace a run's text, preserving its rPr. Caller must have checked run_children_ok."""
    validate_xml_text(s)
    for t in r.findall(w("t")):
        r.remove(t)
    if s == "":
        return
    t = OxmlElement("w:t")
    t.text = s
    t.set(XML_SPACE, "preserve")
    r.append(t)


def clone_run_empty(r):
    """A new run carrying a deep copy of r's rPr and no text."""
    new = OxmlElement("w:r")
    pr = r.find(w("rPr"))
    if pr is not None:
        new.append(copy.deepcopy(pr))
    return new


def ensure_rpr(r):
    pr = r.find(w("rPr"))
    if pr is None:
        pr = OxmlElement("w:rPr")
        r.insert(0, pr)
    return pr


# --------------------------------------------------------------------------
# the document
# --------------------------------------------------------------------------

class ScribeDoc:
    def __init__(self, path):
        self.path = os.path.abspath(path)
        self.doc = None
        self.rev = 0
        self.backup_path = None
        self._style_names = {}
        self._paragraph_style_ids = set()
        self._default_paragraph_style_id = "Normal"
        self._stamped = 0
        self.package_sha256 = None
        # Paragraph ids are identities, not merely values that must be unique in
        # the current XML tree. Remember every id observed or minted during this
        # ScribeDoc session so deleting a paragraph cannot make its id available
        # for a later insert.
        self._issued_ids = set()

    # -- lifecycle ---------------------------------------------------------

    def _word_lock_path(self):
        name = os.path.basename(self.path)
        return os.path.join(os.path.dirname(self.path), "~$" + name[2:])

    def _check_word_lock(self):
        lock = self._word_lock_path()
        if os.path.exists(lock):
            raise DocError(
                "This document is open in Word (a ~$ lock file exists). "
                "Close it in Word and try again.",
                lock=lock,
                word_lock=True)

    def _check_expected_package(self, expected_sha256):
        if expected_sha256 is None:
            return
        if (not isinstance(expected_sha256, str) or
                not re.fullmatch(r"[0-9a-f]{64}", expected_sha256)):
            raise DocError("The expected document package fingerprint is invalid.")
        actual = file_sha256(self.path)
        if actual != expected_sha256:
            raise DocError(
                "The document package changed outside Scribe. Reopen it before editing.",
                package_changed=True,
                expected_package_sha256=expected_sha256,
                actual_package_sha256=actual)

    def open(self):
        self._check_word_lock()
        try:
            with open(self.path, "rb") as package:
                package_bytes = package.read()
        except OSError as error:
            raise DocError(
                "The file is not a readable Word document package.",
                error=str(error))
        self._validate_package(package_bytes)
        self.package_sha256 = hashlib.sha256(package_bytes).hexdigest()
        self.doc = docx.Document(io.BytesIO(package_bytes))
        self._style_names = {}
        self._paragraph_style_ids = set()
        self._default_paragraph_style_id = "Normal"
        for s in self.doc.styles:
            try:
                self._style_names[s.style_id] = s.name
                if s.type == WD_STYLE_TYPE.PARAGRAPH:
                    self._paragraph_style_ids.add(s.style_id)
                    default = s.element.get(w("default"))
                    if (default is not None
                            and default.casefold() not in OOXML_FALSE_VALUES):
                        self._default_paragraph_style_id = s.style_id
            except Exception:
                pass
        self._stamped = self._stamp_ids()
        self._issued_ids.update(self._all_ids())
        return self

    def _validate_package(self, package_bytes=None):
        """Reject corrupt, encrypted, duplicate, or expansion-bomb packages."""
        try:
            source = self.path if package_bytes is None else io.BytesIO(package_bytes)
            with zipfile.ZipFile(source) as archive:
                infos = archive.infolist()
                if len(infos) > MAX_DOCX_PARTS:
                    raise DocError(
                        "The document package has too many parts.",
                        parts=len(infos), limit=MAX_DOCX_PARTS)
                names = [info.filename for info in infos]
                if len(names) != len(set(names)):
                    raise DocError("The document package contains duplicate parts.")
                required = {"[Content_Types].xml", "word/document.xml"}
                missing = sorted(required - set(names))
                if missing:
                    raise DocError(
                        "The file is not a complete Word document package.",
                        missing=missing)

                total = 0
                for info in infos:
                    if info.flag_bits & 0x1:
                        raise DocError("Encrypted Word document packages are not supported.")
                    total += info.file_size
                    if info.file_size > MAX_DOCX_PART_BYTES:
                        raise DocError(
                            "A document package part is too large to open safely.",
                            part=info.filename, bytes=info.file_size,
                            limit=MAX_DOCX_PART_BYTES)
                    if (info.file_size > 1024 * 1024 and
                            info.file_size >
                            max(info.compress_size, 1) * MAX_DOCX_COMPRESSION_RATIO):
                        raise DocError(
                            "The document package has an unsafe compression ratio.",
                            part=info.filename)
                if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise DocError(
                        "The expanded document package is too large to open safely.",
                        bytes=total, limit=MAX_DOCX_UNCOMPRESSED_BYTES)
                corrupt = archive.testzip()
                if corrupt:
                    raise DocError(
                        "The document package contains a corrupt part.",
                        part=corrupt)
        except DocError:
            raise
        except (
                OSError, RuntimeError, ValueError, EOFError, NotImplementedError,
                zipfile.BadZipFile, zipfile.LargeZipFile) as error:
            raise DocError(
                "The file is not a readable Word document package.",
                error=str(error))

    def _body(self):
        return self.doc.element.body

    def paragraphs(self):
        """Every w:p in document order, including those nested inside tables.

        `Document.paragraphs` excludes table cells, which on sibling drafts hides
        up to 213 paragraphs including the per-prompt verdict table. Iterating
        the body catches everything.
        """
        return list(self._body().iter(w("p")))

    def _stamp_ids(self):
        """Give every paragraph a unique nonzero 8-hex-digit w14:paraId.

        Three cases need stamping: no id at all (python-docx generated files have
        none), the reserved 00000000, and a duplicate. Word does produce
        duplicate paraIds after a copy-paste, so the first holder keeps the id
        and later ones are re-stamped.
        """
        keep = set()
        needs = []
        for i, p in enumerate(self.paragraphs()):
            v = p.get(w14("paraId"))
            if v and v != "00000000" and v not in keep:
                keep.add(v)
            else:
                needs.append((i, p))
        for i, p in needs:
            p.set(w14("paraId"), self._fresh_id(keep, i, self.text_of(p)))
        return len(needs)

    def _fresh_id(self, seen, salt, text):
        h = hashlib.sha1(("%d\x00%s" % (salt, text)).encode("utf8")).hexdigest()[:8].upper()
        n = 0
        while h == "00000000" or h in seen or h in self._issued_ids:
            n += 1
            h = hashlib.sha1(("%d\x00%d\x00%s" % (salt, n, text)).encode("utf8")).hexdigest()[:8].upper()
        seen.add(h)
        self._issued_ids.add(h)
        return h

    def _all_ids(self):
        return {p.get(w14("paraId")) for p in self.paragraphs()}

    # -- reading -----------------------------------------------------------

    def text_of(self, p):
        return "".join(run_text(r) for r in p.findall(w("r")))

    def hash_of(self, p):
        return hashlib.sha1(self.text_of(p).encode("utf8")).hexdigest()[:12]

    def _direct_visual_props(self, rpr):
        """Canonical direct visual properties for a paragraph/run rPr.

        OOXML permits run-property elements in different orders even though
        their visual meaning is the same.  Hash a fixed property order and
        sorted attributes so save/reopen normalization cannot create a false
        formatting conflict.
        """
        if rpr is None:
            return []
        out = []
        for tag in VISUAL_RPR_TAGS:
            for element in rpr.findall(w(tag)):
                out.append([
                    tag,
                    sorted((name, value) for name, value in element.attrib.items()),
                ])
        return out

    def format_hash_of(self, p):
        """Formatting-sensitive paragraph guard, separate from the text hash."""
        paragraph_properties = p.findall(w("pPr"))
        payload = [
            "scribe-format-v2",
            self.style_of(p),
            [
                sorted((name, value) for name, value in pstyle.attrib.items())
                for ppr in paragraph_properties
                for pstyle in ppr.findall(w("pStyle"))
            ],
            [
                self._direct_visual_props(rpr)
                for ppr in paragraph_properties
                for rpr in ppr.findall(w("rPr"))
            ],
            [
                [
                    run_text(r),
                    [
                        self._direct_visual_props(rpr)
                        for rpr in r.findall(w("rPr"))
                    ],
                ]
                for r in p.findall(w("r"))
            ],
        ]
        encoded = json.dumps(
            payload, ensure_ascii=False, separators=(",", ":"),
        ).encode("utf8")
        return hashlib.sha256(encoded).hexdigest()[:16]

    def style_of(self, p):
        pr = p.find(w("pPr"))
        if pr is None:
            return self._style_names.get(
                self._default_paragraph_style_id, "Normal")
        st = pr.find(w("pStyle"))
        if st is None:
            return self._style_names.get(
                self._default_paragraph_style_id, "Normal")
        sid = st.get(w("val"))
        return self._style_names.get(sid, sid or "Normal")

    def in_table(self, p):
        return self.table_pos(p) is not None

    def table_pos(self, p):
        """{tbl, row, cell} for a paragraph inside a table, else None.

        The renderer needs this to rebuild the grid. Without it a table's cells
        arrive as a flat run of paragraphs and draw as stray lines, which on the
        sibling drafts would mean 183 to 213 paragraphs rendering wrong.
        Nested tables resolve against their innermost cell, which is what the
        renderer wants.
        """
        tc = p.getparent()
        while tc is not None and tc.tag != w("tc"):
            tc = tc.getparent()
        if tc is None:
            return None
        tr = tc.getparent()
        if tr is None or tr.tag != w("tr"):
            return None
        tbl = tr.getparent()
        if tbl is None or tbl.tag != w("tbl"):
            return None
        # Not cached deliberately: an id()-keyed cache can outlive the element
        # it describes and silently point at a recycled address. There are at
        # most 3 tables in this corpus, so recomputing costs nothing.
        tables = list(self._body().iter(w("tbl")))
        return {
            "tbl": next((i for i, t in enumerate(tables) if t is tbl), 0),
            "row": list(tbl.findall(w("tr"))).index(tr),
            "cell": list(tr.findall(w("tc"))).index(tc),
        }

    def para(self, pid):
        for p in self.paragraphs():
            if p.get(w14("paraId")) == pid:
                return p
        raise DocError("No paragraph with id %s. Call doc_read to get current ids." % pid, pid=pid)

    def _model_props_of_rpr(self, pr):
        """Visible direct properties from one rPr, including explicit off."""
        out = {}
        if pr is None:
            return out

        def toggle(name, tags):
            element = next(
                (pr.find(w(tag)) for tag in tags if pr.find(w(tag)) is not None),
                None,
            )
            if element is None:
                return
            value = (element.get(w("val")) or "true").casefold()
            out[name] = value not in OOXML_FALSE_VALUES

        toggle("b", ("b", "bCs"))
        toggle("i", ("i", "iCs"))
        u = pr.find(w("u"))
        if u is not None:
            value = u.get(w("val")) or "single"
            out["u"] = "" if value.casefold() in OOXML_FALSE_VALUES else value
        c = pr.find(w("color"))
        if c is not None and c.get(w("val")) not in (None, "auto"):
            out["color"] = c.get(w("val"))
        hl = pr.find(w("highlight"))
        if hl is not None:
            out["highlight"] = hl.get(w("val"))
        sz = pr.find(w("sz"))
        if sz is not None:
            try:
                out["size"] = int(sz.get(w("val"))) / 2.0
            except (TypeError, ValueError):
                pass
        return out

    def run_props(self, r):
        return self._model_props_of_rpr(r.find(w("rPr")))

    def para_model(self, p):
        runs = []
        for r in p.findall(w("r")):
            runs.append({"text": run_text(r), **self.run_props(r)})
        ppr = p.find(w("pPr"))
        paragraph_rpr = ppr.find(w("rPr")) if ppr is not None else None
        m = {
            "pid": p.get(w14("paraId")),
            "style": self.style_of(p),
            "text": self.text_of(p),
            "hash": self.hash_of(p),
            "format_hash": self.format_hash_of(p),
            "paragraph_format": self._model_props_of_rpr(paragraph_rpr),
            "runs": runs,
        }
        pos = self.table_pos(p)
        if pos:
            m["table"] = pos
        return m

    def model(self):
        ps = self.paragraphs()
        return {
            "path": self.path,
            "rev": self.rev,
            "stamped": self._stamped,
            "count": len(ps),
            "paragraphs": [self.para_model(p) for p in ps],
        }

    def read(self, start=0, end=None):
        ps = self.paragraphs()
        end = len(ps) - 1 if end is None else min(end, len(ps) - 1)
        start = max(0, start)
        return [self.para_model(p) for p in ps[start:end + 1]]

    def find(self, query, regex=False, limit=50):
        """Locate a string across the document. Returns pid plus char offsets."""
        if not isinstance(query, str) or query == "":
            raise DocError("query must be a non-empty string")
        if len(query) > MAX_FIND_QUERY_CHARS:
            raise DocError("query is too long")
        if not isinstance(regex, bool):
            raise DocError("regex must be a boolean")
        if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= 500:
            raise DocError("limit must be an integer from 1 to 500")
        hits = []
        if regex:
            if timeout_regex is None:
                raise DocError(
                    "Safe regex search is unavailable. Retry as a plain-text search.")
            try:
                pat = timeout_regex.compile(query)
            except timeout_regex.error as e:
                raise DocError("Bad regex: %s" % e, query=query)
            deadline = time.monotonic() + REGEX_SEARCH_TIMEOUT_S
        else:
            pat = re.compile(re.escape(query), re.IGNORECASE)
        for p in self.paragraphs():
            text = self.text_of(p)
            try:
                if regex:
                    remaining = deadline - time.monotonic()
                    if remaining <= 0:
                        raise DocError("Regex search timed out", query=query)
                    matches = pat.finditer(text, timeout=remaining)
                else:
                    matches = pat.finditer(text)
                for match in matches:
                    lo = max(0, match.start() - 40)
                    hits.append({
                        "pid": p.get(w14("paraId")),
                        "start": match.start(),
                        "end": match.end(),
                        "match": match.group(0),
                        "context": text[lo:match.end() + 40],
                        "style": self.style_of(p),
                        "hash": self.hash_of(p),
                        "format_hash": self.format_hash_of(p),
                    })
                    if len(hits) >= limit:
                        return hits
            except TimeoutError:
                raise DocError("Regex search timed out", query=query)
            if regex and time.monotonic() >= deadline:
                raise DocError("Regex search timed out", query=query)
        return hits

    # -- span mechanics ----------------------------------------------------

    def spans(self, p):
        """[(start, end, run, text)] over the paragraph's concatenated text."""
        out, pos = [], 0
        for r in p.findall(w("r")):
            t = run_text(r)
            out.append((pos, pos + len(t), r, t))
            pos += len(t)
        return out

    def _locate(self, p, find, occurrence=None):
        if not isinstance(find, str) or find == "":
            raise DocError("find must be a non-empty string")
        if occurrence is not None and (
                isinstance(occurrence, bool)
                or not isinstance(occurrence, int)
                or occurrence < 1):
            raise DocError("occurrence must be a positive integer")
        text = self.text_of(p)
        starts = []
        i = text.find(find)
        while i >= 0:
            starts.append(i)
            i = text.find(find, i + 1)
        if not starts:
            raise DocError(
                "Text not found in paragraph %s. Current text: %s"
                % (p.get(w14("paraId")), text), pid=p.get(w14("paraId")), current=text)
        if len(starts) > 1 and occurrence is None:
            raise DocError(
                "Ambiguous: %r occurs %d times in paragraph %s. Pass occurrence (1-based), "
                "or give a longer unique phrase." % (find, len(starts), p.get(w14("paraId"))),
                pid=p.get(w14("paraId")), occurrences=len(starts))
        k = 0 if occurrence is None else occurrence - 1
        if k < 0 or k >= len(starts):
            raise DocError("occurrence %s out of range (%d matches)" % (occurrence, len(starts)))
        return starts[k], starts[k] + len(find)

    def _check_hash(self, p, expect_hash):
        if expect_hash is None:
            return
        if not isinstance(expect_hash, str) or not expect_hash:
            raise DocError("expect_hash must be a non-empty string")
        actual = self.hash_of(p)
        if actual != expect_hash:
            raise DocError(
                "Paragraph %s changed since you read it (expected %s, now %s). "
                "Re-read it and retry. Current text: %s"
                % (p.get(w14("paraId")), expect_hash, actual, self.text_of(p)),
                pid=p.get(w14("paraId")), current=self.text_of(p), hash=actual)

    def _check_format_hash(self, p, expect_format_hash):
        if expect_format_hash is None:
            return
        if not isinstance(expect_format_hash, str) or not expect_format_hash:
            raise DocError("expect_format_hash must be a non-empty string")
        actual = self.format_hash_of(p)
        if actual != expect_format_hash:
            raise DocError(
                "Paragraph %s formatting changed since you read it "
                "(expected %s, now %s). Re-read it and retry."
                % (p.get(w14("paraId")), expect_format_hash, actual),
                pid=p.get(w14("paraId")), current=self.text_of(p),
                hash=self.hash_of(p), format_hash=actual)

    def _guard_runs(self, p, at, end):
        for (s, e, r, _t) in self.spans(p):
            if e <= at or s >= end:
                continue
            if not run_children_ok(r):
                raise DocError(
                    "Paragraph %s has a line break or other non-text content inside the "
                    "target span, so it cannot be edited safely. Target a smaller span."
                    % p.get(w14("paraId")), pid=p.get(w14("paraId")))

    def split_at(self, p, offsets):
        """Split runs so that every offset in `offsets` falls on a run boundary.

        Needed by format, which must give a span its own run to carry different
        rPr. Not needed by replace, which folds the replacement into the first
        covered run and so preserves the run count exactly.
        """
        for off in sorted(set(offsets), reverse=True):
            for (s, e, r, t) in self.spans(p):
                if s < off < e:
                    cut = off - s
                    tail = clone_run_empty(r)
                    set_run_text(tail, t[cut:])
                    set_run_text(r, t[:cut])
                    r.addnext(tail)
                    break

    # -- mutations ---------------------------------------------------------

    def replace(self, pid, find, replacement, expect_hash=None, occurrence=None):
        if not isinstance(replacement, str):
            raise DocError("replacement must be a string")
        p = self.para(pid)
        self._check_hash(p, expect_hash)
        at, end = self._locate(p, find, occurrence)
        self._guard_runs(p, at, end)

        before = self.text_of(p)
        if replacement == find:
            return {"pid": pid, "before": before, "after": before,
                    "start": at, "end": end, "hash": self.hash_of(p),
                    "noop": True}
        first = None
        for (s, e, r, t) in self.spans(p):
            if e <= at or s >= end:
                continue
            head = t[:at - s] if s < at else ""
            tail = t[end - s:] if e > end else ""
            if first is None:
                first = r
                set_run_text(r, head + replacement + tail)
            else:
                rest = head + tail
                # Keep every authored run/property shell, even when this
                # replacement consumes all of its visible text. A later edit
                # may need that boundary, and removing it would make a
                # cross-run replacement silently erase formatting.
                set_run_text(r, rest)

        self.rev += 1
        return {"pid": pid, "before": before, "after": self.text_of(p),
                "start": at, "end": at + len(replacement), "hash": self.hash_of(p)}

    def _insertion_target(self, p, at):
        """Return (run, offset) for typed text inserted at a character boundary."""
        text = self.text_of(p)
        if at < 0 or at > len(text):
            raise DocError("Insertion offset %d is outside paragraph %s."
                           % (at, p.get(w14("paraId"))))
        spans = self.spans(p)
        if not spans:
            return None, 0
        # At a run boundary either neighbor is valid. Prefer a normal text run
        # over a w:br/drawing run, whose children must remain untouched.
        candidates = [(r, max(0, min(len(t), at - s)))
                      for (s, e, r, t) in spans if s <= at <= e]
        for r, off in candidates:
            if run_children_ok(r):
                return r, off
        raise DocError(
            "Paragraph %s has a line break or other non-text content at the "
            "typing position, so it cannot be edited safely."
            % p.get(w14("paraId")), pid=p.get(w14("paraId")))

    def _replace_offsets(self, p, at, end, replacement):
        """Replace one offset range while preserving every run/property shell."""
        if at == end:
            r, off = self._insertion_target(p, at)
            if r is None:
                r = OxmlElement("w:r")
                p.append(r)
                set_run_text(r, replacement)
                return
            t = run_text(r)
            set_run_text(r, t[:off] + replacement + t[off:])
            return

        self._guard_runs(p, at, end)
        first = None
        for (s, e, r, t) in self.spans(p):
            if e <= at or s >= end:
                continue
            head = t[:at - s] if s < at else ""
            tail = t[end - s:] if e > end else ""
            if first is None:
                first = r
                set_run_text(r, head + replacement + tail)
            else:
                # Keep empty styled runs instead of removing them. Direct
                # select-all typing must not erase the paragraph's authored
                # run/property structure even when all visible text changes.
                set_run_text(r, head + tail)

    def _validated_text_changes(self, p, text):
        """Return a paragraph's plain-text diff after validating every target.

        Validation is deliberately separate from mutation so compound document
        operations, such as joining two paragraphs, can prove that all supplied
        local text is safe before changing either paragraph.
        """
        validate_xml_text(text)
        before = self.text_of(p)
        changes = [x for x in difflib.SequenceMatcher(
            None, before, text, autojunk=False).get_opcodes() if x[0] != "equal"]
        for _tag, i1, i2, _j1, _j2 in changes:
            if i1 == i2:
                self._insertion_target(p, i1)
            else:
                self._guard_runs(p, i1, i2)
        return before, changes

    def _apply_text_changes(self, p, text, changes):
        """Apply a diff already checked by :meth:`_validated_text_changes`."""
        for _tag, i1, i2, j1, j2 in reversed(changes):
            self._replace_offsets(p, i1, i2, text[j1:j2])

    def set_text(self, pid, text, expect_hash=None):
        """Apply human typing as localized diffs, preserving run formatting.

        The browser supplies the paragraph's final plain text. SequenceMatcher
        isolates disjoint insertions/deletions/replacements, then we apply them
        right-to-left so all offsets remain valid. This avoids treating two
        small corrections as one giant replacement that would flatten every
        formatted run between them.
        """
        p = self.para(pid)
        self._check_hash(p, expect_hash)
        before, changes = self._validated_text_changes(p, text)
        if text == before:
            return {"pid": pid, "before": before, "after": text,
                    "start": 0, "end": 0, "hash": self.hash_of(p),
                    "operations": 0, "noop": True}

        self._apply_text_changes(p, text, changes)

        # One encompassing glow range. The actual mutations above remain
        # localized even when the user changed two distant words.
        start = 0
        while start < len(before) and start < len(text) and before[start] == text[start]:
            start += 1
        suffix = 0
        while (suffix < len(before) - start and suffix < len(text) - start and
               before[len(before) - 1 - suffix] == text[len(text) - 1 - suffix]):
            suffix += 1
        self.rev += 1
        return {"pid": pid, "before": before, "after": self.text_of(p),
                "start": start, "end": len(text) - suffix,
                "hash": self.hash_of(p), "operations": len(changes)}

    def merge(self, first_pid, second_pid, first_text, second_text,
              expect_first_hash=None, expect_second_hash=None):
        """Join two adjacent paragraphs as one formatting-preserving mutation.

        The first paragraph remains and owns the resulting paragraph style.
        Every content node from the second paragraph is moved into it, so run
        formatting, bookmarks, comments, and other inline XML are preserved.
        Both content hashes and both local text snapshots are checked before
        either paragraph changes.
        """
        if not isinstance(first_text, str) or not isinstance(second_text, str):
            raise DocError("first_text and second_text must be strings")
        if not isinstance(expect_first_hash, str) or not expect_first_hash:
            raise DocError("expect_first_hash must be a non-empty string")
        if not isinstance(expect_second_hash, str) or not expect_second_hash:
            raise DocError("expect_second_hash must be a non-empty string")
        if first_pid == second_pid:
            raise DocError("Cannot merge a paragraph with itself.")

        first = self.para(first_pid)
        second = self.para(second_pid)
        self._check_hash(first, expect_first_hash)
        self._check_hash(second, expect_second_hash)

        paragraphs = self.paragraphs()
        first_index = paragraphs.index(first)
        if first_index + 1 >= len(paragraphs) or paragraphs[first_index + 1] is not second:
            raise DocError("Paragraphs must be adjacent and in document order to merge.")
        if first.getparent() is not second.getparent():
            raise DocError(
                "Paragraphs in different document or table containers cannot be merged safely.")
        if first.getnext() is not second:
            raise DocError(
                "Paragraphs separated by another document structure cannot be merged safely.")

        # A paragraph-level section boundary is structural page metadata, not
        # merely paragraph formatting. Removing either mark could collapse a
        # section, so require Word for that uncommon operation.
        for paragraph in (first, second):
            props = paragraph.find(w("pPr"))
            if props is not None and props.find(w("sectPr")) is not None:
                raise DocError("Paragraphs at a section boundary cannot be merged safely.")

        before_first, first_changes = self._validated_text_changes(first, first_text)
        before_second, second_changes = self._validated_text_changes(second, second_text)

        self._apply_text_changes(first, first_text, first_changes)
        self._apply_text_changes(second, second_text, second_changes)

        # Keep the first paragraph's pPr (its style/numbering). Move every
        # content node rather than rebuilding runs, then remove the empty shell.
        for child in list(second):
            if child.tag == w("pPr"):
                continue
            second.remove(child)
            first.append(child)
        second.getparent().remove(second)

        after = self.text_of(first)
        self.rev += 1
        return {
            "pid": first_pid,
            "removed_pid": second_pid,
            "before_first": before_first,
            "before_second": before_second,
            "after": after,
            "start": len(first_text),
            "end": len(after),
            "join_offset": len(first_text),
            "hash": self.hash_of(first),
            "operations": len(first_changes) + len(second_changes) + 1,
            "merged_pids": [first_pid, second_pid],
        }

    def insert(self, after_pid, text, style=None, expect_hash=None):
        if not isinstance(text, str):
            raise DocError("text must be a string")
        if style is not None and (not isinstance(style, str) or not style.strip()):
            raise DocError("style must be a non-empty string")
        ref = self.para(after_pid)
        self._check_hash(ref, expect_hash)
        new = OxmlElement("w:p")

        src_pr = ref.find(w("pPr"))
        if not style and src_pr is not None:
            new.append(copy.deepcopy(src_pr))
        if style:
            # An explicit named style is authoritative for the whole new
            # paragraph. Heading anchors commonly carry direct keepNext,
            # spacing, indentation, numbering, or section metadata in pPr;
            # cloning any of it can push ordinary body prose to the next page
            # or duplicate structural boundaries. Unstyled inserts retain the
            # legacy exact-clone behavior above.
            pr = OxmlElement("w:pPr")
            st = OxmlElement("w:pStyle")
            st.set(w("val"), self._style_id(style))
            pr.append(st)
            new.append(pr)

        r = OxmlElement("w:r")
        # Inherit the reference paragraph's first-run formatting so an inserted
        # paragraph matches the prose around it instead of reverting to default.
        # An explicit paragraph style is authoritative, though: cloning a
        # heading's direct bold/size overrides into a requested Normal paragraph
        # makes the continuation look like another heading.
        first_run = ref.find(w("r"))
        if not style and first_run is not None:
            src_rpr = first_run.find(w("rPr"))
            if src_rpr is not None:
                r.append(copy.deepcopy(src_rpr))
        set_run_text(r, text)
        new.append(r)

        new.set(w14("paraId"), self._fresh_id(self._all_ids(), len(self.paragraphs()), text))
        ref.addnext(new)
        self.rev += 1
        return {"pid": new.get(w14("paraId")), "after_pid": after_pid,
                "text": text, "hash": self.hash_of(new),
                "format_hash": self.format_hash_of(new)}

    def _style_id(self, name):
        wanted = name.casefold()
        for sid, nm in self._style_names.items():
            if (sid in self._paragraph_style_ids
                    and isinstance(nm, str) and nm.casefold() == wanted):
                return sid
        raise DocError("Unknown paragraph style: %s" % name, style=name)

    def delete(self, pid, expect_hash=None):
        p = self.para(pid)
        self._check_hash(p, expect_hash)
        if len(self.paragraphs()) <= 1:
            raise DocError("Refusing to delete the last paragraph in the document.")
        text = self.text_of(p)
        p.getparent().remove(p)
        self.rev += 1
        return {"pid": pid, "text": text}

    def format(self, pid, find, expect_hash=None, expect_format_hash=None,
               occurrence=None, **props):
        """Apply run properties to a span, splitting runs so the span carries its own rPr."""
        known = {"b", "i", "u", "color", "highlight", "size"}
        unknown = set(props) - known
        if unknown:
            raise DocError("Unknown formatting property: %s. Allowed: %s"
                           % (", ".join(sorted(unknown)), ", ".join(sorted(known))))
        if not props:
            raise DocError("format requires at least one formatting property")
        props = self._normalize_format_props(props)
        p = self.para(pid)
        self._check_hash(p, expect_hash)
        self._check_format_hash(p, expect_format_hash)
        at, end = self._locate(p, find, occurrence)
        self._guard_runs(p, at, end)

        self.split_at(p, [at, end])
        touched = 0
        for (s, e, r, _t) in self.spans(p):
            if e <= at or s >= end:
                continue
            self._apply_props(r, props)
            touched += 1

        self.rev += 1
        return {"pid": pid, "start": at, "end": end, "runs": touched,
                "hash": self.hash_of(p),
                "format_hash": self.format_hash_of(p)}

    @staticmethod
    def _drop_rpr_tags(rpr, tags):
        for tag in tags:
            for old in rpr.findall(w(tag)):
                rpr.remove(old)

    @staticmethod
    def _toggle_value(element):
        if element is None:
            return None
        value = element.get(w("val"))
        if value is None:
            return True
        return value.strip().casefold() not in OOXML_FALSE_VALUES

    def _direct_property_value(self, rpr, tag):
        if rpr is None:
            return None
        element = rpr.find(w(tag))
        if element is None:
            return None
        if tag in ("b", "bCs", "i", "iCs"):
            return self._toggle_value(element)
        if tag == "u":
            value = element.get(w("val"))
            if value is None or value == "":
                return "single"
            if value.strip().casefold() in OOXML_FALSE_VALUES:
                return False
            return value
        value = element.get(w("val"))
        return value if value not in (None, "") else False

    def _effective_direct_signature(self, paragraph_rpr, run_rpr, prop):
        values = []
        for tag in CLEARABLE_UNIFORM_DIRECT[prop]:
            run_value = self._direct_property_value(run_rpr, tag)
            values.append(
                run_value
                if run_value is not None
                else self._direct_property_value(paragraph_rpr, tag)
            )
        return tuple(values)

    @staticmethod
    def _signature_is_effective(prop, signature):
        if prop in ("b", "i"):
            # An explicit false can be protecting text from a bold/italic
            # style. Clearing it could add formatting rather than remove it.
            return True in signature and False not in signature
        if prop == "u":
            return bool(signature and signature[0] not in (None, False))
        return any(value not in (None, False) for value in signature)

    def _validate_uniform_direct_clear(self, p, prop):
        ppr = p.find(w("pPr"))
        paragraph_rprs = ppr.findall(w("rPr")) if ppr is not None else []
        if len(paragraph_rprs) > 1:
            raise DocError(
                "Paragraph %s has ambiguous duplicate paragraph run defaults, "
                "so automated formatting is unsafe."
                % p.get(w14("paraId")),
                pid=p.get(w14("paraId")), property=prop)
        paragraph_rpr = paragraph_rprs[0] if paragraph_rprs else None
        visible_runs = [r for r in p.findall(w("r")) if run_text(r) != ""]
        if not visible_runs:
            raise DocError(
                "Paragraph %s has no visible runs whose %s formatting can be "
                "cleared safely." % (p.get(w14("paraId")), prop),
                pid=p.get(w14("paraId")), property=prop)
        for run in p.findall(w("r")):
            run_rprs = run.findall(w("rPr"))
            if len(run_rprs) > 1:
                raise DocError(
                    "Paragraph %s has an ambiguous duplicate run-property "
                    "block, so automated formatting is unsafe."
                    % p.get(w14("paraId")),
                    pid=p.get(w14("paraId")), property=prop)
            rpr = run_rprs[0] if run_rprs else None
            for tag in CLEARABLE_UNIFORM_DIRECT[prop]:
                if rpr is not None and len(rpr.findall(w(tag))) > 1:
                    raise DocError(
                        "Paragraph %s has ambiguous duplicate direct %s "
                        "formatting, so automated formatting is unsafe."
                        % (p.get(w14("paraId")), prop),
                        pid=p.get(w14("paraId")), property=prop)
        if paragraph_rpr is not None:
            for tag in CLEARABLE_UNIFORM_DIRECT[prop]:
                if len(paragraph_rpr.findall(w(tag))) > 1:
                    raise DocError(
                        "Paragraph %s has ambiguous duplicate direct %s "
                        "defaults, so automated formatting is unsafe."
                        % (p.get(w14("paraId")), prop),
                        pid=p.get(w14("paraId")), property=prop)
        signatures = [
            self._effective_direct_signature(
                paragraph_rpr, r.find(w("rPr")), prop)
            for r in visible_runs
        ]
        if (not self._signature_is_effective(prop, signatures[0])
                or any(value != signatures[0] for value in signatures[1:])):
            raise DocError(
                "Paragraph %s does not have uniformly effective direct %s "
                "formatting across every visible run, so clearing it could "
                "damage intentional emphasis."
                % (p.get(w14("paraId")), prop),
                pid=p.get(w14("paraId")), property=prop)

    def _clear_direct_property(self, p, prop):
        tags = CLEARABLE_UNIFORM_DIRECT[prop]
        ppr = p.find(w("pPr"))
        if ppr is not None:
            for paragraph_rpr in ppr.findall(w("rPr")):
                self._drop_rpr_tags(paragraph_rpr, tags)
        for run in p.findall(w("r")):
            for rpr in run.findall(w("rPr")):
                self._drop_rpr_tags(rpr, tags)

    def _set_style_id(self, p, style_id):
        ppr = p.find(w("pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            p.insert(0, ppr)
        for old in ppr.findall(w("pStyle")):
            ppr.remove(old)
        style = OxmlElement("w:pStyle")
        style.set(w("val"), style_id)
        ppr.insert(0, style)

    @staticmethod
    def _restore_paragraph_xml(target, snapshot):
        """Restore one paragraph in place without invalidating element handles."""
        target.attrib.clear()
        target.attrib.update(snapshot.attrib)
        # CT_P defines a read-only aggregate ``text`` property; valid Word
        # paragraph text lives in descendant w:t nodes and is restored with the
        # children below.
        target.tail = snapshot.tail
        for child in list(target):
            target.remove(child)
        for child in snapshot:
            target.append(copy.deepcopy(child))

    def _paragraph_uses_style_id(self, p, style_id):
        ppr = p.find(w("pPr"))
        pstyle = ppr.find(w("pStyle")) if ppr is not None else None
        if pstyle is not None:
            return pstyle.get(w("val")) == style_id
        # An omitted pStyle resolves to the document's declared default
        # paragraph style, whose display name may be localized or renamed.
        return style_id == self._default_paragraph_style_id

    def format_batch(self, changes):
        """Apply a guarded paragraph-formatting plan as one model revision.

        Every target and action is validated before any paragraph is touched.
        The deliberately small operation surface lets an automated formatting
        pass fix paragraph classification and uniform leaked emphasis without
        flattening mixed authored formatting.
        """
        if not isinstance(changes, list):
            raise DocError("changes must be a list")
        if not changes:
            return {
                "changed": 0, "items": [], "noop": True, "rev": self.rev,
            }

        allowed_keys = {
            "pid", "expect_hash", "expect_format_hash",
            "style", "clear_uniform_direct", "reason",
        }
        seen = set()
        plan = []
        for index, change in enumerate(changes):
            if not isinstance(change, dict):
                raise DocError("changes[%d] must be an object" % index)
            unknown = set(change) - allowed_keys
            if unknown:
                raise DocError(
                    "changes[%d] has unknown field(s): %s"
                    % (index, ", ".join(sorted(unknown))))

            pid = change.get("pid")
            if not isinstance(pid, str) or not pid:
                raise DocError("changes[%d].pid must be a non-empty string" % index)
            if pid in seen:
                raise DocError("format_batch requires unique paragraph ids", pid=pid)
            seen.add(pid)

            expect_hash = change.get("expect_hash")
            if not isinstance(expect_hash, str) or not expect_hash:
                raise DocError(
                    "changes[%d].expect_hash must be a non-empty string" % index)
            expect_format_hash = change.get("expect_format_hash")
            if not isinstance(expect_format_hash, str) or not expect_format_hash:
                raise DocError(
                    "changes[%d].expect_format_hash must be a non-empty string"
                    % index)

            has_style = "style" in change
            style_name = change.get("style")
            if has_style and (
                    not isinstance(style_name, str) or not style_name.strip()):
                raise DocError(
                    "changes[%d].style must be a non-empty string" % index)
            style_id = self._style_id(style_name) if has_style else None
            canonical_style = (
                self._style_names.get(style_id, style_name) if has_style else None
            )

            clears = change.get("clear_uniform_direct", [])
            if not isinstance(clears, list):
                raise DocError(
                    "changes[%d].clear_uniform_direct must be a list" % index)
            if any(not isinstance(prop, str) for prop in clears):
                raise DocError(
                    "changes[%d].clear_uniform_direct entries must be strings"
                    % index)
            if len(clears) != len(set(clears)):
                raise DocError(
                    "changes[%d].clear_uniform_direct contains duplicates" % index)
            unknown_clears = set(clears) - set(CLEARABLE_UNIFORM_DIRECT)
            if unknown_clears:
                raise DocError(
                    "Unsupported clear_uniform_direct property: %s. Allowed: %s"
                    % (", ".join(sorted(unknown_clears)),
                       ", ".join(sorted(CLEARABLE_UNIFORM_DIRECT))))
            if not has_style and not clears:
                raise DocError(
                    "changes[%d] requires style or clear_uniform_direct" % index)
            reason = change.get("reason")
            if reason is not None and (
                    not isinstance(reason, str) or len(reason) > 240):
                raise DocError(
                    "changes[%d].reason must be a string of at most 240 characters"
                    % index)

            p = self.para(pid)
            self._check_hash(p, expect_hash)
            self._check_format_hash(p, expect_format_hash)
            if self.in_table(p):
                raise DocError(
                    "Automated paragraph formatting is not allowed inside tables.",
                    pid=pid)
            paragraph_properties = p.findall(w("pPr"))
            if len(paragraph_properties) > 1:
                raise DocError(
                    "Paragraph %s has ambiguous duplicate paragraph-property "
                    "blocks, so automated formatting is unsafe." % pid,
                    pid=pid)
            ppr = paragraph_properties[0] if paragraph_properties else None
            if ppr is not None and len(ppr.findall(w("pStyle"))) > 1:
                raise DocError(
                    "Paragraph %s has ambiguous duplicate paragraph styles, "
                    "so automated formatting is unsafe." % pid,
                    pid=pid)
            if ppr is not None and (
                    ppr.find(w("numPr")) is not None
                    or ppr.find(w("sectPr")) is not None):
                raise DocError(
                    "Numbered or section-boundary paragraphs cannot be "
                    "formatted automatically.", pid=pid)
            for prop in clears:
                self._validate_uniform_direct_clear(p, prop)

            style_changed = (
                has_style
                and not self._paragraph_uses_style_id(p, style_id)
            )
            plan.append({
                "p": p,
                "pid": pid,
                "style_id": style_id,
                "style_name": canonical_style,
                "style_changed": style_changed,
                "clears": list(clears),
            })

        # Mutation starts only after the complete plan has passed validation.
        # Keep in-place snapshots as a final line of defense: although the
        # operations below are deliberately simple, an unexpected XML/library
        # exception must not strand half of an otherwise atomic batch.
        snapshots = [(item["p"], copy.deepcopy(item["p"])) for item in plan]
        changed_items = []
        try:
            for item in plan:
                p = item["p"]
                if item["style_changed"]:
                    self._set_style_id(p, item["style_id"])
                for prop in item["clears"]:
                    self._clear_direct_property(p, prop)
                if item["style_changed"] or item["clears"]:
                    changed_items.append({
                        "pid": item["pid"],
                        "style": self.style_of(p),
                        "style_changed": item["style_changed"],
                        "cleared": item["clears"],
                        "hash": self.hash_of(p),
                        "format_hash": self.format_hash_of(p),
                    })
        except Exception:
            for paragraph, snapshot in snapshots:
                self._restore_paragraph_xml(paragraph, snapshot)
            raise

        if not changed_items:
            return {
                "changed": 0, "items": [], "noop": True, "rev": self.rev,
            }
        self.rev += 1
        return {
            "changed": len(changed_items),
            "items": changed_items,
            "noop": False,
            "rev": self.rev,
        }

    def _normalize_format_props(self, props):
        """Validate every value before format() can split or change a run."""
        normalized = {}
        for key, value in props.items():
            if key in ("b", "i"):
                if not isinstance(value, bool):
                    raise DocError("%s must be a boolean" % key)
                normalized[key] = value
                continue

            if key == "color":
                if value in (None, ""):
                    normalized[key] = None
                elif isinstance(value, str) and re.fullmatch(
                        r"#?[0-9A-Fa-f]{6}", value):
                    normalized[key] = value.lstrip("#").upper()
                else:
                    raise DocError("color must be 6 hexadecimal digits or empty")
                continue

            if key == "highlight":
                if value is None or value is False or (
                        isinstance(value, str) and value == ""):
                    normalized[key] = None
                elif isinstance(value, str):
                    canonical = next(
                        (item for item in HIGHLIGHT_VALUES
                         if item.lower() == value.lower()), None)
                    if canonical is None:
                        raise DocError("Unknown highlight color: %s" % value)
                    normalized[key] = canonical
                else:
                    raise DocError("highlight must be a Word highlight name or empty")
                continue

            if key == "u":
                if value is None or value is False or (
                        isinstance(value, str) and value == ""):
                    normalized[key] = None
                elif isinstance(value, str):
                    canonical = next(
                        (item for item in UNDERLINE_VALUES
                         if item.lower() == value.lower()), None)
                    if canonical is None:
                        raise DocError("Unknown underline style: %s" % value)
                    normalized[key] = canonical
                else:
                    raise DocError("u must be a Word underline style or empty")
                continue

            if key == "size":
                if value is None:
                    normalized[key] = None
                elif isinstance(value, bool) or not isinstance(value, (int, float)):
                    raise DocError("size must be a finite number of points or empty")
                else:
                    number = float(value)
                    if number == 0:
                        normalized[key] = None
                        continue
                    if not math.isfinite(number) or number < 1 or number > 1638:
                        raise DocError("size must be between 1 and 1638 points")
                    normalized[key] = number
                continue

        return normalized

    def _apply_props(self, r, props):
        pr = ensure_rpr(r)

        def drop(tag):
            for old in pr.findall(w(tag)):
                pr.remove(old)

        def toggle(tag, on):
            drop(tag)
            if on:
                pr.append(OxmlElement("w:" + tag))

        def valued(tag, val):
            drop(tag)
            if val is not None and val is not False:
                el = OxmlElement("w:" + tag)
                el.set(w("val"), str(val))
                pr.append(el)

        if "b" in props:
            toggle("b", props["b"])
        if "i" in props:
            toggle("i", props["i"])
        if "u" in props:
            valued("u", props["u"] if props["u"] else None)
        if "color" in props:
            valued("color", (props["color"] or "").lstrip("#").upper() or None)
        if "highlight" in props:
            valued("highlight", props["highlight"] or None)
        if "size" in props:
            valued("sz", int(round(float(props["size"]) * 2)) if props["size"] else None)

    # -- persistence -------------------------------------------------------

    def create_exact_backup(self):
        """Publish an exact sibling backup without exposing partial bytes.

        The copy is built under a private temporary name, flushed, and then
        linked into place.  The hard-link publish is atomic and refuses an
        existing destination on both Windows and POSIX, so even two document
        hosts saving in the same instant cannot overwrite one another's
        recovery point.
        """
        directory = os.path.dirname(self.path)
        root, ext = os.path.splitext(self.path)
        stamp = datetime.now().strftime("%Y-%m-%d-%H%M%S-%f")
        for _attempt in range(100):
            token = secrets.token_hex(6)
            backup = "%s.scribe-backup-%s-p%d-%s%s" % (
                root, stamp, os.getpid(), token, ext)
            fd, tmp = tempfile.mkstemp(
                prefix=".%s.scribe-backup-copy-" % os.path.basename(root),
                suffix=".tmp",
                dir=directory,
            )
            os.close(fd)
            try:
                shutil.copy2(self.path, tmp)
                # copy2 has returned, but explicitly flush the copied inode
                # before publishing its durable name.
                # Windows requires a writable handle for FlushFileBuffers,
                # which is what os.fsync delegates to there.
                with open(tmp, "r+b") as copied:
                    os.fsync(copied.fileno())
                try:
                    os.link(tmp, backup)
                except FileExistsError:
                    # A cryptographic-token collision is extraordinarily
                    # unlikely, but must still never overwrite old recovery
                    # data.
                    continue
                try:
                    os.unlink(tmp)
                except OSError:
                    # Do not report success while leaving an implementation
                    # temporary behind.  Remove the just-published link and
                    # make the whole attempt retryable.
                    try:
                        os.unlink(backup)
                    except OSError:
                        pass
                    raise
                return backup
            finally:
                try:
                    os.unlink(tmp)
                except FileNotFoundError:
                    pass
        raise OSError("Could not allocate a unique Scribe backup filename.")

    def ensure_backup(self):
        """One timestamped sibling backup per session, before the first write.

        Matches the convention already in the project
        (LEAN-Bench_McNair_2026.pre-audit-backup-2026-07-16.docx).
        """
        if self.backup_path:
            if os.path.isfile(self.backup_path):
                return self.backup_path
            # A backup removed outside Scribe must not remain a truthy marker
            # that suppresses recovery protection forever.
            self.backup_path = None
        backup = self.create_exact_backup()
        # Assign only after the exact copy has been durably published.  A
        # failed or partial copy therefore remains retryable.
        self.backup_path = backup
        return backup

    def save(self, expect_package_sha256=None):
        """Atomic: write a temp sibling, then os.replace. Never a partial file."""
        self._check_word_lock()
        self._check_expected_package(expect_package_sha256)
        self.ensure_backup()
        tmp = "%s.scribe-tmp-%d.docx" % (os.path.splitext(self.path)[0], os.getpid())
        try:
            t0 = time.perf_counter()
            self.doc.save(tmp)
            wrote = time.perf_counter() - t0
            package_sha256 = file_sha256(tmp)
            # Recheck both guards immediately before the atomic replacement:
            # the document may have been opened in Word or externally saved
            # while python-docx was constructing the temporary package.
            self._check_word_lock()
            self._check_expected_package(expect_package_sha256)
            os.replace(tmp, self.path)
        finally:
            try:
                os.unlink(tmp)
            except FileNotFoundError:
                pass
        self.package_sha256 = package_sha256
        return {"path": self.path, "backup": self.backup_path,
                "save_ms": round(wrote * 1000, 1), "rev": self.rev,
                "package_sha256": package_sha256}
