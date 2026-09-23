"""Verify (1) w14:paraId can be stamped on python-docx paragraphs and survives a
round trip, and (2) the run-splitting replace preserves formatting on a real
fragmented paragraph. Read-only against the real corpus: writes to scratchpad."""
import shutil, sys, hashlib
from pathlib import Path
import docx
from docx.oxml.ns import qn, nsmap
from lxml import etree
from live_input import required_live_file

SRC = required_live_file("SCRIBE_TEST_DOCX")
TMP = Path(__file__).parent / "paraid_out.docx"

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"

def pid_of(p):
    return p._p.get("{%s}paraId" % W14)

d = docx.Document(str(SRC))
body_ps = d.paragraphs
print(f"paragraphs: {len(body_ps)}")
print(f"existing w14:paraId count: {sum(1 for p in body_ps if pid_of(p))}")

# --- 1. Stamp deterministic unique ids -------------------------------------
seen = set()
stamped = 0
for i, p in enumerate(body_ps):
    if pid_of(p):
        continue
    # Word requires 8 hex digits, nonzero, unique. Derive from content+index.
    h = hashlib.sha1(f"{i}\x00{p.text}".encode("utf8")).hexdigest()[:8].upper()
    if h == "00000000" or h in seen:
        h = f"{(i + 1):08X}"
    seen.add(h)
    p._p.set("{%s}paraId" % W14, h)
    stamped += 1
print(f"stamped: {stamped}")

# --- 2. Run-splitting replace on a known fragmented paragraph ---------------
def para_runs(p):
    return p._p.findall(qn("w:r"))

def run_text(r):
    return "".join(t.text or "" for t in r.findall(qn("w:t")))

def replace_span(p, find, repl):
    """Replace `find` inside paragraph p, splitting runs and preserving each
    surviving run's rPr. Returns True if replaced."""
    runs = para_runs(p)
    spans, pos = [], 0
    for r in runs:
        t = run_text(r)
        spans.append((pos, pos + len(t), r, t))
        pos += len(t)
    full = "".join(s[3] for s in spans)
    at = full.find(find)
    if at < 0:
        return False
    end = at + len(find)

    first_covered = None
    to_remove = []
    for (s, e, r, t) in spans:
        if e <= at or s >= end:
            continue  # untouched
        if first_covered is None:
            first_covered = r
        head = t[: max(0, at - s)] if s < at else ""
        tail = t[end - s:] if e > end else ""
        if r is first_covered:
            # Rebuild this run as head + replacement + (tail if fully inside)
            new_t = head + repl + (tail if e > end else "")
            set_run_text(r, new_t)
        else:
            new_t = head + tail
            if new_t:
                set_run_text(r, new_t)
            else:
                to_remove.append(r)
    for r in to_remove:
        r.getparent().remove(r)
    return True

def set_run_text(r, s):
    for t in r.findall(qn("w:t")):
        r.remove(t)
    t = etree.SubElement(r, qn("w:t"))
    t.text = s
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

# find the most fragmented paragraph
target = max(body_ps, key=lambda p: len(para_runs(p)))
ti = body_ps.index(target)
before_colors = [(r.find(qn("w:rPr")) is not None) for r in para_runs(target)]
print(f"\ntarget paragraph index {ti}: {len(para_runs(target))} runs, "
      f"{sum(before_colors)} with rPr")
print("text:", target.text[:160].replace("\n", " "))

# pick a phrase that spans a run boundary if possible
runs = para_runs(target)
if len(runs) >= 3:
    a, b = run_text(runs[0]), run_text(runs[1])
    phrase = (a[-12:] + b[:12]).strip()
else:
    phrase = target.text[:20]
print("cross-run phrase:", repr(phrase))

ok = replace_span(target, phrase, "[[REPLACED]]")
print("replace ok:", ok)
after = para_runs(target)
print(f"after: {len(after)} runs, {sum(1 for r in after if r.find(qn('w:rPr')) is not None)} with rPr")
print("new text:", target.text[:180].replace("\n", " "))

d.save(str(TMP))

# --- 3. Reload and confirm persistence -------------------------------------
d2 = docx.Document(str(TMP))
ps2 = d2.paragraphs
have = sum(1 for p in ps2 if pid_of(p))
uniq = len({pid_of(p) for p in ps2 if pid_of(p)})
print(f"\nRELOAD: {len(ps2)} paragraphs, {have} carry w14:paraId, {uniq} unique")
print("reloaded target text:", ps2[ti].text[:180].replace("\n", " "))
print("reloaded target runs:", len(ps2[ti]._p.findall(qn('w:r'))))
print("\nSIZE:", SRC.stat().st_size, "->", TMP.stat().st_size)
