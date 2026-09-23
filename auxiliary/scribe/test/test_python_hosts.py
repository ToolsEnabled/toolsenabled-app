"""Robustness tests for Scribe's Python JSON-lines hosts.

All files live in unique operating-system temp directories. Speech inference is
mocked, research uses a tiny temporary corpus/database, and document mutations
touch only a generated DOCX.

Run:
    python test/test_python_hosts.py
"""

import io
import json
import os
import sqlite3
import sys
import tempfile
import time
import types
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import docx
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "engine"))

import dochost  # noqa: E402
import docmodel  # noqa: E402
import research  # noqa: E402
import stt  # noqa: E402
from docmodel import w, w14  # noqa: E402


def run_host_main(module, protocol_input):
    stdin = io.StringIO(protocol_input)
    stdout = io.StringIO()
    stderr = io.StringIO()
    with (
        mock.patch.object(sys, "stdin", stdin),
        mock.patch.object(sys, "stdout", stdout),
        mock.patch.object(sys, "stderr", stderr),
    ):
        module.main()
    replies = [
        json.loads(line)
        for line in stdout.getvalue().splitlines()
        if line.strip()
    ]
    return replies, stderr.getvalue()


class ProtocolBoundaryTests(unittest.TestCase):
    def tearDown(self):
        dochost.STATE["doc"] = None

    def test_malformed_requests_do_not_kill_any_host(self):
        for module in (dochost, research, stt):
            with self.subTest(host=module.__name__):
                malformed = [
                    "{",
                    "[]",
                    "null",
                    '{"id":"bad","cmd":null}',
                    '{"id":"nan","cmd":"ping","value":NaN}',
                ]
                for line in malformed:
                    response = module.handle_line(line)
                    self.assertFalse(response["ok"], response)

                good = module.handle_line(
                    json.dumps({"id": "é😀", "cmd": "ping"}, ensure_ascii=False)
                )
                self.assertTrue(good["ok"], good)
                self.assertEqual(good["id"], "é😀")
                self.assertTrue(good["result"]["pong"])

    def test_request_lines_drain_oversize_and_resume(self):
        for module in (dochost, research, stt):
            with self.subTest(host=module.__name__):
                stream = io.StringIO("x" * 80 + "\n{}\n")
                with mock.patch.object(module, "MAX_REQUEST_CHARS", 16):
                    self.assertEqual(list(module.request_lines(stream)), [None, "{}\n"])
                unicode_stream = io.StringIO("😀" * 5 + "\n{}\n")
                with mock.patch.object(module, "MAX_REQUEST_CHARS", 16):
                    self.assertEqual(
                        list(module.request_lines(unicode_stream)), [None, "{}\n"]
                    )

    def test_main_processes_unicode_then_exits_cleanly_at_eof(self):
        request = json.dumps(
            {"id": "日本語😀", "cmd": "ping"}, ensure_ascii=False
        ) + "\n"
        for module in (dochost, research, stt):
            with self.subTest(host=module.__name__):
                replies, stderr = run_host_main(module, request)
                # STT has one documented unsolicited hello response.
                response = next(reply for reply in replies if reply["id"] == "日本語😀")
                self.assertTrue(response["ok"])
                self.assertIn("stdin closed, exiting", stderr)

    def test_broken_pipe_is_a_clean_stop_signal(self):
        class BrokenOutput:
            def write(self, _):
                raise BrokenPipeError()

            def flush(self):
                raise AssertionError("flush should not follow a broken write")

        for module in (dochost, research, stt):
            with self.subTest(host=module.__name__):
                with mock.patch.object(sys, "stdout", BrokenOutput()):
                    self.assertFalse(module.reply({"ok": True}))


class DocumentHostTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory(prefix="scribe-host-doc-")
        self.addCleanup(self.tempdir.cleanup)
        self.path = Path(self.tempdir.name) / "文書😀.docx"
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("café token token")
        document.save(self.path)
        dochost.STATE["doc"] = None
        opened = dochost.cmd_open({"path": str(self.path)})
        self.assertEqual(opened["count"], 1)
        self.opened = opened
        self.document = dochost.STATE["doc"]
        self.paragraph = self.document.paragraphs()[0]
        self.pid = self.paragraph.get(w14("paraId"))

    def tearDown(self):
        dochost.STATE["doc"] = None

    def test_open_persists_repaired_ids_before_the_first_edit(self):
        self.assertEqual(self.opened["stamped"], 1)
        self.assertTrue(self.opened["stamps_persisted"])
        self.assertEqual(self.document.rev, 0)
        self.assertTrue(Path(self.document.backup_path).is_file())

        raw = docx.Document(self.path)
        raw_ids = [
            paragraph.get(w14("paraId"))
            for paragraph in raw.element.body.iter(w("p"))
        ]
        self.assertEqual(raw_ids, [self.pid])

        reopened = docmodel.ScribeDoc(str(self.path)).open()
        self.assertEqual(reopened._stamped, 0)
        self.assertEqual(
            [paragraph.get(w14("paraId")) for paragraph in reopened.paragraphs()],
            [self.pid],
        )

    def test_reload_persists_ids_from_a_legacy_unstamped_checkpoint(self):
        backup_path = self.document.backup_path
        backup_bytes = Path(backup_path).read_bytes()
        legacy = docx.Document()
        legacy.add_paragraph("cafÃ© token token")
        legacy_path = Path(self.tempdir.name) / "legacy-checkpoint.docx"
        legacy.save(legacy_path)

        # Add standard OOXML that Scribe does not interpret. The recovery copy
        # must retain these exact package bytes even though the active file is
        # normalized through python-docx.
        rewritten = legacy_path.with_suffix(".opaque.tmp")
        with zipfile.ZipFile(legacy_path, "r") as source:
            entries = [
                (info, source.read(info.filename))
                for info in source.infolist()
            ]
        with zipfile.ZipFile(rewritten, "w") as target:
            for info, payload in entries:
                if info.filename == "word/document.xml":
                    root = etree.fromstring(payload)
                    paragraph = root.find(".//" + w("p"))
                    marker = etree.Element(w("bookmarkStart"))
                    marker.set(w("id"), "77")
                    marker.set(w("name"), "ScribeOpaqueSentinel")
                    paragraph.insert(0, marker)
                    marker_end = etree.Element(w("bookmarkEnd"))
                    marker_end.set(w("id"), "77")
                    paragraph.insert(1, marker_end)
                    payload = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                target.writestr(info, payload)
        os.replace(rewritten, legacy_path)
        legacy_bytes = legacy_path.read_bytes()
        self.path.write_bytes(legacy_bytes)

        reloaded = dochost.cmd_reload({})
        self.assertEqual(reloaded["stamped"], 1)
        self.assertTrue(reloaded["stamps_persisted"])
        self.assertEqual(dochost.STATE["doc"].backup_path, backup_path)
        normalization_backup = Path(reloaded["normalization_backup"])
        self.assertNotEqual(normalization_backup, Path(backup_path))
        self.assertTrue(normalization_backup.is_file())
        self.assertEqual(normalization_backup.read_bytes(), legacy_bytes)
        self.assertEqual(Path(backup_path).read_bytes(), backup_bytes)
        self.assertNotEqual(self.path.read_bytes(), legacy_bytes)
        with zipfile.ZipFile(self.path) as active:
            self.assertIn(
                b"ScribeOpaqueSentinel",
                active.read("word/document.xml"),
            )

        persisted = docx.Document(self.path)
        persisted_ids = [
            paragraph.get(w14("paraId"))
            for paragraph in persisted.element.body.iter(w("p"))
        ]
        self.assertEqual(len(persisted_ids), 1)
        self.assertRegex(persisted_ids[0], r"^[0-9A-F]{8}$")

        stable_pid = persisted_ids[0]
        second_reload = dochost.cmd_reload({})
        self.assertEqual(second_reload["stamped"], 0)
        self.assertFalse(second_reload["stamps_persisted"])
        self.assertIsNone(second_reload["normalization_backup"])
        self.assertEqual(
            dochost.STATE["doc"].paragraphs()[0].get(w14("paraId")),
            stable_pid,
        )

    def test_strict_format_validation_is_atomic(self):
        invalid = [
            {"b": "false"},
            {"i": 1},
            {"u": []},
            {"u": True},
            {"color": {}},
            {"color": "not-hex"},
            {"highlight": {}},
            {"highlight": "orange"},
            {"size": "12"},
            {"size": float("nan")},
            {"size": float("inf")},
            {"size": 0},
        ]
        for bad_props in invalid:
            with self.subTest(props=repr(bad_props)):
                before = etree.tostring(self.paragraph)
                before_rev = self.document.rev
                with self.assertRaises(dochost.DocError):
                    dochost.cmd_format({
                        "pid": self.pid,
                        "find": "café",
                        "expect_hash": self.document.hash_of(self.paragraph),
                        **bad_props,
                    })
                self.assertEqual(self.document.rev, before_rev)
                self.assertEqual(etree.tostring(self.paragraph), before)

        # Non-standard JSON constants are rejected before dispatch as well.
        response = dochost.handle_line(
            '{"id":7,"cmd":"format","pid":"%s","find":"café","size":NaN}' % self.pid
        )
        self.assertFalse(response["ok"])
        self.assertIn("Invalid JSON", response["error"])

    def test_valid_unicode_and_clear_format_payloads_still_work(self):
        result = dochost.cmd_format({
            "pid": self.pid,
            "find": "café",
            "expect_hash": self.document.hash_of(self.paragraph),
            "b": True,
            "color": "#A1B2C3",
            "size": 12.5,
        })
        self.assertGreaterEqual(result["runs"], 1)
        marked = [
            run for run in self.paragraph.findall(w("r"))
            if self.document.run_props(run).get("b")
        ]
        self.assertEqual("".join(
            "".join(t.text or "" for t in run.findall(w("t"))) for run in marked
        ), "café")

        cleared = dochost.cmd_format({
            "pid": self.pid,
            "find": "café",
            "expect_hash": self.document.hash_of(self.paragraph),
            "b": False,
            "color": "",
            "highlight": None,
        })
        self.assertGreaterEqual(cleared["runs"], 1)

    def test_wrong_scalar_types_are_refused_without_stringifying(self):
        before = etree.tostring(self.paragraph)
        cases = [
            lambda: dochost.cmd_set_text({"pid": self.pid, "text": None}),
            lambda: dochost.cmd_insert({"after_pid": self.pid, "text": {}}),
            lambda: dochost.cmd_merge({
                "first_pid": self.pid,
                "second_pid": self.pid,
                "first_text": None,
                "second_text": "",
            }),
            lambda: dochost.cmd_replace({
                "pid": self.pid, "find": "café", "replace": []
            }),
            lambda: dochost.cmd_find({"query": "token", "regex": "false"}),
            lambda: dochost.cmd_read({"from": True}),
            lambda: dochost.cmd_read({"from": "1"}),
            lambda: dochost.cmd_read({"from": 1.5}),
        ]
        for call in cases:
            with self.subTest(call=repr(call)):
                with self.assertRaises(dochost.DocError):
                    call()
                self.assertEqual(etree.tostring(self.paragraph), before)

    def test_atomic_merge_command_uses_both_hashes(self):
        inserted = self.document.insert(self.pid, " second")
        second = self.document.para(inserted["pid"])
        first_text = self.document.text_of(self.paragraph)
        second_text = self.document.text_of(second)
        before = etree.tostring(self.document.doc.element.body)
        before_rev = self.document.rev

        for hashes in (
            {},
            {"expect_first_hash": self.document.hash_of(self.paragraph)},
            {"expect_second_hash": self.document.hash_of(second)},
        ):
            with self.subTest(hashes=hashes):
                with self.assertRaisesRegex(dochost.DocError, "non-empty string"):
                    dochost.cmd_merge({
                        "first_pid": self.pid,
                        "second_pid": inserted["pid"],
                        "first_text": first_text,
                        "second_text": second_text,
                        **hashes,
                    })
                self.assertEqual(
                    etree.tostring(self.document.doc.element.body), before)
                self.assertEqual(self.document.rev, before_rev)

        for first_local, second_local in (
            (first_text + "\x00", second_text),
            (first_text, second_text + "\x00"),
        ):
            with self.subTest(invalid_side=(first_local != first_text)):
                with self.assertRaisesRegex(
                        dochost.DocError, "cannot be stored in a Word document"):
                    dochost.cmd_merge({
                        "first_pid": self.pid,
                        "second_pid": inserted["pid"],
                        "first_text": first_local,
                        "second_text": second_local,
                        "expect_first_hash":
                            self.document.hash_of(self.paragraph),
                        "expect_second_hash": self.document.hash_of(second),
                    })
                self.assertEqual(
                    etree.tostring(self.document.doc.element.body), before)
                self.assertEqual(self.document.rev, before_rev)

        current_first_hash = self.document.hash_of(self.paragraph)
        current_second_hash = self.document.hash_of(second)
        for stale_side, first_hash, second_hash in (
            ("first", "deadbeef1234", current_second_hash),
            ("second", current_first_hash, "deadbeef1234"),
        ):
            with self.subTest(stale_side=stale_side):
                with self.assertRaisesRegex(dochost.DocError, "changed since"):
                    dochost.cmd_merge({
                        "first_pid": self.pid,
                        "second_pid": inserted["pid"],
                        "first_text": first_text,
                        "second_text": second_text,
                        "expect_first_hash": first_hash,
                        "expect_second_hash": second_hash,
                    })
                self.assertEqual(
                    etree.tostring(self.document.doc.element.body), before)
                self.assertEqual(self.document.rev, before_rev)

        result = dochost.cmd_merge({
            "first_pid": self.pid,
            "second_pid": inserted["pid"],
            "first_text": first_text,
            "second_text": second_text,
            "expect_first_hash": self.document.hash_of(self.paragraph),
            "expect_second_hash": self.document.hash_of(second),
        })
        self.assertEqual(result["after"], first_text + second_text)
        self.assertEqual(result["removed_pid"], inserted["pid"])
        self.assertEqual(self.document.rev, before_rev + 1)

    def test_document_regex_is_bounded_and_plain_search_stays_available(self):
        self.document.set_text(
            self.pid,
            ("a" * 50000) + "!",
            expect_hash=self.document.hash_of(self.paragraph),
        )
        plain = dochost.cmd_find({"query": "AAAA", "limit": 2})
        self.assertEqual(len(plain["hits"]), 2)

        if docmodel.timeout_regex is None:
            with self.assertRaisesRegex(dochost.DocError, "Safe regex"):
                dochost.cmd_find({"query": "(a+)+$", "regex": True})
        else:
            started = time.monotonic()
            with mock.patch.object(docmodel, "REGEX_SEARCH_TIMEOUT_S", 0.02):
                with self.assertRaisesRegex(dochost.DocError, "timed out"):
                    dochost.cmd_find({"query": "(a+)+$", "regex": True})
            self.assertLess(time.monotonic() - started, 0.5)


class ResearchHostTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory(prefix="scribe-host-research-")
        self.addCleanup(self.tempdir.cleanup)
        base = Path(self.tempdir.name)
        self.cache = base / "cache"
        self.root_a = base / "root-a"
        self.root_b = base / "root-b"
        self.root_a.mkdir()
        self.root_b.mkdir()
        (self.root_a / "z.txt").write_text("zeta semantic edge", encoding="utf-8")
        (self.root_a / "a.txt").write_text("alpha café", encoding="utf-8")
        (self.root_a / "duplicate.txt").write_text("alpha café", encoding="utf-8")
        (self.root_a / "same.txt").write_text("first same-name file", encoding="utf-8")
        (self.root_b / "same.txt").write_text("second same-name file", encoding="utf-8")

        self.database = base / "tiny db.sqlite"
        connection = sqlite3.connect(self.database)
        connection.execute("CREATE TABLE items (n INTEGER, payload TEXT, raw BLOB)")
        for n in range(201):
            payload = ("é" * 320) if n == 0 else "row-%d" % n
            connection.execute(
                "INSERT INTO items VALUES (?, ?, ?)",
                (n, payload, sqlite3.Binary(bytes([n % 256, 0, 255]))),
            )
        connection.commit()
        connection.close()

        self.config = base / "sources.json"
        self.config.write_text(json.dumps({
            "configured": True,
            "roots": [
                {"name": "a", "path": str(self.root_a), "priority": 1},
                {"name": "b", "path": str(self.root_b), "priority": 0},
            ],
            "extensions": [".TXT"],
            "exclude": [],
            "maxFileBytes": 100000,
            "maxFiles": 100,
            "databases": [
                {"name": "tiny", "path": str(self.database), "note": "temporary"}
            ],
        }), encoding="utf-8")

        self.patches = [
            mock.patch.object(research, "CONFIG", str(self.config)),
            mock.patch.object(research, "CACHE", str(self.cache)),
            mock.patch.object(research, "ROOT", str(base)),
        ]
        for patcher in self.patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        research.STATE.clear()
        research.STATE.update({"docs": [], "built_at": None, "stats": {}})

    def test_checked_in_sources_are_empty_and_refuse_before_any_probe(self):
        checked_in = ROOT / "sources.json"
        config = json.loads(checked_in.read_text(encoding="utf-8"))
        self.assertIs(config.get("configured"), False)
        self.assertEqual(config.get("roots"), [])
        self.assertEqual(config.get("databases"), [])

        with (
            mock.patch.object(research, "CONFIG", str(checked_in)),
            mock.patch.object(research.os, "walk") as walk,
            mock.patch.object(research.sqlite3, "connect") as connect,
        ):
            with self.assertRaisesRegex(research.Refused, "not configured"):
                research.cmd_index({"force": False})
            with self.assertRaisesRegex(research.Refused, "not configured"):
                research.cmd_db({"db": "anything", "sql": "SELECT 1"})
        walk.assert_not_called()
        connect.assert_not_called()

    def test_index_is_deterministic_deduplicated_and_cache_safe(self):
        first = research.cmd_index({"force": False})
        self.assertEqual(first["files"], 4)
        self.assertEqual(first["duplicates_skipped"], 1)
        self.assertEqual(
            [Path(doc["path"]).name for doc in research.STATE["docs"]],
            ["a.txt", "same.txt", "z.txt", "same.txt"],
        )
        alpha = next(doc for doc in research.STATE["docs"]
                     if Path(doc["path"]).name == "a.txt")
        self.assertEqual([Path(path).name for path in alpha["also_at"]],
                         ["duplicate.txt"])
        self.assertEqual(list(self.cache.glob("*.tmp-*")), [])

        second = research.cmd_index({"force": False})
        self.assertEqual(second["from_cache"], 5)
        self.assertEqual(
            [Path(doc["path"]).name for doc in research.STATE["docs"]],
            ["a.txt", "same.txt", "z.txt", "same.txt"],
        )

        # Same-size content with a changed nanosecond mtime must not reuse a
        # stale cache entry.
        target = self.root_a / "z.txt"
        old = target.stat()
        target.write_text("ZETA semantic edge", encoding="utf-8")
        # NTFS timestamps have 100 ns granularity; use a 1 ms delta that stays
        # well below the old cache key's one-second resolution.
        os.utime(target, ns=(old.st_atime_ns, old.st_mtime_ns + 1_000_000))
        research.cmd_index({"force": False})
        indexed = next(doc for doc in research.STATE["docs"]
                       if Path(doc["path"]).name == "z.txt")
        self.assertTrue(indexed["text"].startswith("ZETA"))

    def test_empty_built_index_is_not_rebuilt_forever(self):
        research.STATE.update({"docs": [], "built_at": 123.0, "stats": {"files": 0}})
        with mock.patch.object(research, "cmd_index") as index:
            research.ensure_index()
        index.assert_not_called()

    def test_search_and_read_inputs_are_bounded_and_unambiguous(self):
        research.cmd_index({})
        result = research.cmd_search({"query": "same", "limit": -20})
        self.assertLessEqual(len(result["hits"]), 1)
        with self.assertRaises(research.Refused):
            research.cmd_search({"query": ["not", "text"]})
        with self.assertRaises(research.Refused):
            research.cmd_search({"query": "same", "regex": "false"})
        with self.assertRaises(research.Refused):
            research.cmd_search({"query": "same", "limit": "2"})
        with self.assertRaises(research.Refused):
            research.cmd_search({"query": "same", "limit": 1.5})
        with self.assertRaisesRegex(research.Refused, "Ambiguous"):
            research.cmd_read({"path": "same.txt"})

        full_path = str(self.root_b / "same.txt")
        read = research.cmd_read({"path": full_path, "from": 10**9, "to": 10**10})
        self.assertEqual(read["from"], read["chars"])
        self.assertEqual(read["to"], read["chars"])
        self.assertEqual(read["text"], "")

        if research.timeout_regex is None:
            with self.assertRaisesRegex(research.Refused, "unavailable"):
                research.cmd_search({"query": "same", "regex": True})
        else:
            research.STATE["docs"].append({
                "path": str(self.root_a / "hostile.txt"),
                "root": "a",
                "priority": 1,
                "mtime": time.time(),
                "text": "a" * 20000 + "!",
                "lower": "a" * 20000 + "!",
                "also_at": [],
            })
            started = time.monotonic()
            with (
                mock.patch.object(research, "REGEX_SEARCH_TIMEOUT_S", 0.01),
                self.assertRaisesRegex(research.Refused, "timed out"),
            ):
                research.cmd_search({"query": "(a+)+$", "regex": True})
            self.assertLess(time.monotonic() - started, 1.0)

    def test_read_only_db_is_bounded_json_safe_and_times_out(self):
        result = research.cmd_db({
            "db": "tiny",
            "sql": "SELECT n, payload, raw FROM items ORDER BY n",
        })
        self.assertEqual(result["row_count"], 200)
        self.assertTrue(result["truncated"])
        self.assertTrue(result["rows"][0]["payload"].endswith("…(truncated)"))
        self.assertTrue(result["rows"][0]["raw"].startswith("<blob 3 bytes:"))
        json.dumps(result, ensure_ascii=False, allow_nan=False)

        exactly = research.cmd_db({
            "db": "tiny",
            "sql": "SELECT n FROM items WHERE n < 200 ORDER BY n",
        })
        self.assertEqual(exactly["row_count"], 200)
        self.assertFalse(exactly["truncated"])

        recursive = (
            "WITH RECURSIVE numbers(x) AS "
            "(SELECT 1 UNION ALL SELECT x+1 FROM numbers WHERE x < 100000000) "
            "SELECT sum(x) FROM numbers"
        )
        with mock.patch.object(research, "DB_QUERY_TIMEOUT_S", 0):
            with self.assertRaisesRegex(research.Refused, "timed out"):
                research.cmd_db({"db": "tiny", "sql": recursive})

        # mode=ro must leave the database unchanged.
        connection = sqlite3.connect(self.database)
        count = connection.execute("SELECT COUNT(*) FROM items").fetchone()[0]
        connection.close()
        self.assertEqual(count, 201)

    def test_pdf_cleanup_happens_when_extraction_raises(self):
        closed = []

        class BrokenPage:
            def get_text(self):
                raise RuntimeError("broken page")

        class FakePdf:
            page_count = 1

            def __iter__(self):
                return iter([BrokenPage()])

            def close(self):
                closed.append(True)

        fake_fitz = types.SimpleNamespace(open=lambda _: FakePdf())
        with mock.patch.dict(sys.modules, {"fitz": fake_fitz}):
            text, warning = research.extract("fake.pdf", ".pdf")
        self.assertEqual(text, "")
        self.assertIn("extract failed", warning)
        self.assertEqual(closed, [True])


class SpeechHostTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory(prefix="scribe-host-stt-")
        self.addCleanup(self.tempdir.cleanup)
        self.audio = Path(self.tempdir.name) / "audio.wav"
        self.audio.write_bytes(b"RIFF" + b"\0" * 100)
        stt.STATE.clear()
        stt.STATE.update({"model": None, "device": None, "loaded_ms": 0})

    def test_transcribe_validation_and_unicode_with_fake_model(self):
        calls = []

        class Model:
            def transcribe(self, path, **kwargs):
                calls.append((path, kwargs))
                segments = [
                    types.SimpleNamespace(text=" café"),
                    types.SimpleNamespace(text="漢字 😀 "),
                ]
                info = types.SimpleNamespace(duration=2.5, language="fr")
                return iter(segments), info

        stt.STATE.update({"model": Model(), "device": "cpu", "loaded_ms": 7})
        result = stt.cmd_transcribe({
            "path": str(self.audio),
            "beam": 3,
            "language": "fr",
            "prompt": "McNair café",
        })
        self.assertEqual(result["text"], "café 漢字 😀")
        self.assertEqual(result["audio_s"], 2.5)
        self.assertEqual(calls[0][1]["beam_size"], 3)
        self.assertEqual(calls[0][1]["initial_prompt"], "McNair café")

        before_calls = len(calls)
        bad = [
            {"path": None},
            {"path": str(Path(self.tempdir.name) / "missing.wav")},
            {"path": str(self.audio), "beam": True},
            {"path": str(self.audio), "beam": 0},
            {"path": str(self.audio), "beam": 11},
            {"path": str(self.audio), "beam": {}},
            {"path": str(self.audio), "beam": 1.5},
            {"path": str(self.audio), "beam": float("nan")},
            {"path": str(self.audio), "beam": float("inf")},
            {"path": str(self.audio), "language": []},
            {"path": str(self.audio), "language": "../../etc"},
            {"path": str(self.audio), "prompt": []},
            {"path": str(self.audio), "prompt": "x" * (stt.MAX_PROMPT_CHARS + 1)},
        ]
        for payload in bad:
            with self.subTest(payload=repr(payload)):
                with self.assertRaises((stt.Refused, ValueError)):
                    stt.cmd_transcribe(payload)
        self.assertEqual(len(calls), before_calls)

        with mock.patch.object(stt, "MAX_AUDIO_BYTES", 10):
            with self.assertRaisesRegex(stt.Refused, "too large"):
                stt.cmd_transcribe({"path": str(self.audio)})

    def test_transcript_output_is_bounded(self):
        class VerboseModel:
            def transcribe(self, *_args, **_kwargs):
                segments = [types.SimpleNamespace(text="x" * 100)]
                return iter(segments), types.SimpleNamespace(duration=1, language="en")

        stt.STATE.update({"model": VerboseModel(), "device": "cpu", "loaded_ms": 0})
        with mock.patch.object(stt, "MAX_TRANSCRIPT_CHARS", 20):
            with self.assertRaisesRegex(stt.Refused, "output limit"):
                stt.cmd_transcribe({"path": str(self.audio)})

    def test_model_load_falls_back_once_without_real_dependencies(self):
        attempts = []
        cpu_model = object()

        def whisper_model(_name, device, compute_type):
            attempts.append((device, compute_type))
            if device == "cuda":
                raise RuntimeError("no GPU")
            return cpu_model

        fake_whisper = types.SimpleNamespace(WhisperModel=whisper_model)
        with (
            mock.patch.dict(sys.modules, {"faster_whisper": fake_whisper}),
            mock.patch.object(stt, "configure_cuda") as cuda,
        ):
            loaded = stt.load()
            again = stt.load()
        self.assertIs(loaded, cpu_model)
        self.assertIs(again, cpu_model)
        self.assertEqual(attempts, [("cuda", "float16"), ("cpu", "int8")])
        self.assertEqual(stt.STATE["device"], "cpu")
        cuda.assert_called_once_with()


if __name__ == "__main__":
    unittest.main(verbosity=2)
