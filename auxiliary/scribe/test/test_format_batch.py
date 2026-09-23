"""Deterministic safety tests for agentic paragraph-format normalization.

The fixtures are generated in private temporary directories.  No running
Scribe process or user document is opened.

Run:
    python test/test_format_batch.py
"""

import copy
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "engine"))

import dochost  # noqa: E402
from docmodel import DocError, ScribeDoc, VISUAL_RPR_TAGS, w, w14  # noqa: E402


def append_value(parent, tag, value=None):
    element = OxmlElement("w:" + tag)
    if value is not None:
        element.set(w("val"), str(value))
    parent.append(element)
    return element


def ensure_run_rpr(run):
    rpr = run._r.find(w("rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run._r.insert(0, rpr)
    return rpr


def add_paragraph_defaults(paragraph, **values):
    ppr = paragraph._p.get_or_add_pPr()
    rpr = ppr.find(w("rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    for tag, value in values.items():
        append_value(rpr, tag, value)
    return rpr


class FormatHashAndBatchTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory(prefix="scribe-format-batch-")
        self.addCleanup(self.tempdir.cleanup)
        self.path = Path(self.tempdir.name) / "format-batch.docx"

        document = docx.Document()
        document.styles.add_style(
            "Agent Character Only", WD_STYLE_TYPE.CHARACTER)

        first = document.add_paragraph(style="Heading 1")
        first.add_run("A newly authored ")
        first.add_run("body paragraph.")
        defaults = add_paragraph_defaults(
            first, b=None, bCs=None, sz=32, szCs=32,
        )
        append_value(defaults, "rFonts", "IgnoredByFormatHash")
        append_value(first._p.get_or_add_pPr(), "keepNext")

        second = document.add_paragraph()
        for text in ("Second ", "paragraph."):
            run = second.add_run(text)
            rpr = ensure_run_rpr(run)
            append_value(rpr, "i")
            append_value(rpr, "iCs")

        mixed = document.add_paragraph()
        bold = mixed.add_run("Intentional")
        append_value(ensure_run_rpr(bold), "b")
        mixed.add_run(" mixed emphasis.")

        plain = document.add_paragraph("Plain formatting target.")
        document.save(self.path)

        self.doc = ScribeDoc(str(self.path)).open()
        self.first, self.second, self.mixed, self.plain = self.doc.paragraphs()

    def tearDown(self):
        dochost.STATE["doc"] = None

    def guards(self, paragraph):
        return {
            "pid": paragraph.get(w14("paraId")),
            "expect_hash": self.doc.hash_of(paragraph),
            "expect_format_hash": self.doc.format_hash_of(paragraph),
        }

    def test_format_hash_is_stable_and_span_format_checks_both_guards(self):
        paragraph = self.plain
        pid = paragraph.get(w14("paraId"))
        before_text_hash = self.doc.hash_of(paragraph)
        before_format_hash = self.doc.format_hash_of(paragraph)
        model = self.doc.para_model(paragraph)
        self.assertEqual(model["format_hash"], before_format_hash)
        self.assertRegex(before_format_hash, r"^[0-9a-f]{16}$")

        result = self.doc.format(
            pid,
            "formatting",
            expect_hash=before_text_hash,
            expect_format_hash=before_format_hash,
            b=True,
        )
        self.assertEqual(result["hash"], before_text_hash)
        self.assertEqual(result["format_hash"], self.doc.format_hash_of(paragraph))
        self.assertNotEqual(result["format_hash"], before_format_hash)

        before_xml = etree.tostring(self.doc.doc.element.body)
        before_rev = self.doc.rev
        with self.assertRaisesRegex(DocError, "formatting changed since"):
            self.doc.format(
                pid,
                "target",
                expect_hash=before_text_hash,
                expect_format_hash=before_format_hash,
                i=True,
            )
        self.assertEqual(etree.tostring(self.doc.doc.element.body), before_xml)
        self.assertEqual(self.doc.rev, before_rev)

        stable_hashes = {
            item.get(w14("paraId")): self.doc.format_hash_of(item)
            for item in self.doc.paragraphs()
        }
        saved = self.doc.save()
        reopened = ScribeDoc(saved["path"]).open()
        self.assertEqual({
            item.get(w14("paraId")): reopened.format_hash_of(item)
            for item in reopened.paragraphs()
        }, stable_hashes)

    def test_format_hash_is_canonical_but_covers_every_visual_block(self):
        first = copy.deepcopy(self.plain)
        second = copy.deepcopy(self.plain)
        first_rpr = OxmlElement("w:rPr")
        append_value(first_rpr, "b")
        append_value(first_rpr, "i")
        first.find(w("r")).insert(0, first_rpr)
        second_rpr = OxmlElement("w:rPr")
        append_value(second_rpr, "i")
        append_value(second_rpr, "b")
        second.find(w("r")).insert(0, second_rpr)

        self.assertEqual(
            self.doc.format_hash_of(first),
            self.doc.format_hash_of(second),
        )
        nonvisual = copy.deepcopy(first)
        append_value(nonvisual.get_or_add_pPr(), "keepLines")
        self.assertEqual(
            self.doc.format_hash_of(first),
            self.doc.format_hash_of(nonvisual),
        )
        duplicate_rpr = OxmlElement("w:rPr")
        append_value(duplicate_rpr, "bCs")
        second.find(w("r")).insert(1, duplicate_rpr)
        self.assertNotEqual(
            self.doc.format_hash_of(first),
            self.doc.format_hash_of(second),
        )

    def test_explicit_style_insert_does_not_clone_anchor_paragraph_properties(self):
        anchor = self.first
        anchor_pid = anchor.get(w14("paraId"))
        anchor_ppr = anchor.find(w("pPr"))
        anchor_defaults = anchor_ppr.find(w("rPr"))
        for tag in ("i", "iCs", "u", "color", "highlight"):
            append_value(
                anchor_defaults,
                tag,
                {
                    "u": "double",
                    "color": "FF0000",
                    "highlight": "yellow",
                }.get(tag),
            )
        first_run_rpr = ensure_run_rpr(
            docx.text.run.Run(anchor.find(w("r")), None)
        )
        append_value(first_run_rpr, "b")
        append_value(first_run_rpr, "sz", 40)
        duplicate_defaults = OxmlElement("w:rPr")
        append_value(duplicate_defaults, "bCs")
        append_value(duplicate_defaults, "szCs", 44)
        anchor_ppr.append(duplicate_defaults)
        spacing = OxmlElement("w:spacing")
        spacing.set(w("before"), "480")
        spacing.set(w("after"), "240")
        anchor_ppr.append(spacing)
        anchor_ppr.append(OxmlElement("w:numPr"))
        anchor_ppr.append(OxmlElement("w:sectPr"))

        unstyled = self.doc.insert(
            anchor_pid,
            "Legacy unstyled inheritance.",
            expect_hash=self.doc.hash_of(anchor),
        )
        unstyled_paragraph = self.doc.para(unstyled["pid"])
        self.assertEqual(
            etree.tostring(unstyled_paragraph.find(w("pPr"))),
            etree.tostring(anchor_ppr),
        )
        self.assertEqual(
            etree.tostring(unstyled_paragraph.find(w("r")).find(w("rPr"))),
            etree.tostring(anchor.find(w("r")).find(w("rPr"))),
        )

        inserted = self.doc.insert(
            anchor_pid,
            "This must be normal body prose.",
            style="Normal",
            expect_hash=self.doc.hash_of(anchor),
        )
        paragraph = self.doc.para(inserted["pid"])
        ppr = paragraph.find(w("pPr"))

        self.assertEqual(self.doc.style_of(paragraph), "Normal")
        self.assertEqual([child.tag for child in ppr], [w("pStyle")])
        for tag in (
                "keepNext", "spacing", "numPr", "sectPr", "rPr",
                "ind", "pageBreakBefore", "keepLines"):
            self.assertIsNone(ppr.find(w(tag)), tag)
        inserted_run = paragraph.find(w("r"))
        inserted_rpr = inserted_run.find(w("rPr"))
        self.assertTrue(
            inserted_rpr is None
            or all(inserted_rpr.find(w(tag)) is None for tag in VISUAL_RPR_TAGS)
        )

    def test_format_batch_is_atomic_and_one_revision(self):
        first_guard = self.guards(self.first)
        second_guard = self.guards(self.second)
        before_body = etree.tostring(self.doc.doc.element.body)
        before_rev = self.doc.rev

        with self.assertRaisesRegex(DocError, "formatting changed since"):
            self.doc.format_batch([
                {
                    **first_guard,
                    "style": "Normal",
                    "clear_uniform_direct": ["b", "size"],
                },
                {
                    **second_guard,
                    "expect_format_hash": "stale-format-guard",
                    "clear_uniform_direct": ["i"],
                },
            ])
        self.assertEqual(etree.tostring(self.doc.doc.element.body), before_body)
        self.assertEqual(self.doc.rev, before_rev)

        ids_before = [
            paragraph.get(w14("paraId")) for paragraph in self.doc.paragraphs()
        ]
        texts_before = [
            self.doc.text_of(paragraph) for paragraph in self.doc.paragraphs()
        ]
        first_format_before = self.doc.format_hash_of(self.first)
        second_format_before = self.doc.format_hash_of(self.second)
        result = self.doc.format_batch([
            {
                **first_guard,
                "style": "Normal",
                "clear_uniform_direct": ["b", "size"],
            },
            {
                **second_guard,
                "style": "Normal",
                "clear_uniform_direct": ["i"],
            },
        ])

        self.assertEqual(result["changed"], 2)
        self.assertFalse(result["noop"])
        self.assertEqual(result["rev"], before_rev + 1)
        self.assertEqual(self.doc.rev, before_rev + 1)
        self.assertEqual(
            [paragraph.get(w14("paraId")) for paragraph in self.doc.paragraphs()],
            ids_before,
        )
        self.assertEqual(
            [self.doc.text_of(paragraph) for paragraph in self.doc.paragraphs()],
            texts_before,
        )
        self.assertEqual(self.doc.style_of(self.first), "Normal")
        self.assertNotEqual(self.doc.format_hash_of(self.first), first_format_before)
        self.assertNotEqual(self.doc.format_hash_of(self.second), second_format_before)
        self.assertEqual(
            {item["pid"] for item in result["items"]},
            {first_guard["pid"], second_guard["pid"]},
        )

        for paragraph, tags in (
            (self.first, ("b", "bCs", "sz", "szCs")),
            (self.second, ("i", "iCs")),
        ):
            ppr = paragraph.find(w("pPr"))
            rprs = (
                ([ppr.find(w("rPr"))] if ppr is not None else [])
                + [run.find(w("rPr")) for run in paragraph.findall(w("r"))]
            )
            for rpr in (item for item in rprs if item is not None):
                for tag in tags:
                    self.assertIsNone(rpr.find(w(tag)), (paragraph, tag))

        rev_after = self.doc.rev
        noop = self.doc.format_batch([{
            **self.guards(self.first),
            "style": "Normal",
        }])
        self.assertTrue(noop["noop"])
        self.assertEqual(noop["items"], [])
        self.assertEqual(self.doc.rev, rev_after)
        self.assertTrue(self.doc.format_batch([])["noop"])
        self.assertEqual(self.doc.rev, rev_after)

    def test_mixed_emphasis_and_bad_plan_leave_every_target_untouched(self):
        before = etree.tostring(self.doc.doc.element.body)
        rev = self.doc.rev
        with self.assertRaisesRegex(DocError, "uniformly effective direct b"):
            self.doc.format_batch([
                {
                    **self.guards(self.first),
                    "style": "Normal",
                },
                {
                    **self.guards(self.mixed),
                    "clear_uniform_direct": ["b"],
                },
            ])
        self.assertEqual(etree.tostring(self.doc.doc.element.body), before)
        self.assertEqual(self.doc.rev, rev)

        invalid_plans = [
            [{
                **self.guards(self.first),
                "style": "Definitely Missing Style",
            }],
            [{
                **self.guards(self.first),
                "style": "Agent Character Only",
            }],
            [{
                **self.guards(self.first),
                "style": "Normal",
            }, {
                **self.guards(self.first),
                "style": "Normal",
            }],
            [{
                "pid": self.first.get(w14("paraId")),
                "expect_hash": self.doc.hash_of(self.first),
                "style": "Normal",
            }],
            [{
                **self.guards(self.first),
                "clear_uniform_direct": ["color"],
            }],
        ]
        for plan in invalid_plans:
            with self.subTest(plan=plan):
                with self.assertRaises(DocError):
                    self.doc.format_batch(plan)
                self.assertEqual(etree.tostring(self.doc.doc.element.body), before)
                self.assertEqual(self.doc.rev, rev)

    def test_duplicate_direct_tags_are_refused_and_apply_failures_roll_back(self):
        mixed_run = self.mixed.find(w("r"))
        append_value(mixed_run.find(w("rPr")), "b", "false")
        before = etree.tostring(self.doc.doc.element.body)
        rev = self.doc.rev
        with self.assertRaisesRegex(DocError, "ambiguous duplicate direct b"):
            self.doc.format_batch([{
                **self.guards(self.mixed),
                "clear_uniform_direct": ["b"],
            }])
        self.assertEqual(etree.tostring(self.doc.doc.element.body), before)
        self.assertEqual(self.doc.rev, rev)

        # Use fresh, valid targets for a forced failure after the first
        # paragraph has already changed. The in-memory transaction must restore
        # both paragraph objects in place.
        first_before = etree.tostring(self.first)
        second_before = etree.tostring(self.second)
        body_before = etree.tostring(self.doc.doc.element.body)
        original_clear = self.doc._clear_direct_property
        calls = 0

        def fail_on_second(paragraph, prop):
            nonlocal calls
            calls += 1
            original_clear(paragraph, prop)
            if calls == 2:
                raise RuntimeError("forced apply failure")

        with mock.patch.object(
                self.doc, "_clear_direct_property", side_effect=fail_on_second):
            with self.assertRaisesRegex(RuntimeError, "forced apply failure"):
                self.doc.format_batch([
                    {
                        **self.guards(self.first),
                        "style": "Normal",
                        "clear_uniform_direct": ["b"],
                    },
                    {
                        **self.guards(self.second),
                        "clear_uniform_direct": ["i"],
                    },
                ])
        self.assertEqual(etree.tostring(self.first), first_before)
        self.assertEqual(etree.tostring(self.second), second_before)
        self.assertEqual(etree.tostring(self.doc.doc.element.body), body_before)
        self.assertEqual(self.doc.rev, rev)

    def test_table_and_numbering_boundaries_are_refused(self):
        with tempfile.TemporaryDirectory(prefix="scribe-format-structure-") as tmp:
            path = Path(tmp) / "structure.docx"
            document = docx.Document()
            numbered = document.add_paragraph("Numbered structure")
            num_pr = OxmlElement("w:numPr")
            numbered._p.get_or_add_pPr().append(num_pr)
            add_paragraph_defaults(numbered, b=None)
            table_paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
            table_paragraph.add_run("Table structure")
            add_paragraph_defaults(table_paragraph, b=None)
            document.save(path)

            model = ScribeDoc(str(path)).open()
            targets = [model.paragraphs()[0], model.paragraphs()[1]]
            before = etree.tostring(model.doc.element.body)
            for target in targets:
                with self.subTest(pid=target.get(w14("paraId"))):
                    with self.assertRaises(DocError):
                        model.format_batch([{
                            "pid": target.get(w14("paraId")),
                            "expect_hash": model.hash_of(target),
                            "expect_format_hash": model.format_hash_of(target),
                            "clear_uniform_direct": ["b"],
                        }])
                    self.assertEqual(etree.tostring(model.doc.element.body), before)
                    self.assertEqual(model.rev, 0)

    def test_renamed_default_paragraph_style_remains_an_exact_noop(self):
        with tempfile.TemporaryDirectory(prefix="scribe-format-default-") as tmp:
            path = Path(tmp) / "renamed-default.docx"
            document = docx.Document()
            document.styles["Normal"].name = "Localized Body"
            document.add_paragraph("Default-styled body text.")
            document.save(path)

            model = ScribeDoc(str(path)).open()
            paragraph = model.paragraphs()[0]
            self.assertEqual(model.style_of(paragraph), "Localized Body")
            before = etree.tostring(paragraph)
            result = model.format_batch([{
                "pid": paragraph.get(w14("paraId")),
                "expect_hash": model.hash_of(paragraph),
                "expect_format_hash": model.format_hash_of(paragraph),
                "style": "Localized Body",
            }])
            self.assertTrue(result["noop"])
            self.assertEqual(model.rev, 0)
            self.assertEqual(etree.tostring(paragraph), before)

    def test_host_accepts_empty_underline_clear_and_dispatches_batch(self):
        dochost.STATE["doc"] = self.doc
        pid = self.plain.get(w14("paraId"))
        with self.assertRaises(DocError):
            dochost.cmd_format({
                "pid": pid,
                "find": "Plain",
                "expect_hash": self.doc.hash_of(self.plain),
                "expect_format_hash": self.doc.format_hash_of(self.plain),
                "u": 0,
            })
        added = dochost.cmd_format({
            "pid": pid,
            "find": "Plain",
            "expect_hash": self.doc.hash_of(self.plain),
            "expect_format_hash": self.doc.format_hash_of(self.plain),
            "u": "single",
        })
        cleared = dochost.cmd_format({
            "pid": pid,
            "find": "Plain",
            "expect_hash": added["hash"],
            "expect_format_hash": added["format_hash"],
            "u": "",
        })
        self.assertIn("format_hash", cleared)
        plain_run = next(
            run for run in self.plain.findall(w("r"))
            if "Plain" in "".join(t.text or "" for t in run.findall(w("t")))
        )
        self.assertIsNone(plain_run.find(w("rPr")).find(w("u")))

        result = dochost.cmd_format_batch({
            "changes": [{
                **self.guards(self.first),
                "style": "Normal",
                "reason": "Heading styling leaked into body prose.",
            }],
        })
        self.assertEqual(result["changed"], 1)
        self.assertEqual(self.doc.style_of(self.first), "Normal")


if __name__ == "__main__":
    unittest.main(verbosity=2)
