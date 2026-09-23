"""Deterministic property tests for the Scribe document engine.

The suite is intentionally independent of every project/content fixture. Each
test creates a synthetic DOCX in a unique operating-system temp directory and
removes it afterwards.

Run:
    python test/test_docmodel_fuzz.py
"""

import hashlib
import json
import os
import random
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest import mock

import docx
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "engine"))

import docmodel  # noqa: E402
from docmodel import DocError, ScribeDoc, run_children_ok, w, w14  # noqa: E402


UNICODE_FRAGMENTS = (
    "café",
    "naïve",
    "漢字",
    "e\u0301",
    "Ω",
    "😀",
    "Привет",
    "مرحبا",
)


def add_run(paragraph, text, *, bold=None, italic=None, color=None, size=None):
    run = paragraph.add_run(text)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    return run


def build_fixture(path):
    document = docx.Document()

    styled = document.add_paragraph()
    add_run(styled, "Alpha ", bold=True, color="0033CC", size=11)
    add_run(styled, "café 漢字 ", italic=True, color="008000", size=12)
    add_run(styled, "😀 omega", bold=False, color="AA00AA", size=13)
    add_run(styled, "")  # authored empty shell

    repeated = document.add_paragraph()
    add_run(repeated, "token ", bold=True)
    add_run(repeated, "token ", italic=True)
    add_run(repeated, "token")

    # Both empty shapes matter. python-docx emits no run for add_paragraph()
    # and an explicit empty w:r for add_run("").
    document.add_paragraph()
    empty_run = document.add_paragraph()
    empty_run.add_run("")

    boundary = document.add_paragraph()
    boundary.add_run("left")
    special = OxmlElement("w:r")
    special.append(OxmlElement("w:br"))
    boundary._p.append(special)
    boundary.add_run("right")

    unicode_paragraph = document.add_paragraph()
    add_run(unicode_paragraph, "naïve e\u0301 ", color="112233")
    add_run(unicode_paragraph, "Привет مرحبا 😀", bold=True)

    plain = document.add_paragraph("ordinary body text with spaces at both ends ")
    plain.runs[0].font.size = Pt(10.5)

    document.save(path)


def zip_hashes(path):
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
        }


def rpr_xml(paragraph):
    values = []
    for run in paragraph.findall(w("r")):
        prop = run.find(w("rPr"))
        values.append(None if prop is None else etree.tostring(prop))
    return values


def body_xml(document):
    return etree.tostring(document.doc.element.body)


def occurrences(text, phrase):
    found = []
    at = text.find(phrase)
    while at >= 0:
        found.append(at)
        at = text.find(phrase, at + 1)
    return found


def occurrence_at(text, phrase, at):
    starts = occurrences(text, phrase)
    return starts.index(at) + 1 if len(starts) > 1 else None


def char_properties(document, paragraph):
    """One canonical direct-formatting value per visible character."""
    values = []
    for _start, _end, run, text in document.spans(paragraph):
        props = tuple(sorted(document.run_props(run).items()))
        values.extend([props] * len(text))
    return values


def assert_model(test, document, expected_order, expected_text):
    paragraphs = document.paragraphs()
    ids = [p.get(w14("paraId")) for p in paragraphs]
    test.assertEqual(ids, expected_order)
    test.assertEqual(len(ids), len(set(ids)))
    test.assertNotIn("00000000", ids)
    test.assertEqual(
        {pid: document.text_of(document.para(pid)) for pid in ids},
        expected_text,
    )


class SyntheticDocTestCase(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory(prefix="scribe-docmodel-fuzz-")
        self.addCleanup(self.tempdir.cleanup)
        self.path = Path(self.tempdir.name) / "synthetic.docx"
        build_fixture(self.path)
        self.original_parts = zip_hashes(self.path)

    def open(self):
        return ScribeDoc(self.path).open()

    def assert_archive_ok(self):
        self.assertTrue(zipfile.is_zipfile(self.path))
        with zipfile.ZipFile(self.path) as archive:
            self.assertIsNone(archive.testzip())
            self.assertIn("[Content_Types].xml", archive.namelist())
            self.assertIn("word/document.xml", archive.namelist())
            xml = etree.fromstring(archive.read("word/document.xml"))
            self.assertTrue(xml.tag.endswith("}document"))
        self.assertEqual(
            list(self.path.parent.glob("*.scribe-tmp-*.docx")),
            [],
            "an atomic-save temp file was left behind",
        )


class DeterministicDocModelProperties(SyntheticDocTestCase):
    def test_failed_backup_copy_is_clean_and_retryable(self):
        document = self.open()
        source_bytes = self.path.read_bytes()

        def partial_copy(_source, destination):
            Path(destination).write_bytes(b"partial backup")
            raise OSError("forced backup copy failure")

        with mock.patch.object(docmodel.shutil, "copy2", side_effect=partial_copy):
            with self.assertRaisesRegex(OSError, "forced backup copy failure"):
                document.ensure_backup()

        self.assertIsNone(document.backup_path)
        self.assertEqual(
            list(self.path.parent.glob(".*.scribe-backup-copy-*.tmp")),
            [],
            "a failed backup left a partial implementation file behind",
        )
        self.assertEqual(
            list(self.path.parent.glob("*.scribe-backup-*.docx")),
            [],
            "a failed backup was published under a recovery filename",
        )

        backup = Path(document.ensure_backup())
        self.assertTrue(backup.is_file())
        self.assertEqual(backup.read_bytes(), source_bytes)
        self.assertEqual(document.ensure_backup(), str(backup))

    def test_rapid_new_sessions_never_overwrite_an_older_backup(self):
        class FrozenDatetime:
            @classmethod
            def now(cls):
                return datetime(2026, 7, 20, 12, 34, 56, 123456)

        original_bytes = self.path.read_bytes()
        with mock.patch.object(docmodel, "datetime", FrozenDatetime):
            first = self.open()
            first_backup = Path(first.ensure_backup())
            paragraph = first.paragraphs()[0]
            first.set_text(
                paragraph.get(w14("paraId")),
                first.text_of(paragraph) + " changed",
                expect_hash=first.hash_of(paragraph),
            )
            first.save()
            changed_bytes = self.path.read_bytes()

            # A host restart creates a fresh ScribeDoc and therefore has no
            # in-memory knowledge of the first host's backup.
            restarted = self.open()
            restarted_backup = Path(restarted.ensure_backup())

        self.assertNotEqual(first_backup, restarted_backup)
        self.assertEqual(first_backup.read_bytes(), original_bytes)
        self.assertEqual(restarted_backup.read_bytes(), changed_bytes)

    def test_synthetic_shape_unicode_and_archive(self):
        document = self.open()
        paragraphs = document.paragraphs()
        texts = [document.text_of(p) for p in paragraphs]
        ids = [p.get(w14("paraId")) for p in paragraphs]

        self.assertEqual(len(paragraphs), 7)
        self.assertEqual(document._stamped, 7)
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(all(pid and len(pid) == 8 for pid in ids))
        self.assertIn("Alpha café 漢字 😀 omega", texts)
        self.assertIn("naïve e\u0301 Привет مرحبا 😀", texts)
        self.assertEqual(sum(text == "" for text in texts), 2)
        self.assertEqual(
            sorted(len(p.findall(w("r"))) for p in paragraphs if not document.text_of(p)),
            [0, 1],
        )

        boundary = next(p for p in paragraphs if document.text_of(p) == "leftright")
        unsafe = [r for r in boundary.findall(w("r")) if not run_children_ok(r)]
        self.assertEqual(len(unsafe), 1)
        self.assertIsNotNone(unsafe[0].find(w("br")))
        self.assert_archive_ok()

    def test_unsafe_package_expansion_is_rejected_before_parsing(self):
        bomb = Path(self.tempdir.name) / "expansion.docx"
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr(
                "word/document.xml",
                "<document>" + ("x" * 4096) + "</document>",
            )
        with mock.patch("docmodel.MAX_DOCX_UNCOMPRESSED_BYTES", 1024):
            with self.assertRaisesRegex(DocError, "expanded document package"):
                ScribeDoc(bomb).open()

    def test_merge_stays_inside_one_table_cell(self):
        source = docx.Document()
        table = source.add_table(rows=1, cols=2)
        left_cell, right_cell = table.rows[0].cells
        left_cell.paragraphs[0].add_run("left ", style=None)
        left_cell.add_paragraph("right")
        right_cell.paragraphs[0].add_run("other")
        source.save(self.path)

        document = self.open()
        table_paragraphs = [p for p in document.paragraphs() if document.in_table(p)]
        first, second, other = table_paragraphs
        first_pid = first.get(w14("paraId"))
        second_pid = second.get(w14("paraId"))
        other_pid = other.get(w14("paraId"))
        result = document.merge(
            first_pid,
            second_pid,
            document.text_of(first),
            document.text_of(second),
            expect_first_hash=document.hash_of(first),
            expect_second_hash=document.hash_of(second),
        )
        self.assertEqual(result["after"], "left right")
        self.assertEqual(document.table_pos(document.para(first_pid))["cell"], 0)
        self.assertNotIn(second_pid, [
            p.get(w14("paraId")) for p in document.paragraphs()
        ])

        before = body_xml(document)
        rev = document.rev
        with self.assertRaisesRegex(DocError, "different document or table containers"):
            document.merge(
                first_pid,
                other_pid,
                document.text_of(document.para(first_pid)),
                document.text_of(document.para(other_pid)),
                expect_first_hash=document.hash_of(document.para(first_pid)),
                expect_second_hash=document.hash_of(document.para(other_pid)),
            )
        self.assertEqual(body_xml(document), before)
        self.assertEqual(document.rev, rev)

    def test_merge_keeps_only_the_first_paragraph_properties(self):
        source = docx.Document()
        first_source = source.add_paragraph("heading")
        second_source = source.add_paragraph("list item")

        first_properties = first_source._p.get_or_add_pPr()
        first_style = OxmlElement("w:pStyle")
        first_style.set(w("val"), "Heading2")
        first_spacing = OxmlElement("w:spacing")
        first_spacing.set(w("before"), "240")
        first_properties.extend([first_style, first_spacing])

        second_properties = second_source._p.get_or_add_pPr()
        second_style = OxmlElement("w:pStyle")
        second_style.set(w("val"), "ListParagraph")
        numbering = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(w("val"), "2")
        number_id = OxmlElement("w:numId")
        number_id.set(w("val"), "7")
        numbering.extend([level, number_id])
        alignment = OxmlElement("w:jc")
        alignment.set(w("val"), "center")
        second_properties.extend([second_style, numbering, alignment])
        source.save(self.path)

        document = self.open()
        first, second = document.paragraphs()
        first_pid = first.get(w14("paraId"))
        second_pid = second.get(w14("paraId"))
        first_properties_before = etree.tostring(first.find(w("pPr")))
        second_properties_before = etree.tostring(second.find(w("pPr")))
        self.assertNotEqual(first_properties_before, second_properties_before)

        document.merge(
            first_pid,
            second_pid,
            document.text_of(first),
            document.text_of(second),
            expect_first_hash=document.hash_of(first),
            expect_second_hash=document.hash_of(second),
        )
        merged = document.para(first_pid)
        self.assertEqual(len(merged.findall(w("pPr"))), 1)
        self.assertEqual(
            etree.tostring(merged.find(w("pPr"))),
            first_properties_before,
        )
        self.assertIsNone(merged.find(".//" + w("numPr")))
        self.assertIsNone(merged.find(w("pPr")).find(w("jc")))

        document.save()
        reopened = self.open()
        reopened_merged = reopened.para(first_pid)
        self.assertEqual(
            etree.tostring(reopened_merged.find(w("pPr"))),
            first_properties_before,
        )
        self.assertIsNone(reopened_merged.find(".//" + w("numPr")))
        with self.assertRaisesRegex(DocError, "No paragraph"):
            reopened.para(second_pid)

    def test_merge_refuses_structure_and_preserves_opaque_inline_xml(self):
        document = self.open()
        first, second = document.paragraphs()[:2]
        first_pid = first.get(w14("paraId"))
        second_pid = second.get(w14("paraId"))

        for paragraph in (first, second):
            props = paragraph.find(w("pPr"))
            if props is None:
                props = OxmlElement("w:pPr")
                paragraph.insert(0, props)
            section = OxmlElement("w:sectPr")
            props.append(section)
            before = body_xml(document)
            rev = document.rev
            with self.assertRaisesRegex(DocError, "section boundary"):
                document.merge(
                    first_pid,
                    second_pid,
                    document.text_of(first),
                    document.text_of(second),
                    expect_first_hash=document.hash_of(first),
                    expect_second_hash=document.hash_of(second),
                )
            self.assertEqual(body_xml(document), before)
            self.assertEqual(document.rev, rev)
            props.remove(section)

        between = OxmlElement("w:bookmarkStart")
        between.set(w("id"), "90")
        between.set(w("name"), "between-paragraphs")
        first.addnext(between)
        before = body_xml(document)
        rev = document.rev
        with self.assertRaisesRegex(DocError, "another document structure"):
            document.merge(
                first_pid,
                second_pid,
                document.text_of(first),
                document.text_of(second),
                expect_first_hash=document.hash_of(first),
                expect_second_hash=document.hash_of(second),
            )
        self.assertEqual(body_xml(document), before)
        self.assertEqual(document.rev, rev)
        between.getparent().remove(between)

        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(w("id"), "91")
        bookmark_start.set(w("name"), "inline-marker")
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(w("id"), "91")
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(w("anchor"), "inline-marker")
        link_run = OxmlElement("w:r")
        link_text = OxmlElement("w:t")
        link_text.text = "opaque linked text"
        link_run.append(link_text)
        hyperlink.append(link_run)
        second.insert(0, bookmark_start)
        second.append(hyperlink)
        second.append(bookmark_end)

        first_text = document.text_of(first)
        second_text = document.text_of(second)
        result = document.merge(
            first_pid,
            second_pid,
            first_text,
            second_text,
            expect_first_hash=document.hash_of(first),
            expect_second_hash=document.hash_of(second),
        )
        merged = document.para(first_pid)
        self.assertEqual(result["after"], first_text + second_text)
        self.assertIsNotNone(merged.find(w("bookmarkStart")))
        self.assertIsNotNone(merged.find(w("bookmarkEnd")))
        moved_link = merged.find(w("hyperlink"))
        self.assertIsNotNone(moved_link)
        self.assertEqual(
            moved_link.find(".//" + w("t")).text,
            "opaque linked text",
        )

        document.save()
        reopened = self.open()
        reopened_merged = reopened.para(first_pid)
        self.assertIsNotNone(reopened_merged.find(w("bookmarkStart")))
        self.assertIsNotNone(reopened_merged.find(w("bookmarkEnd")))
        self.assertEqual(
            reopened_merged.find(w("hyperlink")).find(".//" + w("t")).text,
            "opaque linked text",
        )
        with self.assertRaisesRegex(DocError, "No paragraph"):
            reopened.para(second_pid)

    def test_merge_handles_empty_paragraph_combinations(self):
        for index, (first_text, second_text) in enumerate((
            ("", "tail"),
            ("head", ""),
            ("", ""),
        )):
            with self.subTest(first=first_text, second=second_text):
                path = Path(self.tempdir.name) / ("empty-merge-%d.docx" % index)
                source = docx.Document()
                source.add_paragraph().add_run(first_text).bold = True
                source.add_paragraph().add_run(second_text).italic = True
                source.save(path)

                document = ScribeDoc(path).open()
                first, second = document.paragraphs()
                first_pid = first.get(w14("paraId"))
                second_pid = second.get(w14("paraId"))
                props = rpr_xml(first) + rpr_xml(second)
                result = document.merge(
                    first_pid,
                    second_pid,
                    first_text,
                    second_text,
                    expect_first_hash=document.hash_of(first),
                    expect_second_hash=document.hash_of(second),
                )
                self.assertEqual(result["after"], first_text + second_text)
                self.assertEqual(
                    document.text_of(document.para(first_pid)),
                    first_text + second_text,
                )
                self.assertEqual(rpr_xml(document.para(first_pid)), props)
                self.assertEqual(len(document.paragraphs()), 1)
                document.save()
                reopened = ScribeDoc(path).open()
                self.assertEqual(
                    reopened.text_of(reopened.para(first_pid)),
                    first_text + second_text,
                )
                self.assertEqual(rpr_xml(reopened.para(first_pid)), props)
                with self.assertRaisesRegex(DocError, "No paragraph"):
                    reopened.para(second_pid)

    def test_refusals_are_atomic_and_safe_boundaries_survive(self):
        document = self.open()
        repeated = next(p for p in document.paragraphs()
                        if document.text_of(p) == "token token token")
        repeated_pid = repeated.get(w14("paraId"))

        def atomic_refusal(call, message):
            before = body_xml(document)
            rev = document.rev
            with self.assertRaisesRegex(DocError, message):
                call()
            self.assertEqual(document.rev, rev)
            self.assertEqual(body_xml(document), before)

        atomic_refusal(
            lambda: document.replace(repeated_pid, "token", "x"),
            "Ambiguous",
        )
        atomic_refusal(
            lambda: document.format(repeated_pid, "token", b=True),
            "Ambiguous",
        )
        atomic_refusal(
            lambda: document.replace(repeated_pid, "", "x", occurrence=1),
            "non-empty",
        )
        atomic_refusal(
            lambda: document.format(repeated_pid, "token"),
            "at least one",
        )
        for occurrence in (True, 1.5, 0, -1):
            atomic_refusal(
                lambda occurrence=occurrence: document.replace(
                    repeated_pid, "token", "x", occurrence=occurrence
                ),
                "positive integer",
            )
        for props, message in (
            ({"b": "true"}, "boolean"),
            ({"color": "GG00GG"}, "hexadecimal"),
            ({"highlight": "laser"}, "Unknown highlight"),
            ({"highlight": 0}, "highlight must"),
            ({"u": "zigzag"}, "Unknown underline"),
            ({"u": 0}, "u must"),
            ({"size": False}, "finite number"),
            ({"size": float("nan")}, "between 1 and 1638"),
            ({"size": 5000}, "between 1 and 1638"),
        ):
            atomic_refusal(
                lambda props=props: document.format(
                    repeated_pid, "token", occurrence=1, **props
                ),
                message,
            )
        atomic_refusal(
            lambda: document.replace(repeated_pid, "token", object(), occurrence=1),
            "replacement must be a string",
        )
        atomic_refusal(
            lambda: document.insert(
                repeated_pid,
                "must not land",
                style="Definitely Missing",
                expect_hash=document.hash_of(repeated),
            ),
            "Unknown paragraph style",
        )
        atomic_refusal(
            lambda: document.replace(
                repeated_pid, "token", "x", expect_hash="deadbeef1234", occurrence=1
            ),
            "changed since",
        )
        atomic_refusal(
            lambda: document.set_text(
                repeated_pid, "must not land", expect_hash="deadbeef1234"
            ),
            "changed since",
        )
        atomic_refusal(
            lambda: document.insert(
                repeated_pid, "must not land", expect_hash="deadbeef1234"
            ),
            "changed since",
        )
        atomic_refusal(
            lambda: document.delete(repeated_pid, expect_hash="deadbeef1234"),
            "changed since",
        )
        atomic_refusal(
            lambda: document.replace(
                repeated_pid, "token", "\x00", occurrence=1,
                expect_hash=document.hash_of(repeated),
            ),
            "cannot be stored in a Word document",
        )
        atomic_refusal(
            lambda: document.set_text(
                repeated_pid,
                document.text_of(repeated) + "\x00",
                expect_hash=document.hash_of(repeated),
            ),
            "cannot be stored in a Word document",
        )
        atomic_refusal(
            lambda: document.insert(
                repeated_pid, "\x00", expect_hash=document.hash_of(repeated)
            ),
            "cannot be stored in a Word document",
        )

        merge_first, merge_second, merge_third = document.paragraphs()[:3]
        merge_first_pid = merge_first.get(w14("paraId"))
        merge_second_pid = merge_second.get(w14("paraId"))
        merge_third_pid = merge_third.get(w14("paraId"))
        merge_args = (
            document.text_of(merge_first),
            document.text_of(merge_second),
        )
        for missing_hashes in (
            {},
            {"expect_first_hash": document.hash_of(merge_first)},
            {"expect_second_hash": document.hash_of(merge_second)},
        ):
            atomic_refusal(
                lambda missing_hashes=missing_hashes: document.merge(
                    merge_first_pid, merge_second_pid,
                    merge_args[0], merge_args[1],
                    **missing_hashes,
                ),
                "must be a non-empty string",
            )
        for first_text, second_text in (
            (merge_args[0] + "\x00", merge_args[1]),
            (merge_args[0], merge_args[1] + "\x00"),
        ):
            atomic_refusal(
                lambda first_text=first_text, second_text=second_text:
                    document.merge(
                        merge_first_pid, merge_second_pid,
                        first_text, second_text,
                        expect_first_hash=document.hash_of(merge_first),
                        expect_second_hash=document.hash_of(merge_second),
                    ),
                "cannot be stored in a Word document",
            )
        atomic_refusal(
            lambda: document.merge(
                merge_first_pid, merge_first_pid,
                merge_args[0], merge_args[0],
                expect_first_hash=document.hash_of(merge_first),
                expect_second_hash=document.hash_of(merge_first),
            ),
            "itself",
        )
        atomic_refusal(
            lambda: document.merge(
                merge_second_pid, merge_first_pid,
                merge_args[1], merge_args[0],
                expect_first_hash=document.hash_of(merge_second),
                expect_second_hash=document.hash_of(merge_first),
            ),
            "adjacent",
        )
        atomic_refusal(
            lambda: document.merge(
                merge_first_pid, merge_third_pid,
                merge_args[0], document.text_of(merge_third),
                expect_first_hash=document.hash_of(merge_first),
                expect_second_hash=document.hash_of(merge_third),
            ),
            "adjacent",
        )
        atomic_refusal(
            lambda: document.merge(
                merge_first_pid, merge_second_pid,
                merge_args[0], merge_args[1],
                expect_first_hash="deadbeef1234",
                expect_second_hash=document.hash_of(merge_second),
            ),
            "changed since",
        )
        atomic_refusal(
            lambda: document.merge(
                merge_first_pid, merge_second_pid,
                merge_args[0], merge_args[1],
                expect_first_hash=document.hash_of(merge_first),
                expect_second_hash="deadbeef1234",
            ),
            "changed since",
        )

        boundary = next(p for p in document.paragraphs()
                        if document.text_of(p) == "leftright")
        boundary_pid = boundary.get(w14("paraId"))
        atomic_refusal(
            lambda: document.replace(boundary_pid, "leftr", "unsafe"),
            "non-text content",
        )
        atomic_refusal(
            lambda: document.format(boundary_pid, "leftr", color="FF0000"),
            "non-text content",
        )
        atomic_refusal(
            lambda: document.set_text(boundary_pid, "XXXXXXXXX"),
            "non-text content",
        )

        # Inserting exactly beside the w:br is safe. The text engine chooses a
        # normal neighboring run and leaves the special child untouched.
        document.set_text(
            boundary_pid,
            "left|right",
            expect_hash=document.hash_of(boundary),
        )
        self.assertEqual(document.text_of(boundary), "left|right")
        self.assertEqual(len(boundary.findall(".//" + w("br"))), 1)
        document.replace(
            boundary_pid,
            "left",
            "LEFT",
            expect_hash=document.hash_of(boundary),
        )
        self.assertEqual(document.text_of(boundary), "LEFT|right")
        self.assertEqual(len(boundary.findall(".//" + w("br"))), 1)

    def test_shell_preservation_empty_shapes_and_id_lifetime(self):
        document = self.open()
        styled = next(p for p in document.paragraphs()
                      if document.text_of(p).startswith("Alpha "))
        styled_pid = styled.get(w14("paraId"))
        before_props = rpr_xml(styled)
        before_runs = len(styled.findall(w("r")))
        before_text = document.text_of(styled)

        document.replace(
            styled_pid,
            before_text,
            "Entirely replacement text",
            expect_hash=document.hash_of(styled),
        )
        self.assertEqual(len(styled.findall(w("r"))), before_runs)
        self.assertEqual(rpr_xml(styled), before_props)
        self.assertEqual(
            [span[3] for span in document.spans(styled)],
            ["Entirely replacement text", "", "", ""],
        )

        empty_paragraphs = [p for p in document.paragraphs()
                            if document.text_of(p) == ""]
        zero_run = next(p for p in empty_paragraphs if not p.findall(w("r")))
        empty_run = next(p for p in empty_paragraphs if len(p.findall(w("r"))) == 1)
        for paragraph in (zero_run, empty_run):
            pid = paragraph.get(w14("paraId"))
            document.set_text(pid, "typed 😀", expect_hash=document.hash_of(paragraph))
            self.assertEqual(document.text_of(paragraph), "typed 😀")
            document.set_text(pid, "", expect_hash=document.hash_of(paragraph))
            self.assertEqual(document.text_of(paragraph), "")
            self.assertEqual(len(paragraph.findall(w("r"))), 1)

        anchor = document.paragraphs()[-1]
        anchor_pid = anchor.get(w14("paraId"))
        first = document.insert(anchor_pid, "same text")
        document.delete(first["pid"])
        second = document.insert(anchor_pid, "same text")
        self.assertNotEqual(first["pid"], second["pid"])

        ids_before = [p.get(w14("paraId")) for p in document.paragraphs()]
        props_before = {
            p.get(w14("paraId")): rpr_xml(p) for p in document.paragraphs()
        }
        document.save()
        self.assert_archive_ok()
        reopened = self.open()
        ids_after = [p.get(w14("paraId")) for p in reopened.paragraphs()]
        self.assertEqual(ids_after, ids_before)
        self.assertEqual(reopened._stamped, 0)
        self.assertEqual(
            {p.get(w14("paraId")): rpr_xml(p) for p in reopened.paragraphs()},
            props_before,
        )

    def test_seeded_random_mutations_and_round_trips(self):
        rng = random.Random(0x5C12BE)
        document = self.open()
        expected_order = [p.get(w14("paraId")) for p in document.paragraphs()]
        expected_text = {
            p.get(w14("paraId")): document.text_of(p)
            for p in document.paragraphs()
        }
        retired_ids = set()
        insert_serial = 0

        operations = (
            ["set_text"] * 16
            + ["replace"] * 16
            + ["format"] * 14
            + ["insert"] * 14
            + ["delete"] * 8
            + ["merge"] * 6
        )
        rng.shuffle(operations)

        for step, operation in enumerate(operations):
            paragraphs = document.paragraphs()
            safe = [
                p for p in paragraphs
                if all(run_children_ok(run) for run in p.findall(w("r")))
            ]
            nonempty_safe = [p for p in safe if document.text_of(p)]

            if operation == "set_text":
                paragraph = rng.choice(safe)
                pid = paragraph.get(w14("paraId"))
                before = document.text_of(paragraph)
                at = rng.randrange(len(before) + 1)
                drop = rng.randrange(min(4, len(before) - at) + 1)
                fragment = UNICODE_FRAGMENTS[rng.randrange(len(UNICODE_FRAGMENTS))]
                replacement = f"§{step}:{fragment}"
                desired = before[:at] + replacement + before[at + drop:]
                props_before = rpr_xml(paragraph)
                runs_before = len(paragraph.findall(w("r")))

                result = document.set_text(
                    pid, desired, expect_hash=document.hash_of(paragraph)
                )
                expected_text[pid] = desired
                self.assertEqual(result["after"], desired)
                if runs_before:
                    self.assertEqual(rpr_xml(paragraph), props_before)
                else:
                    self.assertEqual(rpr_xml(paragraph), [None])
                self.assertEqual(
                    len(paragraph.findall(w("r"))),
                    runs_before if runs_before else 1,
                )

            elif operation == "replace":
                paragraph = rng.choice(nonempty_safe)
                pid = paragraph.get(w14("paraId"))
                before = document.text_of(paragraph)
                at = rng.randrange(len(before))
                width = rng.randrange(1, min(5, len(before) - at) + 1)
                phrase = before[at:at + width]
                occurrence = occurrence_at(before, phrase, at)
                replacement = f"R{step}{UNICODE_FRAGMENTS[step % len(UNICODE_FRAGMENTS)]}"
                props_before = rpr_xml(paragraph)
                runs_before = len(paragraph.findall(w("r")))

                result = document.replace(
                    pid,
                    phrase,
                    replacement,
                    expect_hash=document.hash_of(paragraph),
                    occurrence=occurrence,
                )
                desired = before[:at] + replacement + before[at + width:]
                expected_text[pid] = desired
                self.assertEqual(result["after"], desired)
                self.assertEqual(len(paragraph.findall(w("r"))), runs_before)
                self.assertEqual(rpr_xml(paragraph), props_before)

            elif operation == "format":
                paragraph = rng.choice(nonempty_safe)
                pid = paragraph.get(w14("paraId"))
                text = document.text_of(paragraph)
                at = rng.randrange(len(text))
                width = rng.randrange(1, min(6, len(text) - at) + 1)
                phrase = text[at:at + width]
                occurrence = occurrence_at(text, phrase, at)
                char_props_before = char_properties(document, paragraph)
                runs_before = len(paragraph.findall(w("r")))
                prop = rng.choice([
                    {"b": bool(step % 2)},
                    {"i": bool((step + 1) % 2)},
                    {"color": rng.choice(["FF0000", "0055AA", ""])},
                    {"highlight": rng.choice(["yellow", "cyan", None])},
                    {"size": rng.choice([9, 11.5, 14])},
                    {"u": rng.choice([False, "single", "double"])},
                ])

                document.format(
                    pid,
                    phrase,
                    expect_hash=document.hash_of(paragraph),
                    occurrence=occurrence,
                    **prop,
                )
                char_props_after = char_properties(document, paragraph)
                self.assertEqual(document.text_of(paragraph), text)
                self.assertGreaterEqual(len(paragraph.findall(w("r"))), runs_before)
                self.assertEqual(char_props_after[:at], char_props_before[:at])
                self.assertEqual(char_props_after[at + width:],
                                 char_props_before[at + width:])

            elif operation == "insert":
                anchor = rng.choice(paragraphs)
                anchor_pid = anchor.get(w14("paraId"))
                anchor_xml = etree.tostring(anchor)
                ids_before = set(expected_order)
                insert_serial += 1
                text = (
                    f"insert-{insert_serial}-{step} "
                    f"{UNICODE_FRAGMENTS[rng.randrange(len(UNICODE_FRAGMENTS))]}"
                )
                result = document.insert(
                    anchor_pid,
                    text,
                    style="Normal" if step % 3 == 0 else None,
                    expect_hash=document.hash_of(anchor),
                )
                pid = result["pid"]
                self.assertNotIn(pid, ids_before)
                self.assertNotIn(pid, retired_ids)
                self.assertEqual(etree.tostring(anchor), anchor_xml)
                index = expected_order.index(anchor_pid) + 1
                expected_order.insert(index, pid)
                expected_text[pid] = text

            elif operation == "delete":
                candidates = [p for p in safe if len(paragraphs) > 2]
                if candidates:
                    paragraph = rng.choice(candidates)
                    pid = paragraph.get(w14("paraId"))
                    text = document.text_of(paragraph)
                    result = document.delete(pid, expect_hash=document.hash_of(paragraph))
                    self.assertEqual(result["text"], text)
                    expected_order.remove(pid)
                    del expected_text[pid]
                    retired_ids.add(pid)

            elif operation == "merge":
                safe_ids = {p.get(w14("paraId")) for p in safe}
                pairs = []
                if len(paragraphs) > 2:
                    for first, second in zip(paragraphs, paragraphs[1:]):
                        first_props = first.find(w("pPr"))
                        second_props = second.find(w("pPr"))
                        if (first.get(w14("paraId")) in safe_ids and
                                second.get(w14("paraId")) in safe_ids and
                                first.getparent() is second.getparent() and
                                first.getnext() is second and
                                not (first_props is not None and
                                     first_props.find(w("sectPr")) is not None) and
                                not (second_props is not None and
                                     second_props.find(w("sectPr")) is not None)):
                            pairs.append((first, second))
                if pairs:
                    first, second = rng.choice(pairs)
                    first_pid = first.get(w14("paraId"))
                    second_pid = second.get(w14("paraId"))
                    first_text = document.text_of(first)
                    second_text = document.text_of(second)
                    combined_props = rpr_xml(first) + rpr_xml(second)
                    result = document.merge(
                        first_pid,
                        second_pid,
                        first_text,
                        second_text,
                        expect_first_hash=document.hash_of(first),
                        expect_second_hash=document.hash_of(second),
                    )
                    self.assertEqual(result["after"], first_text + second_text)
                    self.assertEqual(result["removed_pid"], second_pid)
                    self.assertEqual(rpr_xml(document.para(first_pid)), combined_props)
                    expected_text[first_pid] = first_text + second_text
                    expected_order.remove(second_pid)
                    del expected_text[second_pid]
                    retired_ids.add(second_pid)

            assert_model(self, document, expected_order, expected_text)

            if step % 11 == 10:
                ids_before = list(expected_order)
                props_before = {
                    p.get(w14("paraId")): rpr_xml(p) for p in document.paragraphs()
                }
                document.save()
                self.assert_archive_ok()
                document = self.open()
                self.assertEqual(document._stamped, 0)
                self.assertEqual(
                    [p.get(w14("paraId")) for p in document.paragraphs()],
                    ids_before,
                )
                self.assertEqual(
                    {p.get(w14("paraId")): rpr_xml(p) for p in document.paragraphs()},
                    props_before,
                )
                assert_model(self, document, expected_order, expected_text)

        document.save()
        self.assert_archive_ok()
        reopened = self.open()
        self.assertEqual(reopened._stamped, 0)
        assert_model(self, reopened, expected_order, expected_text)

        changed_parts = {
            name
            for name in set(self.original_parts) | set(zip_hashes(self.path))
            if self.original_parts.get(name) != zip_hashes(self.path).get(name)
        }
        self.assertEqual(changed_parts, {"word/document.xml"})


if __name__ == "__main__":
    unittest.main(verbosity=2)
